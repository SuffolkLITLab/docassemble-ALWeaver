"""SME wording workbooks with source-checked, scalar-only YAML imports."""

import ast
import base64
from collections import Counter
import hashlib
from io import BytesIO
import json
import re
import zipfile
from typing import Any, Dict, List

import yaml
from mako.lexer import Lexer

from .editor_utils import source_revision
from .interview_scan import FIELD_OPTIONS

# Only display text, never expressions, answer values, metadata or code blocks.
SCREEN_TEXT = {
    "question",
    "subquestion",
    "help",
    "under",
    "pre",
    "post",
    "right",
    "continue button label",
    "back button label",
    "resume button label",
    "note",
    "html",
    "raw html",
    "label",
    "hint",
    "under text",
    "edit header",
    "not available label",
    "help label",
    "corner back button label",
}
FIELD_TEXT = {"label", "hint", "help", "note", "html", "raw html", "under text"}
CHOICE_OPTIONS = {
    "image",
    "help",
    "default",
    "disable",
    "css class",
    "color",
    "code",
    "show if",
    "group",
    "label",
    "value",
    "url",
}
SCHEMA = "ALWeaver wording workbook 1"
MAX_ROWS = 10000
MAX_UNCOMPRESSED = 80 * 1024 * 1024


def protected_parts(text: str) -> List[tuple[str, int]]:
    """Keep source intact, with Dashboard's green control / blue expression colors.

    Mako's own lexer handles nested braces and quoted closing braces. Unlike the
    webapp translation helper, this does not rewrite emoji or HTML placeholders.
    """
    parts: List[tuple[str, int]] = []
    pattern = re.compile(
        r"\$\{|<%|</%|^[ \t]*%(?!%)|^[ \t]*##|<[/!]?[A-Za-z][^>]*>"
        # Display directives, including those with arguments: [FILE x.png, 50%].
        r"|\[(?:FILE|QR|YOUTUBE|VIMEO|TARGET|EMOJI|FIELD)\s[^\]\n]*\]"
        r"|\[(?:[A-Z][A-Z _-]*|:[\w-]+:)\]",
        re.M,
    )
    position = 0
    while match := pattern.search(text, position):
        start = match.start()
        if start > position:
            parts.append((text[position:start], 0))
        token = match.group()
        kind = 2
        if token == "${":
            lexer = Lexer(text)
            lexer.textlength = len(text)
            lexer.match_position = match.end()
            lexer.parse_until_text(True, r"}")
            end = lexer.match_position
        elif token in ("<%", "</%"):
            # Use Mako's quote-aware scanner so a literal '%>' inside Python
            # cannot make part of a code block appear to be editable prose.
            is_tag = token == "</%" or bool(
                re.match(
                    r"(?:include|namespace|inherit|page|def|block|call|text|doc|filter)\b|[\w]+:",
                    text[match.end() :],
                )
            )
            lexer = Lexer(text)
            lexer.textlength = len(text)
            lexer.match_position = match.end()
            lexer.parse_until_text(False, r">" if is_tag else r"%>")
            end = lexer.match_position
            kind = 1
        elif token.lstrip().startswith(("%", "##")):
            end = text.find("\n", match.end())
            end = len(text) if end < 0 else end + 1
            kind = 1
        else:
            end = match.end()
        parts.append((text[start:end], kind))
        position = end
    if position < len(text):
        parts.append((text[position:], 0))
    return parts or [(text, 0)]


def validate_wording(original: str, edited: str) -> None:
    """Allow moving intact expressions within a control region, not across it."""
    old, new = protected_parts(original), protected_parts(edited)

    def regions(parts: List[tuple[str, int]]) -> tuple[List[str], List[Counter]]:
        controls = []
        groups: List[Counter] = [Counter()]
        for text, kind in parts:
            if kind == 1:
                controls.append(text)
                groups.append(Counter())
            elif kind == 2:
                groups[-1][text] += 1
        return controls, groups

    if regions(old) != regions(new):
        raise ValueError(
            "Highlighted code was changed, removed, duplicated, or moved across a control line."
        )
    # No execution. Parsing catches broken Mako introduced through prose edits.
    Lexer(edited).parse()


def workbook_context(scan: Dict[str, Any], sources: Dict[str, str]) -> Dict[str, Any]:
    """Map the main-order walk to workbook tabs, including inherited screens."""
    nodes = {
        name: list(yaml.compose_all(source, Loader=yaml.SafeLoader))
        for name, source in sources.items()
    }
    by_scan_id = {block["scan_id"]: block for block in scan["blocks"]}
    screen_ids = []
    inherited = []
    title_screen = None
    for scan_id in scan.get("screen_order", []):
        block = by_scan_id[scan_id]
        name = block["sourceFile"]
        for doc, node in enumerate(nodes.get(name, [])):
            if (
                node is None
                or not block["line_start"]
                <= node.start_mark.line + 1
                <= block["line_end"]
            ):
                continue
            screen_id = f"{name}#{doc}"
            if screen_id not in screen_ids:
                screen_ids.append(screen_id)
            if ":" in name:
                inherited.append(screen_id)
            data = block.get("data") or {}
            if title_screen is None and any(
                "interview_short_title" in str(data.get(k, ""))
                for k in ("question", "subquestion")
            ):
                title_screen = screen_id
            break
    return {
        "screen_ids": screen_ids,
        "inherited": inherited,
        "title_screen": title_screen,
    }


def _workbook_inventory(
    sources: Dict[str, str], context: Any = None
) -> List[Dict[str, Any]]:
    context = context or {}
    inherited = set(context.get("inherited", []))
    result = []
    for filename, source in sources.items():
        for item in text_inventory(filename, source):
            screen_id = f"{filename}#{item['document']}"
            if ":" in filename and screen_id not in inherited:
                continue
            item["readonly"] = ":" in filename
            item["screen_id"] = screen_id
            if item["label"] == "Action title" and context.get("title_screen"):
                item["screen_id"] = context["title_screen"]
            result.append(item)
    return result


def workbook_screens(
    sources: Dict[str, str], context: Any = None
) -> List[Dict[str, Any]]:
    """Use the same tab inventory for required PNGs and workbook generation."""
    documents = {
        name: list(yaml.safe_load_all(source)) for name, source in sources.items()
    }
    grouped = {item["screen_id"] for item in _workbook_inventory(sources, context)}
    order = list((context or {}).get("screen_ids", []))
    result = []
    for screen_id in sorted(
        grouped, key=lambda key: (order.index(key) if key in order else len(order), key)
    ):
        filename, doc = screen_id.rsplit("#", 1)
        data = documents[filename][int(doc)]
        if isinstance(data, dict) and any(
            key in data for key in ("question", "review", "table")
        ):
            result.append({"id": screen_id, "data": data})
    return result


def _action_title(node: Any, source: str) -> Any:
    """Locate only an unconditional string literal, using UTF-8 AST offsets."""
    if not isinstance(node, yaml.ScalarNode) or node.style != "|":
        return None
    try:
        tree = ast.parse(node.value)
    except SyntaxError:
        return None
    matches = []
    for statement in tree.body:
        value: Any
        if isinstance(statement, ast.Assign) and len(statement.targets) == 1:
            target, value = statement.targets[0], statement.value
        elif isinstance(statement, ast.AnnAssign):
            target, value = statement.target, statement.value
        else:
            continue
        if isinstance(target, ast.Name) and target.id == "interview_short_title":
            if not isinstance(value, ast.Constant) or not isinstance(value.value, str):
                return None
            matches.append(value)
    if len(matches) != 1:
        return None
    value = matches[0]
    if value.end_lineno is None or value.end_col_offset is None:
        return None
    raw_lines = source[node.start_mark.index : node.end_mark.index].splitlines(
        keepends=True
    )
    code_lines = node.value.splitlines(keepends=True)

    def offset(line: int, column: int) -> int:
        raw = raw_lines[line]
        decoded = code_lines[line - 1]
        if raw.rstrip("\r\n").endswith(decoded.rstrip("\r\n")):
            indent = len(raw.rstrip("\r\n")) - len(decoded.rstrip("\r\n"))
        else:
            raise ValueError("Could not locate the action title safely.")
        char_column = len(decoded.encode("utf-8")[:column].decode("utf-8"))
        return (
            node.start_mark.index
            + sum(len(text) for text in raw_lines[:line])
            + indent
            + char_column
        )

    return (
        value.value,
        offset(value.lineno, value.col_offset),
        offset(value.end_lineno, value.end_col_offset),
    )


def text_inventory(filename: str, source: str) -> List[Dict[str, Any]]:
    """Return editable scalar locations, including shorthand field/choice labels."""
    result: List[Dict[str, Any]] = []
    revision = source_revision(source)
    documents = list(yaml.compose_all(source, Loader=yaml.SafeLoader))
    seen: set[int] = set()
    repeated: set[int] = set()

    def count(node: Any) -> None:
        if node is None:
            return
        if id(node) in seen:
            repeated.add(id(node))
            return
        seen.add(id(node))
        if isinstance(node, yaml.MappingNode):
            for k, v in node.value:
                count(k)
                count(v)
        elif isinstance(node, yaml.SequenceNode):
            for item in node.value:
                count(item)

    for document in documents:
        count(document)

    def add(node: Any, document: int, path: list, label: str, screen: str) -> None:
        if not isinstance(node, yaml.ScalarNode) or node.tag != "tag:yaml.org,2002:str":
            return
        if id(node) in repeated or source[
            node.start_mark.index : node.start_mark.index + 1
        ] in {"&", "*"}:
            return  # Aliases affect more than one location and are not a scalar edit.
        if len(node.value) > 32767:
            raise ValueError(
                f"{filename}: text is longer than Excel permits in a cell."
            )
        protected_parts(node.value)
        identity = json.dumps([filename, revision, document, path], ensure_ascii=False)
        result.append(
            {
                "id": hashlib.sha256(identity.encode()).hexdigest(),
                "filename": filename,
                "revision": revision,
                "document": document,
                "path": path,
                "label": label,
                "screen": screen,
                "original": node.value,
                "start": node.start_mark.index,
                "end": node.end_mark.index,
                "style": node.style,
                "line": node.start_mark.line + 1,
            }
        )

    def choices(node: Any, doc: int, path: list, screen: str) -> None:
        if not isinstance(node, yaml.SequenceNode) or id(node) in repeated:
            return
        for i, item in enumerate(node.value):
            if (
                isinstance(item, yaml.ScalarNode)
                and item.tag == "tag:yaml.org,2002:str"
            ):
                before = len(result)
                add(item, doc, path + [i], "Answer label", screen)
                if len(result) > before:
                    result[-1]["choice_value"] = item.value
            if isinstance(item, yaml.MappingNode) and id(item) not in repeated:
                keys = {k.value for k, _ in item.value}
                # With both "label" and "value", the label is a value, not a key.
                labelled = {"label", "value"} <= keys
                for j, (k, v) in enumerate(item.value):
                    if labelled and k.value == "label":
                        add(v, doc, path + [i, j, "value"], "Answer label", screen)
                    elif k.value not in CHOICE_OPTIONS and not labelled:
                        add(k, doc, path + [i, j, "key"], "Answer label", screen)
                    elif k.value == "help":
                        add(v, doc, path + [i, j, "value"], "Answer help", screen)

    for doc, node in enumerate(documents):
        if not isinstance(node, yaml.MappingNode) or id(node) in repeated:
            continue
        pairs = {k.value: v for k, v in node.value}
        if (
            "code" in pairs
            and id(pairs["code"]) not in repeated
            and ":" not in filename
        ):
            located = _action_title(pairs["code"], source)
            if located:
                original, start, end = located
                identity = json.dumps(
                    [filename, revision, doc, ["action_title"]], ensure_ascii=False
                )
                result.append(
                    {
                        "id": hashlib.sha256(identity.encode()).hexdigest(),
                        "filename": filename,
                        "revision": revision,
                        "document": doc,
                        "path": ["action_title"],
                        "label": "Action title",
                        "screen": "Interview action title",
                        "original": original,
                        "start": start,
                        "end": end,
                        "style": "python",
                        "line": source.count("\n", 0, start) + 1,
                    }
                )
        if not any(k in pairs for k in ("question", "review", "table", "template")):
            continue
        screen_node = pairs.get("question") or pairs.get("subject")
        screen = (
            screen_node.value
            if isinstance(screen_node, yaml.ScalarNode)
            else f"Screen {doc + 1}"
        )
        screen = screen.split("\n")[0][:120]
        for j, (k, v) in enumerate(node.value):
            path = [j, "value"]
            if k.value in SCREEN_TEXT or (
                "template" in pairs and k.value in {"content", "subject"}
            ):
                add(
                    v,
                    doc,
                    path,
                    (
                        "Action title"
                        if pairs.get("template")
                        and pairs["template"].value == "interview_short_title"
                        and k.value == "content"
                        else k.value.capitalize()
                    ),
                    screen,
                )
            elif k.value == "columns" and isinstance(v, yaml.SequenceNode):
                for column_index, column in enumerate(v.value):
                    if isinstance(column, yaml.MappingNode):
                        for key_index, (column_label, expression) in enumerate(
                            column.value
                        ):
                            if column_label.value not in {"code", "header"}:
                                add(
                                    column_label,
                                    doc,
                                    path + [column_index, key_index, "key"],
                                    "Column heading",
                                    screen,
                                )
                            elif column_label.value == "header":
                                add(
                                    expression,
                                    doc,
                                    path + [column_index, key_index, "value"],
                                    "Column heading",
                                    screen,
                                )
            elif k.value in {"buttons", "choices"}:
                choices(v, doc, path, screen)
            elif (
                k.value in {"fields", "review"}
                and isinstance(v, yaml.SequenceNode)
                and id(v) not in repeated
            ):
                for i, field in enumerate(v.value):
                    if not isinstance(field, yaml.MappingNode) or id(field) in repeated:
                        continue
                    for n, (fk, fv) in enumerate(field.value):
                        field_path = path + [i, n]
                        if fk.value in FIELD_TEXT or (
                            k.value == "review" and fk.value == "button"
                        ):
                            add(
                                fv,
                                doc,
                                field_path + ["value"],
                                fk.value.capitalize(),
                                screen,
                            )
                        elif fk.value == "choices":
                            choices(fv, doc, field_path + ["value"], screen)
                        elif (
                            k.value == "fields"
                            and fk.value not in FIELD_OPTIONS | {"no label"}
                            and not any(
                                key.value == "field" for key, value in field.value
                            )
                        ):
                            add(fk, doc, field_path + ["key"], "Field label", screen)
    if len(result) > MAX_ROWS:
        raise ValueError("Too many wording rows in this interview.")
    return result


def export_workbook(
    sources: Dict[str, str], previews: Dict[str, str], context: Any = None
) -> bytes:
    """Write instructions, screen tabs with PNGs, and hidden round-trip metadata."""
    import xlsxwriter
    from PIL import Image

    output = BytesIO()
    workbook = xlsxwriter.Workbook(
        output,
        {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False},
    )
    fixed = workbook.add_format({"text_wrap": True, "valign": "top"})
    editable = workbook.add_format(
        {"text_wrap": True, "valign": "top", "locked": False, "bg_color": "#FFF2CC"}
    )
    heading = workbook.add_format(
        {"bold": True, "bg_color": "#17365D", "font_color": "white", "text_wrap": True}
    )
    colors = [
        workbook.add_format({}),
        workbook.add_format({"font_color": "#008000", "bold": True}),
        workbook.add_format({"font_color": "#0000FF", "bold": True}),
    ]
    instructions = workbook.add_worksheet("Instructions")
    instructions.set_column("A:A", 110)
    instructions.write(0, 0, "Review and update the interview wording", heading)
    notes = [
        "Each screen tab shows an illustrative PNG preview beside its wording. Yellow cells are yours to edit. Only display text is included; interview logic and metadata stay in the source.",
        "Edit the Revised wording column. Leave the original wording, row IDs, sheet names and other cells alone. Do not add or delete rows or sheets. To remove wording, clear its yellow cell.",
        "Green text is a control line or code tag. Blue text is an expression, HTML tag or display directive. Leave highlighted code exactly as written, including spaces, punctuation and capitalization.",
        "You may move a whole blue expression within the same paragraph/control region: Hello ${ users[0] } → ${ users[0] }, hello. Keep the entire ${ ... } together.",
        "Do not change ${ users[0] } to ${ user[0] }, remove it, or copy it twice. Do not move text or expressions across green % if / % else / % endif lines; that would change when they appear.",
        "Example: % if has_children: and % endif must remain on their own lines, in the same order. Change only the wording between them. The importer checks highlighted code even if your spreadsheet app loses its colors.",
        "Previews are static illustrations, not a live interview. Expressions and conditional fields may appear differently at runtime. Template text has no independent screen preview. Shared YAML aliases are excluded because changing one can affect multiple locations. Inherited screens from the main order are shown read-only. Their wording cannot be changed here. A yellow Action title cell changes only this interview’s title, even when it appears on an inherited intro screen.",
        "Save as XLSX. In the Weaver choose Import wording workbook, inspect the proposed changes, then Apply. If source files changed after export, export a fresh workbook and transfer your edits; nothing is applied when validation fails.",
    ]
    for row, note in enumerate(notes, 2):
        instructions.write_string(row, 0, note, fixed)
        instructions.set_row(row, 60)
    manifest = workbook.add_worksheet("_ALWeaver")
    manifest.write_string(0, 0, SCHEMA)
    manifest.hide()
    manifest_row = 1
    groups: Dict[str, List[Dict[str, Any]]] = {}
    inventory = _workbook_inventory(sources, context)
    order = (context or {}).get("screen_ids", [])
    for item in sorted(
        inventory,
        key=lambda item: (
            order.index(item["screen_id"]) if item["screen_id"] in order else len(order)
        ),
    ):
        groups.setdefault(item["screen_id"], []).append(item)
    if not groups:
        raise ValueError("No editable user-facing wording was found.")
    if sum(len(items) for items in groups.values()) > MAX_ROWS:
        raise ValueError("Too many wording rows for one workbook.")
    for index, (screen_id, items) in enumerate(groups.items(), 1):
        filename, document_text = screen_id.rsplit("#", 1)
        document = int(document_text)
        screen_title = next(
            (item["screen"] for item in items if item["label"] != "Action title"),
            items[0]["screen"],
        )
        sheet = workbook.add_worksheet(f"Screen {index:03d}")
        sheet.freeze_panes(3, 2)
        sheet.hide_gridlines(2)
        sheet.set_landscape()
        sheet.fit_to_pages(1, 0)
        sheet.set_row(0, 36)
        sheet.set_row(1, 30)
        sheet.set_column("A:A", 18)
        sheet.set_column("B:C", 55)
        sheet.set_column("D:D", 70)
        sheet.set_column("E:E", 10, None, {"hidden": True})
        sheet.merge_range("A1:D1", screen_title, heading)
        sheet.merge_range(
            "A2:D2",
            f"{filename} · block {document + 1} · "
            + (
                "Inherited wording is read-only; the yellow action title is editable"
                if ":" in filename
                else "static preview"
            ),
            fixed,
        )
        sheet.write_row(
            2,
            0,
            [
                "Text part",
                "Original wording",
                "Revised wording",
                "Screen preview",
                "Row ID",
            ],
            heading,
        )
        png = previews.get(screen_id)
        if png:
            raw = base64.b64decode(png.split(",")[-1], validate=True)
            image = Image.open(BytesIO(raw))
            if image.format != "PNG" or image.width * image.height > 16000000:
                raise ValueError("Invalid or oversized preview PNG.")
            image.verify()
            scale = min(1, 480 / image.width)
            sheet.insert_image(
                "D4",
                "screen.png",
                {
                    "image_data": BytesIO(raw),
                    "description": items[0]["screen"] + " — static screen preview",
                    "x_scale": scale,
                    "y_scale": scale,
                    "object_position": 1,
                },
            )
        else:
            sheet.write_string(
                3, 3, "No standalone screen preview for this text.", fixed
            )
        for row, item in enumerate(items, 3):
            sheet.write_string(row, 0, item["label"], fixed)
            sheet.write_string(row, 4, item["id"])
            parts = protected_parts(item["original"])
            for col, fmt in [(1, fixed), (2, fixed if item["readonly"] else editable)]:
                runs: List[Any] = []
                for text, kind in parts:
                    if text:
                        runs.extend([colors[kind], text])
                if len(runs) > 2:
                    sheet.write_rich_string(row, col, *runs, fmt)
                else:
                    single_format = workbook.add_format(
                        {
                            "text_wrap": True,
                            "valign": "top",
                            "locked": col != 2 or item["readonly"],
                            "bg_color": (
                                "#FFF2CC"
                                if col == 2 and not item["readonly"]
                                else "#FFFFFF"
                            ),
                            "font_color": ["#000000", "#008000", "#0000FF"][
                                parts[0][1]
                            ],
                            "bold": bool(parts[0][1]),
                        }
                    )
                    sheet.write_string(row, col, item["original"], single_format)
            sheet.set_row(
                row,
                min(
                    300,
                    max(
                        65,
                        18
                        * (
                            len(item["original"]) // 55
                            + item["original"].count("\n")
                            + 2
                        ),
                    ),
                ),
            )
            metadata = {
                k: v
                for k, v in item.items()
                if k not in {"start", "end", "style", "original"}
            }
            metadata["sheet"] = sheet.name
            manifest.write_string(
                manifest_row, 0, json.dumps(metadata, ensure_ascii=False)
            )
            manifest_row += 1
        sheet.protect("", {"select_locked_cells": True, "select_unlocked_cells": True})
    workbook.close()
    return output.getvalue()


def _block_scalar(source: str, item: Dict[str, Any], edited: str) -> str:
    """Rewrite a block scalar as a literal at its original indentation.

    PyYAML's range for a block scalar runs through the blank lines after it, so
    those are kept as separators. Text a literal cannot hold safely is written
    as a quoted scalar on the same line instead.
    """
    original_token = source[item["start"] : item["end"]]
    header, _, body = original_token.partition("\n")
    lines = body.split("\n")
    # Trailing whitespace-only lines, minus those a keep chomp makes part of the value.
    trailing = 0
    for line in reversed(lines[:-1]):
        if line.strip():
            break
        trailing += 1
    original = item["original"]
    if header.split("#", 1)[0].strip().endswith("+") and original.endswith("\n"):
        trailing -= len(original) - len(original.rstrip("\n")) - 1
    separators = "\n" * max(trailing, 0)
    first_line = next((line for line in edited.split("\n") if line.strip()), "")
    if re.search(
        r"\d", header.split("#", 1)[0]
    ) or first_line[  # explicit indentation indicator
        :1
    ] in {
        " ",
        "\t",
    }:  # leading spaces would read as indentation
        return json.dumps(edited, ensure_ascii=False) + "\n" + separators
    content = next((line for line in lines if line.strip()), None)
    if content is not None:
        indent = len(content) - len(content.lstrip(" "))
    else:
        # The key's column, past any "- " sequence markers, plus two.
        line = source[: item["start"]].rsplit("\n", 1)[-1]
        indent = re.match(r"[ \t]*(?:-[ \t]+)*", line).end() + 2  # type: ignore[union-attr]
    comment = " " + header[header.index("#") :] if "#" in header else ""
    # A literal style faithfully preserves edited newlines, including trailing ones.
    chomp = "+" if edited.endswith("\n\n") else "" if edited.endswith("\n") else "-"
    if chomp == "+":
        separators = ""  # Blank lines after a kept block would join its value.
    text = "".join(
        (" " * indent + line if line else "") + "\n" for line in edited.splitlines()
    )
    return "|" + chomp + comment + "\n" + text + separators


def _replacement(source: str, item: Dict[str, Any], edited: str) -> str:
    """Replace one scalar, leaving all surrounding source untouched."""
    if item["style"] == "python":
        return repr(edited)
    if "choice_value" in item:
        return json.dumps({edited: item["choice_value"]}, ensure_ascii=False)
    if item["style"] == "'":
        return (
            "'" + edited.replace("'", "''") + "'"
            if "\n" not in edited
            else json.dumps(edited, ensure_ascii=False)
        )
    if item["style"] in {"|", ">"}:
        return _block_scalar(source, item, edited)
    if item["style"] is None and "\n" not in edited:
        try:
            if yaml.safe_load(edited) == edited and not re.search(
                r'[:#\[\]{},&*!|>\'"%@`]', edited
            ):
                return edited
        except yaml.YAMLError:
            pass
    return json.dumps(edited, ensure_ascii=False)


def _validate_unique_keys(source: str) -> None:
    seen: set[int] = set()

    def visit(node: Any) -> None:
        if node is None or id(node) in seen:
            return
        seen.add(id(node))
        if isinstance(node, yaml.MappingNode):
            keys = [k.value for k, _ in node.value if isinstance(k, yaml.ScalarNode)]
            if len(set(keys)) != len(keys):
                raise ValueError("An edited label would duplicate a YAML key.")
            for k, value in node.value:
                visit(value)
        elif isinstance(node, yaml.SequenceNode):
            for value in node.value:
                visit(value)

    for document in yaml.compose_all(source, Loader=yaml.SafeLoader):
        visit(document)


def import_workbook(
    content: bytes, sources: Dict[str, str], context: Any = None
) -> Dict[str, Any]:
    """Preflight every row, then return exact source patches without writing."""
    import openpyxl

    with zipfile.ZipFile(BytesIO(content)) as archive:
        if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED:
            raise ValueError("The workbook is too large.")
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False, read_only=True)
    try:
        if "_ALWeaver" not in workbook or workbook["_ALWeaver"]["A1"].value != SCHEMA:
            raise ValueError("This is not a Weaver wording workbook.")
        inventory = {item["id"]: item for item in _workbook_inventory(sources, context)}
        metadata: Dict[str, Dict[str, Any]] = {}
        for row in workbook["_ALWeaver"].iter_rows(min_row=2, max_col=1):
            if row[0].value is None:
                continue
            item = json.loads(row[0].value)
            if len(metadata) >= MAX_ROWS or item["id"] in metadata:
                raise ValueError("Invalid or duplicate workbook row IDs.")
            current = inventory.get(item["id"])
            if not current or any(
                item.get(k) != current[k]
                for k in ("filename", "revision", "document", "path")
            ):
                raise ValueError(
                    "Source changed since export, or workbook metadata was edited. Export a fresh workbook."
                )
            metadata[item["id"]] = item
        if set(metadata) != set(inventory):
            raise ValueError(
                "Workbook rows or metadata are missing, or the interview scope changed. Export a fresh workbook."
            )
        if not metadata:
            raise ValueError("The workbook has no wording rows.")
        seen = set()
        edits: Dict[str, List[tuple]] = {}
        changes = []
        for sheet in workbook:
            if sheet.max_row > MAX_ROWS + 4:
                raise ValueError("Too many rows in the workbook.")
            if sheet.title in {"Instructions", "_ALWeaver"}:
                continue
            for row in sheet.iter_rows(min_row=4, max_col=5):
                if all(cell.value is None for cell in row):
                    continue
                identity = row[4].value
                if (
                    identity not in metadata
                    or identity in seen
                    or metadata[identity]["sheet"] != sheet.title
                ):
                    raise ValueError(
                        f"{sheet.title}: missing, duplicate, or moved row ID."
                    )
                seen.add(identity)
                item = inventory[identity]
                if row[1].value != (item["original"] or None):
                    raise ValueError(f"{sheet.title}: original wording was changed.")
                if row[2].data_type == "f":
                    raise ValueError(
                        f"{sheet.title}: revised wording must be text, not a formula."
                    )
                edited = row[2].value if row[2].value is not None else ""
                if not isinstance(edited, str):
                    raise ValueError(f"{sheet.title}: revised wording must be text.")
                if edited == item["original"]:
                    continue
                if item["readonly"]:
                    raise ValueError(
                        f"{sheet.title}: inherited wording is read-only. Edit the local Action title instead."
                    )
                if item["label"] == "Field label" and edited in FIELD_OPTIONS:
                    raise ValueError(
                        "A field label cannot become an interview directive."
                    )
                if item["label"] in {
                    "Answer label",
                    "Column heading",
                } and edited in CHOICE_OPTIONS | {"header"}:
                    raise ValueError(
                        "This label is reserved for interview code or options."
                    )
                validate_wording(item["original"], edited)
                filename = item["filename"]
                edits.setdefault(filename, []).append(
                    (
                        item["start"],
                        item["end"],
                        _replacement(sources[filename], item, edited),
                    )
                )
                changes.append(
                    {
                        "filename": filename,
                        "screen": item["screen"],
                        "part": item["label"],
                        "original": item["original"],
                        "edited": edited,
                    }
                )
        if seen != set(metadata):
            raise ValueError("Rows or screen sheets were removed from the workbook.")
        updated = {}
        for filename, replacements in edits.items():
            source = sources[filename]
            for start, end, text in sorted(replacements, reverse=True):
                source = source[:start] + text + source[end:]
            # Refuse duplicate keys: a renamed shorthand label must never silently
            # overwrite another field, choice, or directive.
            _validate_unique_keys(source)
            updated[filename] = source
        return {"updated": updated, "changes": changes}
    finally:
        workbook.close()
