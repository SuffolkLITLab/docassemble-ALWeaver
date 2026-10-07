from io import BytesIO
import json

import openpyxl
import pytest
import yaml

from .text_workbook import (
    export_workbook,
    import_workbook,
    protected_parts,
    text_inventory,
    validate_wording,
)

SOURCE = """# Keep my comment
---
id: greeting
question: Hello ${ users[0] }
subquestion: | # Keep this too
  % if eligible:
  Welcome, ${ users[0] }.
  % endif
fields:
  - Your name: user_name
    hint: First and last name
    required: true
  - label: Favorite color
    field: color
    choices:
      - Red: red_value
      - Blue: blue_value
---
metadata:
  title: Not editable metadata
---
code: |
  secret = 'Not editable Python'
"""


def edit_workbook(raw, edits):
    book = openpyxl.load_workbook(BytesIO(raw))
    for sheet in book:
        if sheet.title.startswith("Screen "):
            for row in sheet.iter_rows(min_row=4):
                if row[1].value in edits:
                    row[2].value = edits[row[1].value]
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def test_inventory_only_display_text():
    items = text_inventory("main.yml", SOURCE)
    values = [i["original"] for i in items]
    assert "Your name" in values and "Red" in values
    assert "red_value" not in values and "user_name" not in values
    assert not any("Not editable" in text for text in values)


def test_round_trip_preserves_source_and_allows_expression_movement():
    raw = export_workbook({"main.yml": SOURCE}, {})
    assert import_workbook(raw, {"main.yml": SOURCE})["updated"] == {}
    changed = edit_workbook(
        raw, {"Hello ${ users[0] }": "${ users[0] }, hello!", "Your name": "Full name"}
    )
    result = import_workbook(changed, {"main.yml": SOURCE})
    updated = result["updated"]["main.yml"]
    assert updated == SOURCE.replace(
        "Hello ${ users[0] }", '"${ users[0] }, hello!"'
    ).replace("Your name:", "Full name:")
    assert len(result["changes"]) == 2


def test_rejects_code_changes_and_stale_source():
    raw = export_workbook({"main.yml": SOURCE}, {})
    changed = edit_workbook(raw, {"Hello ${ users[0] }": "Hello ${ users[1] }"})
    with pytest.raises(ValueError, match="Highlighted"):
        import_workbook(changed, {"main.yml": SOURCE})
    with pytest.raises(ValueError, match="Source changed"):
        import_workbook(raw, {"main.yml": SOURCE + "# external change\n"})


def test_nested_expression_and_control_regions():
    text = 'Hi ${ {"}": users[0]}["}"] }!'
    assert "".join(p for p, _ in protected_parts(text)) == text
    validate_wording(text, 'Welcome ${ {"}": users[0]}["}"] }!')
    with pytest.raises(ValueError):
        validate_wording(
            "% if ok:\n${ name }\n% endif\n", "${ name }\n% if ok:\n\n% endif\n"
        )


def test_multiline_blank_and_formula_text():
    source = 'question: "Hello" # comment\nsubquestion: |\n  Old text.\nyesno: done\n'
    raw = export_workbook({"main.yml": source}, {})
    edited = edit_workbook(
        raw, {"Hello": "", "Old text.\n": "New text.\n\nMore text.\n"}
    )
    updated = import_workbook(edited, {"main.yml": source})["updated"]["main.yml"]
    assert "# comment" in updated
    data = yaml.safe_load(updated)
    assert data["question"] == ""
    assert data["subquestion"] == "New text.\n\nMore text.\n"
    bad = edit_workbook(raw, {"Hello": '=HYPERLINK("https://example.org")'})
    with pytest.raises(ValueError, match="formula"):
        import_workbook(bad, {"main.yml": source})


def test_missing_rows_and_duplicate_ids_rejected():
    raw = export_workbook({"main.yml": SOURCE}, {})
    book = openpyxl.load_workbook(BytesIO(raw))
    book["Screen 001"].delete_rows(4)
    out = BytesIO()
    book.save(out)
    with pytest.raises(ValueError, match="removed"):
        import_workbook(out.getvalue(), {"main.yml": SOURCE})


def test_scalar_choice_preserves_stored_value_and_reserved_labels_rejected():
    source = "question: Choose\nfield: color\nchoices:\n  - Red\n  - Blue\n"
    raw = export_workbook({"main.yml": source}, {})
    changed = edit_workbook(raw, {"Red": "Scarlet"})
    updated = import_workbook(changed, {"main.yml": source})["updated"]["main.yml"]
    assert yaml.safe_load(updated)["choices"][0] == {"Scarlet": "Red"}
    raw = export_workbook({"main.yml": SOURCE}, {})
    with pytest.raises(ValueError, match="directive"):
        import_workbook(edit_workbook(raw, {"Your name": "code"}), {"main.yml": SOURCE})


def test_png_is_embedded_and_highlighting_survives_export():
    import base64
    from PIL import Image

    image = BytesIO()
    Image.new("RGB", (320, 200), "white").save(image, format="PNG")
    raw = export_workbook(
        {"main.yml": SOURCE},
        {"main.yml#0": base64.b64encode(image.getvalue()).decode()},
    )
    book = openpyxl.load_workbook(BytesIO(raw), rich_text=True)
    assert len(book["Screen 001"]._images) == 1
    assert book["Screen 001"]["C4"].value[1].font.color.rgb == "FF0000FF"
    assert book["_ALWeaver"].sheet_state == "hidden"
    assert book["Screen 001"]["C4"].protection.locked is False
    assert book["Screen 001"]["B4"].protection.locked is True


def test_duplicate_edited_labels_do_not_overwrite_fields():
    source = "question: Name\nfields:\n  - First: first\n    hint: Your given name\n"
    raw = export_workbook({"main.yml": source}, {})
    with pytest.raises(ValueError):
        import_workbook(edit_workbook(raw, {"First": "hint"}), {"main.yml": source})


def test_unicode_comments_multiple_files_and_duplicate_wording():
    sources = {
        "one.yml": "question: Café\nsubquestion: Hello\nyesno: ready\n",
        "two.yml": "# दूसरे\nquestion: Hello\nyesno: other\n",
    }
    raw = export_workbook(sources, {})
    changed = edit_workbook(raw, {"Hello": "Good morning"})
    result = import_workbook(changed, sources)
    assert len(result["changes"]) == 2
    assert "# दूसरे" in result["updated"]["two.yml"]
    assert "Café" in result["updated"]["one.yml"]


def test_python_block_delimiter_in_string_is_still_protected():
    original = "<% x = '%>'; y = 2 %>Hello ${ x }"
    assert protected_parts(original)[0] == ("<% x = '%>'; y = 2 %>", 1)
    with pytest.raises(ValueError, match="Highlighted"):
        validate_wording(original, "<% x = '%>'; y = 3 %>Hello ${ x }")


def test_legacy_intro_readonly_and_action_title_roundtrip():
    from .test_interview_scan import LEGACY
    from .interview_scan import scan_interview
    from .text_workbook import workbook_context, workbook_screens

    context = workbook_context(
        scan_interview(LEGACY.__getitem__, "standalone.yml"), LEGACY
    )
    screens = workbook_screens(LEGACY, context)
    assert screens[0]["data"]["question"] == "${ interview_short_title }"
    assert screens[-1]["data"]["question"] == "Download your forms"
    raw = export_workbook(LEGACY, {}, context=context)
    edited = edit_workbook(raw, {"Ask for help": "Get help today"})
    result = import_workbook(edited, LEGACY, context=context)
    assert result["updated"] == {
        "reusable.yml": LEGACY["reusable.yml"].replace(
            "'Ask for help'", "'Get help today'"
        )
    }
    edited = edit_workbook(raw, {"Shared instructions.": "Changed shared wording"})
    with pytest.raises(ValueError, match="read.only"):
        import_workbook(edited, LEGACY, context=context)


def test_unicode_action_title_exact_patch_and_computed_title_excluded():
    source = 'code: | # Keep\n  café = 1; interview_short_title = "Ask for help" # Keep too\n'
    raw = export_workbook({"main.yml": source}, {})
    result = import_workbook(
        edit_workbook(raw, {"Ask for help": "Help me"}), {"main.yml": source}
    )
    assert result["updated"]["main.yml"] == source.replace(
        '"Ask for help"', "'Help me'"
    )
    assert not text_inventory(
        "main.yml", "code: |\n  interview_short_title = make_title()\n"
    )
