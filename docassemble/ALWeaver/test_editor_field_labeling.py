# do not pre-load
"""HTTP review/worker flow and Assistant label/apply integration."""

import base64
from contextlib import ExitStack, nullcontext
import io
from pathlib import Path
import tempfile
import types
import unittest
from unittest.mock import patch

from .test_editor_api import api_editor
from .test_field_labeling import make_pdf, SOURCE
from .editor_agent_models import (
    AgentCandidate,
    AgentToolCall,
    WeaverAgentSession,
    store_agent_session,
    load_agent_session,
)
from .editor_agent_tools import ToolContext, execute_tool
from . import editor_agent_validation, editor_agent_labeling, editor_utils
from .field_labeling import template_fields, LabelingError
from .test_editor_agent_api import FakeRedis
from .source_document import source_revision


class EditorFieldLabelingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.original = self.directory / "form.pdf"
        make_pdf(self.original, ["First Name", "Last Name"])
        self.original_bytes = self.original.read_bytes()
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        from docassemble.ALDashboard import pdf_field_labeler as dashboard

        self.stack.enter_context(
            patch.object(
                dashboard,
                "_generate_ai_relabel_mapping_for_unique_fields",
                return_value={
                    "First Name": "user_name_first",
                    "Last Name": "user_name_last",
                },
            )
        )
        self.stack.enter_context(
            patch.object(api_editor, "_editor_auth_check", return_value=True)
        )
        self.stack.enter_context(
            patch.object(api_editor, "_actor_user_id", return_value=7)
        )
        self.stack.enter_context(
            patch.object(api_editor, "_editor_async_is_configured", return_value=True)
        )
        self.stack.enter_context(
            patch.object(api_editor.r, "set", return_value=True, create=True)
        )
        self.stack.enter_context(patch.object(api_editor, "_release_ai_owner"))
        self.stack.enter_context(
            patch.object(
                api_editor, "_load_job_state", return_value={"owner_user_id": 7}
            )
        )

    def request(self, action, approved=False):
        data = {"file": (io.BytesIO(self.original_bytes), "form.pdf"), "action": action}
        if approved:
            data["author_requested"] = "true"
        with api_editor.app.test_request_context(
            "/al/editor/api/field-labeling", method="POST", data=data
        ):
            response = api_editor.editor_api_field_labeling()
            return response.get_json(), response.status_code

    def test_upload_inspect_requires_choice_then_worker_returns_reviewable_copy(self):
        assessment, code = self.request("inspect")
        self.assertEqual(code, 200)
        self.assertTrue(assessment["data"]["suggest_labeling"])
        refused, code = self.request("propose")
        self.assertEqual(code, 422)
        states = {}

        def store(kind, job, value):
            states[job] = value

        def update(kind, job, **changes):
            states[job].update(changes)

        self.stack.enter_context(
            patch.object(api_editor, "_store_job_state", side_effect=store)
        )
        self.stack.enter_context(
            patch.object(api_editor, "_update_job_state", side_effect=update)
        )
        send = self.stack.enter_context(patch.object(api_editor.workerapp, "send_task"))
        queued, code = self.request("propose", approved=True)
        self.assertEqual(code, 202)
        api_editor._complete_field_labeling_job(**send.call_args.kwargs["kwargs"])
        job_id = send.call_args.kwargs["task_id"]
        self.assertEqual(states[job_id]["status"], "succeeded")
        proposal = states[job_id]["result"]
        self.assertEqual(proposal["mappings"][0]["old_name"], "First Name")
        accepted = self.directory / "accepted.pdf"
        accepted.write_bytes(base64.b64decode(proposal["content_base64"]))
        from .interview_generator import generate_interview_from_path

        draft = generate_interview_from_path(
            str(accepted),
            output_dir=str(self.directory / "draft"),
            create_package_zip=False,
            include_next_steps=False,
            use_llm_assist=False,
        )
        self.assertIn("users[0].name.first", draft.yaml_text)
        self.assertEqual(self.original.read_bytes(), self.original_bytes)
        self.stack.enter_context(
            patch.object(api_editor, "_load_job_state", return_value=states[job_id])
        )
        with patch.object(api_editor, "_actor_user_id", return_value=8):
            with api_editor.app.test_request_context(
                "/al/editor/api/field-labeling/jobs/" + job_id
            ):
                self.assertEqual(
                    api_editor.editor_api_field_labeling_job(job_id).status_code, 404
                )

    def test_assistant_tool_stages_binary_and_remaps_source_then_materializes(self):
        area = types.SimpleNamespace(finalize=lambda: None)
        self.stack.enter_context(
            patch.object(
                editor_agent_labeling,
                "template_directory",
                return_value=(area, self.directory),
            )
        )
        self.stack.enter_context(
            patch.object(
                editor_utils,
                "playground_list_yaml_files",
                return_value=[{"filename": "main.yml"}],
            )
        )
        self.stack.enter_context(
            patch.object(
                editor_agent_validation, "dayamlchecker_findings", return_value=[]
            )
        )
        candidate = AgentCandidate.from_source(SOURCE)
        context = ToolContext(
            project="default", filename="main.yml", owner_user_id=7, candidate=candidate
        )
        result = execute_tool(
            context,
            AgentToolCall(
                tool="relabel_template",
                arguments={"template_filename": "form.pdf", "author_requested": True},
            ),
        )
        self.assertEqual(result.status, "success", str(result))
        self.assertEqual(len(candidate.template_assets), 1)
        self.assertIn("users.gather()", candidate.raw_source)
        self.assertIn('"user_name_last": ${ users[0].name.last }', candidate.raw_source)
        self.assertEqual(self.original.read_bytes(), self.original_bytes)
        asset = candidate.template_assets[0]
        self.assertFalse((self.directory / asset["filename"]).exists())
        session = WeaverAgentSession(
            session_id="labels",
            owner_user_id=7,
            project="default",
            filename="main.yml",
            base_saved_revision=source_revision(SOURCE),
            original_working_source=SOURCE,
            candidate_source=SOURCE,
            candidate_revision=source_revision(SOURCE),
        )
        session.store_candidate(candidate)
        redis = FakeRedis()
        store_agent_session(redis, session)
        restored = load_agent_session(redis, "labels", 7)
        self.assertEqual(
            restored.candidate().template_assets, candidate.template_assets
        )
        self.assertNotIn("content_base64", str(restored.public_dict()))
        editor_agent_labeling.materialize_templates(
            7, "default", restored.template_assets
        )
        self.assertEqual(
            [
                field["name"]
                for field in template_fields(str(self.directory / asset["filename"]))
            ],
            ["user_name_first", "user_name_last"],
        )
        self.assertEqual(self.original.read_bytes(), self.original_bytes)
        restored.reset_candidate()
        self.assertFalse(restored.template_assets)

    def test_assistant_decline_does_not_call_labeler(self):
        candidate = AgentCandidate.from_source(SOURCE)
        context = ToolContext(
            project="default", filename="main.yml", owner_user_id=7, candidate=candidate
        )
        with patch.object(editor_agent_labeling, "stage_relabel") as stage:
            result = execute_tool(
                context,
                AgentToolCall(
                    tool="relabel_template",
                    arguments={
                        "template_filename": "form.pdf",
                        "author_requested": False,
                    },
                ),
            )
        self.assertEqual(result.status, "rejected")
        stage.assert_not_called()

    def test_labeling_job_respects_account_concurrency_limit(self):
        with patch.object(api_editor.r, "set", return_value=False):
            body, status = self.request("propose", approved=True)
        self.assertEqual(status, 429)
        self.assertEqual(body["error"]["code"], "ai_job_in_progress")

    def test_relabel_does_not_strand_a_previously_staged_word_template(self):
        self.stack.enter_context(
            patch.object(
                editor_agent_labeling,
                "template_directory",
                return_value=(
                    types.SimpleNamespace(finalize=lambda: None),
                    self.directory,
                ),
            )
        )
        self.stack.enter_context(
            patch.object(
                editor_utils,
                "playground_list_yaml_files",
                return_value=[{"filename": "main.yml"}],
            )
        )
        candidate = AgentCandidate.from_source(SOURCE)
        candidate.template_assets.append(
            {
                "filename": "previous.docx",
                "mappings": [{"old_name": "old_first", "new_name": "first_name"}],
                "added_fields": [],
            }
        )
        context = ToolContext(
            project="default", filename="main.yml", owner_user_id=7, candidate=candidate
        )
        result = execute_tool(
            context,
            AgentToolCall(
                tool="relabel_template",
                arguments={"template_filename": "form.pdf", "author_requested": True},
            ),
        )
        self.assertEqual(result.status, "rejected")
        self.assertIn("staged Word template", result.message)
        self.assertEqual(candidate.raw_source, SOURCE)
        self.assertEqual(candidate.raw_source, SOURCE)

    def test_stale_template_never_overwrites_original(self):
        area = types.SimpleNamespace(finalize=lambda: None)
        with patch.object(
            editor_agent_labeling,
            "template_directory",
            return_value=(area, self.directory),
        ):
            with self.assertRaisesRegex(LabelingError, "changed since labeling"):
                editor_agent_labeling.materialize_templates(
                    7,
                    "default",
                    [
                        {
                            "original_filename": "form.pdf",
                            "original_revision": "wrong",
                            "filename": "new.pdf",
                            "content_base64": base64.b64encode(b"bad").decode(),
                        }
                    ],
                )
        self.assertEqual(self.original.read_bytes(), self.original_bytes)
        self.assertFalse((self.directory / "new.pdf").exists())

    def test_assistant_unlabeled_pdf_adds_mapping_and_questions(self):
        from docassemble.ALDashboard import pdf_field_labeler as dashboard

        area = types.SimpleNamespace(finalize=lambda: None)
        self.stack.enter_context(
            patch.object(
                editor_agent_labeling,
                "template_directory",
                return_value=(area, self.directory),
            )
        )
        self.stack.enter_context(
            patch.object(
                editor_utils,
                "playground_list_yaml_files",
                return_value=[{"filename": "main.yml"}],
            )
        )
        self.stack.enter_context(
            patch.object(
                editor_agent_validation, "dayamlchecker_findings", return_value=[]
            )
        )
        make_pdf(self.original, [])
        source = """include:
  - docassemble.AssemblyLine:assembly_line.yml
---
id: attachment
attachment:
  variable name: form_attachment
  pdf template file: form.pdf
  fields: {}
"""

        def detect(**kwargs):
            make_pdf(kwargs["output_pdf_path"], ["favorite_color"])
            return {}

        self.stack.enter_context(
            patch.object(
                dashboard,
                "detect_pdf_fields_and_optionally_relabel",
                side_effect=detect,
            )
        )
        candidate = AgentCandidate.from_source(source)
        result = execute_tool(
            ToolContext(
                project="default",
                filename="main.yml",
                owner_user_id=7,
                candidate=candidate,
            ),
            AgentToolCall(
                tool="relabel_template",
                arguments={"template_filename": "form.pdf", "author_requested": True},
            ),
        )
        self.assertEqual(result.status, "success", str(result))
        self.assertIn("favorite_color: ${ favorite_color }", candidate.raw_source)
        self.assertIn('"Favorite color": favorite_color', candidate.raw_source)
        self.assertEqual(
            candidate.template_assets[0]["added_fields"], ["favorite_color"]
        )

    def test_assistant_docx_relabels_and_remaps_questions_order_and_review(self):
        import docx
        from docassemble.ALDashboard import docx_wrangling as dashboard

        word = self.directory / "form.docx"
        document = docx.Document()
        document.add_paragraph("{{ first_name }}")
        document.add_paragraph("{{ last_name }}")
        document.save(word)
        original = word.read_bytes()
        area = types.SimpleNamespace(finalize=lambda: None)
        self.stack.enter_context(
            patch.object(
                editor_agent_labeling,
                "template_directory",
                return_value=(area, self.directory),
            )
        )
        self.stack.enter_context(
            patch.object(
                editor_utils,
                "playground_list_yaml_files",
                return_value=[{"filename": "main.yml"}],
            )
        )
        self.stack.enter_context(
            patch.object(
                editor_agent_validation, "dayamlchecker_findings", return_value=[]
            )
        )
        self.stack.enter_context(
            patch.object(
                dashboard,
                "get_labeled_docx_runs",
                return_value=[
                    (0, 0, "{{ users[0].name.first }}", 0),
                    (1, 0, "{{ users[0].name.last }}", 0),
                ],
            )
        )
        source = (
            SOURCE.split("  pdf template file:")[0]
            + "  docx template file: form.docx\n"
        )
        candidate = AgentCandidate.from_source(source)
        result = execute_tool(
            ToolContext(
                project="default",
                filename="main.yml",
                owner_user_id=7,
                candidate=candidate,
            ),
            AgentToolCall(
                tool="relabel_template",
                arguments={"template_filename": "form.docx", "author_requested": True},
            ),
        )
        self.assertEqual(result.status, "success", str(result))
        self.assertIn("First: users[0].name.first", candidate.raw_source)
        self.assertIn("Last: users[0].name.last", candidate.raw_source)
        self.assertIn("Edit: users[0].name.first", candidate.raw_source)
        self.assertIn("users.gather()", candidate.raw_source)
        self.assertNotIn("first_name", candidate.raw_source)
        self.assertEqual(word.read_bytes(), original)
        self.assertEqual(len(candidate.template_assets[0]["mappings"]), 2)
        before = candidate.raw_source
        rename = execute_tool(
            ToolContext(
                project="default",
                filename="main.yml",
                owner_user_id=7,
                candidate=candidate,
            ),
            AgentToolCall(
                tool="rename_variables",
                arguments={
                    "renames": [
                        {"old_name": "users[0].name.first", "new_name": "renamed_first"}
                    ]
                },
            ),
        )
        self.assertEqual(rename.reason, "staged_template_reference")
        self.assertEqual(candidate.raw_source, before)
