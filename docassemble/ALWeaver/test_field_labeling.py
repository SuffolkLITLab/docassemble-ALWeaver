# do not pre-load
"""Real PDF/DOCX bytes through labeling, drafting and Assistant candidate edits."""

from contextlib import ExitStack
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import docx
import pikepdf

from .field_labeling import (
    assess_template,
    propose_labels,
    LabelingError,
    template_fields,
)
from .field_remapping import remap_interview


def make_pdf(path, names, signature=False):
    with pikepdf.new() as pdf:
        page = pdf.add_blank_page()
        fields = []
        for index, name in enumerate(names):
            widget = pdf.make_indirect(
                pikepdf.Dictionary(
                    Type=pikepdf.Name.Annot,
                    Subtype=pikepdf.Name.Widget,
                    FT=pikepdf.Name.Sig if signature else pikepdf.Name.Tx,
                    T=name,
                    Rect=pikepdf.Array([10, 10 + index * 30, 200, 30 + index * 30]),
                    P=page.obj,
                )
            )
            fields.append(widget)
        page.Annots = pikepdf.Array(fields)
        pdf.Root.AcroForm = pikepdf.Dictionary(Fields=pikepdf.Array(fields))
        pdf.save(path)


SOURCE = """# preserve this comment
include:
  - docassemble.AssemblyLine:assembly_line.yml
---
id: first
question: First name
fields:
  - First: first_name
---
id: last
question: Last name
fields:
  - Last: last_name
---
id: review
question: Review ${ first_name } ${ last_name }
review:
  - Edit: first_name
---
id: order
mandatory: True
code: |
  first_name
  last_name
  download
---
id: attachment
attachment:
  variable name: form_attachment
  pdf template file: form.pdf
  fields:
    "First Name": ${ first_name }
    "Last Name": ${ last_name }
"""


class FieldLabelingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)

    def test_assessment_naming_signature_and_unlabeled(self):
        path = self.directory / "form.pdf"
        make_pdf(path, ["First Name", "person[0].name", "widget4", "user_signature"])
        issues = assess_template(str(path))["issues"]
        codes = {issue["code"] for issue in issues}
        self.assertTrue(
            {
                "spaces",
                "mixed_case",
                "pdf_object_path",
                "unrecognized_noun",
                "signature_type",
            }
            <= codes
        )
        make_pdf(path, ["user_signature"], signature=True)
        self.assertNotIn(
            "signature_type", {i["code"] for i in assess_template(str(path))["issues"]}
        )
        make_pdf(path, [])
        self.assertEqual(assess_template(str(path))["issues"][0]["code"], "unlabeled")
        word = self.directory / "form.docx"
        document = docx.Document()
        document.add_paragraph("{{ users[0].name.first }}")
        document.save(word)
        self.assertFalse(assess_template(str(word))["issues"])

    def test_pdf_remap_preserves_source_and_gathers_people(self):
        updated = remap_interview(
            SOURCE,
            "main.yml",
            "form.pdf",
            "labeled.pdf",
            [
                {"old_name": "First Name", "new_name": "user_name_first"},
                {"old_name": "Last Name", "new_name": "user_name_last"},
            ],
            [],
        )
        self.assertIn("# preserve this comment", updated)
        self.assertIn("First: users[0].name.first", updated)
        self.assertIn("Edit: users[0].name.first", updated)
        self.assertIn(
            "Review ${ users[0].name.first } ${ users[0].name.last }", updated
        )
        self.assertIn("pdf template file: labeled.pdf", updated)
        self.assertIn('"user_name_first": ${ users[0].name.first }', updated)
        self.assertEqual(updated.count("users.gather()"), 1)
        self.assertNotIn("users: ALPeopleList", updated)
        self.assertNotIn("first_name", updated)

    def test_computed_mapping_rejected(self):
        with self.assertRaises(LabelingError):
            remap_interview(
                SOURCE.replace("${ first_name }", "${ first_name.upper() }"),
                "main.yml",
                "form.pdf",
                "labeled.pdf",
                [{"old_name": "First Name", "new_name": "user_name_first"}],
                [],
            )

    def test_people_preserve_baseline_custom_classes_and_idempotency(self):
        from .field_remapping import consolidate_people

        source = SOURCE.replace("first_name", "users[0].name.first")
        once = consolidate_people(source, "main.yml", ["users[0].name.first"])
        self.assertEqual(
            once, consolidate_people(once, "main.yml", ["users[0].name.first"])
        )
        self.assertNotIn("target_number", once)
        custom = "objects:\n  - people: ALChildList\n---\nmandatory: True\ncode: |\n  people[0].email\n"
        updated = consolidate_people(custom, "main.yml", ["people[0].email"])
        self.assertIn("people: ALChildList", updated)
        self.assertIn("people.gather()", updated)
        commented = "# objects:\n#   - people: ALPeopleList\n---\nmandatory: True\ncode: |\n  people[0].name.first\n"
        updated = consolidate_people(commented, "main.yml", ["people[0].name.first"])
        self.assertIn("\n  - people: ALPeopleList.using(ask_number=True)", updated)

    def test_pdf_display_expression_is_not_replaced_by_question_trigger(self):
        updated = remap_interview(
            SOURCE,
            "main.yml",
            "form.pdf",
            "new.pdf",
            [{"old_name": "First Name", "new_name": "user_name_full"}],
            [],
        )
        self.assertIn('"user_name_full": ${ users[0] }', updated)
        self.assertIn("First: users[0].name.first", updated)

    def test_unmapped_pdf_widgets_are_left_unmapped(self):
        updated = remap_interview(
            SOURCE,
            "main.yml",
            "form.pdf",
            "new.pdf",
            [
                {"old_name": "First Name", "new_name": "user_name_first"},
                {"old_name": "Court Only", "new_name": "court_only"},
            ],
            [],
        )
        self.assertNotIn("court_only:", updated)

    def test_pdf_added_fields_keep_comments_and_block_scalar_boundaries(self):
        import yaml

        for last in ["${ last_name }  # retained", "|\n      ${ last_name }"]:
            source = (
                SOURCE.replace('"Last Name": ${ last_name }', '"Last Name": ' + last)
                + "---\nquestion: Next\n"
            )
            updated = remap_interview(
                source, "main.yml", "form.pdf", "new.pdf", [], ["favorite_color"]
            )
            documents = list(yaml.safe_load_all(updated))
            self.assertEqual(
                documents[-2]["attachment"]["fields"]["favorite_color"],
                "${ favorite_color }",
            )
            self.assertEqual(documents[-1]["question"], "Next")
            if "# retained" in source:
                self.assertIn("${ last_name }  # retained", updated)

    def test_docx_display_methods_are_valid_and_remap_to_assignable_answers(self):
        from docassemble.ALDashboard import docx_wrangling as dashboard

        original, labeled = self.directory / "methods.docx", self.directory / "new.docx"
        word = docx.Document()
        word.add_paragraph("{{ FirstName }}")
        word.save(original)
        with patch.object(
            dashboard,
            "get_labeled_docx_runs",
            return_value=[(0, 0, "{{ users[0].name.full() }}", 0)],
        ):
            proposal = propose_labels(str(original), str(labeled))
        self.assertFalse(
            any(
                issue["code"] == "invalid_reference"
                for issue in proposal["assessment"]["issues"]
            )
        )
        source = (
            SOURCE.replace("first_name", "FirstName").split("  pdf template file:")[0]
            + "  docx template file: methods.docx\n"
        )
        updated = remap_interview(
            source, "main.yml", "methods.docx", "new.docx", proposal["mappings"], []
        )
        self.assertIn("First: users[0].name.first", updated)

    def test_pdf_repeated_appearance_suffix_is_supported(self):
        from docassemble.ALDashboard import pdf_field_labeler as dashboard

        original, output = self.directory / "repeat.pdf", self.directory / "new.pdf"
        make_pdf(original, ["Name One", "Name Two"])
        with patch.object(
            dashboard,
            "_generate_ai_relabel_mapping_for_unique_fields",
            return_value={
                "Name One": "user_name_first",
                "Name Two": "user_name_first__2",
            },
        ):
            proposal = propose_labels(str(original), str(output))
        self.assertEqual(len(proposal["mappings"]), 2)

    def test_real_pdf_labeler_and_generator(self):
        from docassemble.ALDashboard import pdf_field_labeler as dashboard
        from .interview_generator import generate_interview_from_path

        original = self.directory / "original.pdf"
        labeled = self.directory / "labeled.pdf"
        make_pdf(original, ["First Name", "Last Name"])
        original_bytes = original.read_bytes()
        # Deterministic AI answer; the actual Dashboard PDF writer and Weaver
        # generator run on real files, not stubbed generation results.
        with patch.object(
            dashboard,
            "_generate_ai_relabel_mapping_for_unique_fields",
            return_value={
                "First Name": "user_name_first",
                "Last Name": "user_name_last",
            },
        ):
            proposal = propose_labels(str(original), str(labeled))
        self.assertEqual(len(proposal["mappings"]), 2)
        result = generate_interview_from_path(
            str(labeled),
            output_dir=str(self.directory / "draft"),
            create_package_zip=False,
            include_next_steps=False,
            use_llm_assist=False,
        )
        self.assertIn("users[0].name.first", result.yaml_text)
        self.assertEqual(original.read_bytes(), original_bytes)
        self.assertEqual(
            [f["name"] for f in template_fields(str(labeled))],
            ["user_name_first", "user_name_last"],
        )

    def test_real_docx_labeler_tracks_body_and_table_fields(self):
        from docassemble.ALDashboard import docx_wrangling as dashboard

        original = self.directory / "original.docx"
        labeled = self.directory / "labeled.docx"
        document = docx.Document()
        document.add_paragraph("{{ FirstName }}")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "{{ LastName }}"
        document.save(original)
        original_bytes = original.read_bytes()
        with patch.object(
            dashboard,
            "get_labeled_docx_runs",
            return_value=[
                (0, 0, "{{ users[0].name.first }}", 0),
                (1, 0, "{{ users[0].name.last }}", 0),
            ],
        ):
            proposal = propose_labels(str(original), str(labeled))
        self.assertEqual(
            proposal["mappings"],
            [
                {"old_name": "FirstName", "new_name": "users[0].name.first"},
                {"old_name": "LastName", "new_name": "users[0].name.last"},
            ],
        )
        self.assertEqual(original.read_bytes(), original_bytes)

    def test_missing_dashboard_is_actionable(self):
        import sys

        original = self.directory / "original.pdf"
        make_pdf(original, ["First Name"])
        with patch.dict(
            sys.modules, {"docassemble.ALDashboard.pdf_field_labeler": None}
        ):
            with self.assertRaisesRegex(LabelingError, "Install or update ALDashboard"):
                propose_labels(str(original), str(self.directory / "output.pdf"))

    def test_generated_sequence_pdf_mappings_and_review_strings(self):
        source = SOURCE.replace('    "First Name":', '    - "First Name":').replace(
            '    "Last Name":', '    - "Last Name":'
        )
        source = source.replace(
            "Review ${ first_name }", "Review ${ showifdef('first_name') }"
        )
        updated = remap_interview(
            source,
            "main.yml",
            "form.pdf",
            "new.pdf",
            [
                {"old_name": "First Name", "new_name": "user_name_first"},
                {"old_name": "Last Name", "new_name": "user_name_last"},
            ],
            ["favorite_color"],
        )
        import yaml

        attachment = list(yaml.safe_load_all(updated))[5]["attachment"]
        self.assertEqual(
            attachment["fields"][-1], {"favorite_color": "${ favorite_color }"}
        )
        self.assertIn("showifdef('users[0].name.first')", updated)
        self.assertIn("users.gather()", updated)

    def test_signature_type_warning_reaches_draft_result(self):
        from .interview_generator import generate_interview_from_path

        path = self.directory / "signature.pdf"
        make_pdf(path, ["user_signature"])
        result = generate_interview_from_path(
            str(path),
            output_dir=str(self.directory / "draft"),
            create_package_zip=False,
            include_next_steps=False,
            use_llm_assist=False,
        )
        self.assertTrue(
            any("not a PDF signature field" in warning for warning in result.warnings)
        )

    def test_unlabeled_word_proposal_generates_questions(self):
        from docassemble.ALDashboard import docx_wrangling as dashboard
        from .interview_generator import generate_interview_from_path

        path = self.directory / "blank.docx"
        labeled = self.directory / "labeled.docx"
        document = docx.Document()
        document.add_paragraph("Favorite color: ______")
        document.save(path)
        self.assertEqual(assess_template(str(path))["issues"][0]["code"], "unlabeled")
        with patch.object(
            dashboard,
            "get_labeled_docx_runs",
            return_value=[(0, 0, "Favorite color: {{ favorite_color }}", 0)],
        ):
            proposal = propose_labels(str(path), str(labeled))
        self.assertEqual(proposal["added_fields"], ["favorite_color"])
        result = generate_interview_from_path(
            str(labeled),
            output_dir=str(self.directory / "draft"),
            create_package_zip=False,
            include_next_steps=False,
            use_llm_assist=False,
        )
        self.assertIn("favorite_color", result.yaml_text)

    def test_relabel_an_actual_generated_interview_without_stale_scalar_names(self):
        from .interview_generator import generate_interview_from_path

        path = self.directory / "form.pdf"
        make_pdf(path, ["First Name", "Last Name"])
        result = generate_interview_from_path(
            str(path),
            output_dir=str(self.directory / "draft"),
            create_package_zip=False,
            include_next_steps=False,
            use_llm_assist=False,
        )
        updated = remap_interview(
            result.yaml_text,
            "main.yml",
            "form.pdf",
            "labeled.pdf",
            [
                {"old_name": "First Name", "new_name": "user_name_first"},
                {"old_name": "Last Name", "new_name": "user_name_last"},
            ],
            [],
        )
        self.assertNotIn("First_Name", updated)
        self.assertNotIn("Last_Name", updated)
        self.assertIn("users.gather()", updated)
        self.assertIn("showifdef('users[0].name.first')", updated)

    def test_split_word_runs_keep_field_identity(self):
        from docassemble.ALDashboard import docx_wrangling as dashboard

        path = self.directory / "split.docx"
        labeled = self.directory / "labeled.docx"
        document = docx.Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("{{ First")
        paragraph.add_run("Name }}")
        document.save(path)
        with patch.object(
            dashboard,
            "get_labeled_docx_runs",
            return_value=[(0, 0, "{{ users[0].name.first }}", 0)],
        ):
            proposal = propose_labels(str(path), str(labeled))
        self.assertEqual(
            proposal["mappings"],
            [{"old_name": "FirstName", "new_name": "users[0].name.first"}],
        )

    def test_partial_word_rename_is_rejected(self):
        from docassemble.ALDashboard import docx_wrangling as dashboard

        path = self.directory / "repeat.docx"
        document = docx.Document()
        document.add_paragraph("{{ FirstName }}")
        document.add_paragraph("{{ FirstName }}")
        document.save(path)
        with patch.object(
            dashboard,
            "get_labeled_docx_runs",
            return_value=[(0, 0, "{{ users[0].name.first }}", 0)],
        ):
            with self.assertRaisesRegex(LabelingError, "still occur"):
                propose_labels(str(path), str(self.directory / "output.docx"))

    def test_assistant_context_has_exact_template_filenames(self):
        from .editor_agent_context import build_agent_context

        context = build_agent_context(filename="main.yml", raw_source=SOURCE)
        self.assertEqual(context["template_filenames"], ["form.pdf"])
