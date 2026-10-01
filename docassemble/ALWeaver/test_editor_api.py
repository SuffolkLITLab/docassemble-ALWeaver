# do not pre-load

from io import BytesIO
from contextlib import ExitStack, nullcontext
import hashlib
from pathlib import Path
import os
import importlib
import importlib.util
import sys
import tempfile
import threading
import types
import unittest
from types import SimpleNamespace
from unittest.mock import patch

from flask import Flask, jsonify
from werkzeug.datastructures import FileStorage
from .project_filenames import safe_project_filename


class _TestRedisLock:
    _locks: dict[str, threading.RLock] = {}
    _guard = threading.Lock()

    def __init__(self, name):
        with self._guard:
            self._lock = self._locks.setdefault(name, threading.RLock())

    def acquire(self, blocking=True, blocking_timeout=None):
        if blocking_timeout is None:
            return self._lock.acquire(blocking=blocking)
        return self._lock.acquire(blocking=blocking, timeout=blocking_timeout)

    def release(self):
        self._lock.release()


class _TestRedis:
    def lock(self, name, **_kwargs):
        return _TestRedisLock(name)


def _load_api_editor_for_tests():
    module_path = Path(__file__).with_name("api_editor.py")
    package_name = __package__ or "docassemble.ALWeaver"
    app = Flask("alweaver-api-editor-tests")

    class _CSRF:
        def exempt(self, fn):
            return fn

    current_user = types.SimpleNamespace(is_authenticated=False, id=None)

    app_object = types.ModuleType("docassemble.webapp.app_object")
    app_object.app = app
    app_object.csrf = _CSRF()

    server_mod = types.ModuleType("docassemble.webapp.server")

    def jsonify_with_status(payload, status):
        response = jsonify(payload)
        response.status_code = status
        return response

    server_mod.jsonify_with_status = jsonify_with_status
    server_mod.r = _TestRedis()

    worker_common = types.ModuleType("docassemble.webapp.worker_common")
    worker_common.bg_context = nullcontext
    worker_common.workerapp = types.SimpleNamespace(
        send_task=lambda *args, **kwargs: None,
        AsyncResult=lambda *args, **kwargs: None,
    )

    flask_cors = types.ModuleType("flask_cors")
    flask_cors.cross_origin = lambda *args, **kwargs: (lambda fn: fn)

    flask_login = types.ModuleType("flask_login")
    flask_login.current_user = current_user

    base_util = types.ModuleType("docassemble.base.util")
    base_util.log = lambda *args, **kwargs: None

    api_utils = types.ModuleType(f"{package_name}.api_utils")
    api_utils.generate_interview_from_bytes = lambda *args, **kwargs: {}
    api_utils.parse_bool = lambda value, default=False: (
        default
        if value is None
        else str(value).strip().lower() in {"1", "true", "yes", "on"}
    )
    api_utils.validate_upload_metadata = lambda **kwargs: (kwargs["filename"], ".docx")
    api_utils.validate_document_content = lambda filename, content_bytes: None

    editor_utils = types.ModuleType(f"{package_name}.editor_utils")
    for name, func in {
        "canonical_block_yaml": lambda block: "id: block\n",
        "canonicalize_block_yaml": lambda yaml_text: yaml_text.strip(),
        "comment_out_block_in_yaml": lambda content, block_id: content,
        "delete_block_from_yaml": lambda content, block_id: content,
        "delete_saved_file": lambda *args, **kwargs: None,
        "generate_draft_order": lambda *args, **kwargs: {},
        "add_object_declaration": lambda content, block_id, name, expression: content,
        "insert_block_in_yaml": lambda content, block_yaml, insert_after_id=None: content,
        "inserted_block_id_by_position": lambda blocks, insert_after_id: None,
        "is_comment_only_yaml": lambda text: bool(
            [line for line in text.splitlines() if line.strip()]
        )
        and all(
            line.lstrip().startswith("#") for line in text.splitlines() if line.strip()
        ),
        "parse_interview_yaml": lambda *args, **kwargs: {
            "blocks": [],
            "metadata_blocks": [],
        },
        "_safe_load_interview_document": lambda raw: __import__("yaml").safe_load(raw),
        "metadata_source_slice": lambda *args, **kwargs: "",
        "parse_order_code": lambda *args, **kwargs: {},
        "playground_get_variables": lambda *args, **kwargs: {},
        "playground_interview_url": lambda *args, **kwargs: "/interview",
        "playground_list_projects": lambda *args, **kwargs: [],
        "playground_list_yaml_files": lambda *args, **kwargs: [],
        "playground_read_yaml": lambda *args, **kwargs: "",
        "playground_write_yaml": lambda *args, **kwargs: None,
        "rename_saved_file": lambda *args, **kwargs: None,
        "serialize_blocks_to_yaml": lambda *args, **kwargs: "",
        "serialize_order_steps": lambda *args, **kwargs: "",
        "validate_order_steps": lambda *args, **kwargs: None,
        "source_revision": lambda text: "test-revision",
        "enable_commented_block_in_yaml": lambda content, block_id: content,
        "reorder_blocks_in_yaml": lambda content, order: content,
        "update_block_in_yaml": lambda content, block_id, new_yaml, **kwargs: content,
        "update_metadata_documents_in_yaml": lambda content, edited: content,
    }.items():
        setattr(editor_utils, name, func)
    # `document_bundles` and `template_analysis` are imported for real by
    # `api_editor`, and they read these from `editor_utils`.
    for name, value in {
        "BLOCK_TYPE_ATTACHMENT": "attachment",
        "BLOCK_TYPE_CODE": "code",
        "BLOCK_TYPE_OBJECTS": "objects",
        "BLOCK_TYPE_QUESTION": "question",
        "BLOCK_TYPE_TEMPLATE": "template",
        "_split_top_level_commas": lambda text: [part for part in str(text).split(",")],
    }.items():
        setattr(editor_utils, name, value)

    editor_ai_utils = types.ModuleType(f"{package_name}.editor_ai_utils")
    editor_ai_utils.DEFAULT_FIELD_TYPES = []
    editor_ai_utils.normalize_generated_fields = lambda *args, **kwargs: []
    editor_ai_utils.normalize_generated_screen = lambda *args, **kwargs: {}
    editor_ai_utils.pick_small_model_name = lambda *args, **kwargs: "gpt-5-nano"
    editor_ai_utils.validate_yaml_with_dayamlchecker = lambda *args, **kwargs: (
        True,
        "",
    )

    playground_publish = types.ModuleType(f"{package_name}.playground_publish")
    playground_publish.SECTION_TO_STORAGE = {
        "templates": "templates",
        "modules": "modules",
        "static": "static",
        "sources": "sources",
    }
    playground_publish._copy_files_to_section = lambda *args, **kwargs: None
    playground_publish.delete_project = lambda *args, **kwargs: None
    playground_publish.create_project = lambda *args, **kwargs: None
    playground_publish.get_list_of_projects = lambda *args, **kwargs: []
    playground_publish.find_project_github_sync = lambda *args, **kwargs: None
    playground_publish.import_github_snapshot = lambda *args, **kwargs: {}
    playground_publish.merge_github_snapshot = lambda *args, **kwargs: {}
    playground_publish.next_available_project_name = (
        lambda base_name, existing=None: base_name
    )
    playground_publish.normalize_github_package_name = lambda raw_name: str(
        raw_name
    ).strip()
    playground_publish.github_project_name = lambda repository, branch="": (
        "GitHubProject"
    )
    playground_publish.normalize_project_name = lambda raw_name, **kwargs: str(
        raw_name
    ).strip()
    playground_publish.prepare_project_github_package = lambda **kwargs: {
        "package": kwargs["package_name"],
        "repository": "docassemble-" + kwargs["package_name"],
    }
    playground_publish.load_project_github_manifest = lambda **kwargs: ({}, "")
    playground_publish.record_project_github_sync = lambda *args, **kwargs: None
    playground_publish.rename_project = lambda *args, **kwargs: None

    stubs = {
        "docassemble.base.util": base_util,
        "docassemble.webapp.app_object": app_object,
        "docassemble.webapp.server": server_mod,
        "docassemble.webapp.worker_common": worker_common,
        "flask_cors": flask_cors,
        "flask_login": flask_login,
        f"{package_name}.api_utils": api_utils,
        f"{package_name}.editor_utils": editor_utils,
        f"{package_name}.editor_ai_utils": editor_ai_utils,
        f"{package_name}.playground_publish": playground_publish,
    }
    # `api_editor` imports this one for real. Import it before the stubs go in,
    # so it binds to the real `editor_utils` and the document endpoints can be
    # tested against actual YAML instead of a stub that returns nothing.
    importlib.import_module(f"{package_name}.document_bundles")
    previous = {name: sys.modules.get(name) for name in stubs}
    module_name = f"{package_name}._test_api_editor"
    try:
        sys.modules.update(stubs)
        spec = importlib.util.spec_from_file_location(module_name, module_path)
        module = importlib.util.module_from_spec(spec)
        sys.modules[module_name] = module
        if spec is None or spec.loader is None:
            raise RuntimeError("Unable to load api_editor test module")
        spec.loader.exec_module(module)
        return module
    finally:
        for name, original in previous.items():
            if original is None:
                sys.modules.pop(name, None)
            else:
                sys.modules[name] = original


api_editor = _load_api_editor_for_tests()


class TestEditorNavigationRoutes(unittest.TestCase):
    def test_deep_editor_pages_serve_the_shell_and_keep_auth_guard(self):
        paths = [
            "/al/editor/",
            "/al/editor/projects",
            "/al/editor/projects/",
            "/al/editor/projects/FaxCoverSheet",
            "/al/editor/projects/FaxCoverSheet/interviews",
            "/al/editor/projects/FaxCoverSheet/interviews/",
            "/al/editor/projects/FaxCoverSheet/interviews/filename.yml",
            "/al/editor/projects/FaxCoverSheet/interviews/filename.yml/blocks",
            "/al/editor/projects/FaxCoverSheet/interviews/filename.yml/blocks/",
            "/al/editor/projects/FaxCoverSheet/interviews/filename.yml/blocks/block-id",
            "/al/editor/projects/FaxCoverSheet/interviews/filename.yml/blocks/block-id/",
            "/al/editor/projects/FaxCoverSheet/interviews/filename.yml/blocks/~intake%252Fperson",
            "/al/editor/projects/FaxCoverSheet/interviews/filename.yml/blocks/~intake%255Cperson",
            "/al/editor/projects/FaxCoverSheet/interviews/filename.yml/blocks/~..",
            "/al/editor/projects/FaxCoverSheet/interviews/filename.yml/blocks/~~",
            "/al/editor/projects/FaxCoverSheet/interviews/filename.yml/source",
            "/al/editor/projects/FaxCoverSheet/interviews/filename.yml/order",
            "/al/editor/projects/FaxCoverSheet/interviews/filename.yml/settings",
            "/al/editor/projects/FaxCoverSheet/interviews/filename.yml/tests",
            "/al/editor/projects/FaxCoverSheet/interviews/filename.yml/debug",
            "/al/editor/projects/FaxCoverSheet/interviews/filename.yml/documents",
            "/al/editor/projects/FaxCoverSheet/templates",
            "/al/editor/projects/FaxCoverSheet/templates/form.docx",
            "/al/editor/projects/FaxCoverSheet/modules",
            "/al/editor/projects/FaxCoverSheet/modules/main.py",
            "/al/editor/projects/FaxCoverSheet/static",
            "/al/editor/projects/FaxCoverSheet/static/app.js",
            "/al/editor/projects/FaxCoverSheet/sources",
            "/al/editor/projects/FaxCoverSheet/sources/data.csv",
            "/al/editor/projects/FaxCoverSheet/documents",
            "/al/editor/create",
        ]
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(
                api_editor, "_render_editor_page", return_value="<main>SPA</main>"
            ),
        ):
            client = api_editor.app.test_client()
            for path in paths:
                with self.subTest(path=path):
                    response = client.get(path)
                    self.assertEqual(response.status_code, 200)
                    self.assertIn(b"<main>SPA</main>", response.data)

        with (
            patch.object(api_editor, "_editor_auth_check", return_value=False),
            patch.object(
                api_editor,
                "_editor_auth_urls",
                return_value=("/user/sign-in?next=%2Fal%2Feditor", "/user/sign-out"),
            ),
        ):
            response = api_editor.app.test_client().get(
                "/al/editor/projects/FaxCoverSheet/interviews/filename.yml/blocks/block-id"
            )
        self.assertEqual(response.status_code, 302)
        self.assertIn("/user/sign-in", response.headers["Location"])

    def test_api_and_static_routes_are_not_captured_by_page_routes(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=1),
            patch.object(api_editor, "playground_list_projects", return_value=[]),
            patch.object(api_editor, "_project_github_sync_summaries", return_value={}),
            patch.object(api_editor, "_get_static_content", return_value="body{}"),
        ):
            client = api_editor.app.test_client()
            api_response = client.get("/al/editor/api/projects")
            static_response = client.get("/al/editor/static/editor.css")
        self.assertEqual(api_response.status_code, 200)
        self.assertTrue(api_response.is_json)
        self.assertEqual(static_response.status_code, 200)
        self.assertEqual(static_response.mimetype, "text/css")
        self.assertEqual(static_response.get_data(as_text=True), "body{}")

    def test_unknown_nested_interview_paths_do_not_fall_through_to_spa(self):
        response = api_editor.app.test_client().get(
            "/al/editor/projects/FaxCoverSheet/interviews/filename.yml/unknown"
        )
        self.assertEqual(response.status_code, 404)


class _FakeRedis:
    """Just enough Redis for the editor's job-state records."""

    def __init__(self):
        self.store = {}

    def get(self, key):
        return self.store.get(key)

    def pipeline(self):
        store = self.store

        class _Pipe:
            def __init__(self):
                self.pending = []

            def set(self, key, value):
                self.pending.append((key, value))
                return self

            def expire(self, key, seconds):
                return self

            def execute(self):
                for key, value in self.pending:
                    store[key] = value
                self.pending = []

        return _Pipe()


class TestEditorGithubApi(unittest.TestCase):
    def test_github_branches_lists_selected_repository(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=42),
            patch.object(
                api_editor,
                "get_native_github_integration",
                return_value={"enabled": True, "connected": True},
            ),
            patch.object(
                api_editor,
                "get_github_repository_branches",
                return_value={
                    "repository_exists": True,
                    "default_branch": "main",
                    "branches": ["main", "feature/housing"],
                },
            ) as branches,
        ):
            response = api_editor.app.test_client().get(
                "/al/editor/api/github/branches",
                query_string={
                    "project": "Housing",
                    "owner": "LegalAid",
                    "package": "HousingForms",
                },
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.get_json()["data"]["branches"], ["main", "feature/housing"]
        )
        branches.assert_called_once_with(
            owner="LegalAid", repository="docassemble-HousingForms", user_id=42
        )

    def test_github_branches_lists_any_repository_by_url_without_a_connection(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=42),
            patch.object(api_editor, "get_native_github_integration") as integration,
            patch.object(
                api_editor,
                "get_github_repository_branches",
                return_value={
                    "repository_exists": True,
                    "default_branch": "main",
                    "branches": ["main", "draft"],
                },
            ) as branches,
        ):
            response = api_editor.app.test_client().get(
                "/al/editor/api/github/branches",
                query_string={
                    "repository_url": "https://github.com/LegalAid/docassemble-Forms.git"
                },
            )
        self.assertEqual(response.status_code, 200, response.get_json())
        data = response.get_json()["data"]
        self.assertEqual(data["branches"], ["main", "draft"])
        self.assertEqual(
            data["repository_url"], "https://github.com/LegalAid/docassemble-Forms"
        )
        branches.assert_called_once_with(
            owner="LegalAid",
            repository="docassemble-Forms",
            user_id=42,
            allow_anonymous=True,
        )
        integration.assert_not_called()

    def test_github_branches_requires_connected_account(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=42),
            patch.object(
                api_editor,
                "get_native_github_integration",
                return_value={"enabled": True, "connected": False},
            ),
            patch.object(api_editor, "get_github_repository_branches") as branches,
        ):
            response = api_editor.app.test_client().get(
                "/al/editor/api/github/branches",
                query_string={
                    "project": "Housing",
                    "owner": "LegalAid",
                    "package": "HousingForms",
                },
            )
        self.assertEqual(response.status_code, 409)
        branches.assert_not_called()

    def test_github_publish_preview_shows_target_and_text_diff_without_persisting_manifest(
        self,
    ):
        prepared = {
            "manifest": {"interview_files": ["main.yml"]},
            "manifest_path": "",
        }
        local_files = {
            "docassemble/forms/data/questions/main.yml": {
                "content": b"question: Updated\n",
                "mode": "100644",
            }
        }
        remote = {
            "sha": "remote-head",
            "files": {
                "docassemble/forms/data/questions/main.yml": b"question: Old\n",
                "remote-only.txt": b"remove me\n",
            },
            "missing": False,
        }
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "get_native_github_integration",
                return_value={"enabled": True, "connected": True},
            ),
            patch.object(
                api_editor,
                "get_github_publish_owners",
                return_value=[{"login": "ada", "type": "user"}],
            ),
            patch.object(api_editor, "_editor_user_designator", return_value="Ada"),
            patch.object(
                api_editor,
                "prepare_project_github_package",
                return_value=prepared,
            ) as prepare,
            patch.object(api_editor, "repository_dependency_names", return_value=[]),
            patch.object(
                api_editor,
                "repository_publish_files",
                return_value={"files": {}, "managed_paths": []},
            ),
            patch.object(
                api_editor, "build_github_package_snapshot", return_value=local_files
            ),
            patch.object(
                api_editor, "get_github_repository_snapshot", return_value=remote
            ) as get_remote,
            patch.object(api_editor, "find_project_github_sync", return_value=None),
            patch.object(
                api_editor,
                "_sign_github_publish_preview",
                return_value="signed-preview",
            ),
        ):
            api_editor.current_user.email = "ada@example.com"
            with api_editor.app.test_request_context(
                "/al/editor/api/github/publish/preview",
                method="POST",
                json={
                    "project": "Housing",
                    "owner": "ada",
                    "package": "forms",
                    "branch": "main",
                },
            ):
                response = api_editor.editor_api_github_publish_preview()

        self.assertEqual(response.status_code, 200)
        data = response.get_json()["data"]
        self.assertEqual(data["remote_sha"], "remote-head")
        self.assertEqual(data["preview_token"], "signed-preview")
        self.assertEqual(data["files"], ["docassemble/forms/data/questions/main.yml"])
        by_path = {entry["path"]: entry for entry in data["changes"]}
        self.assertEqual(
            by_path["docassemble/forms/data/questions/main.yml"]["change"],
            "modified",
        )
        self.assertIn(
            "-question: Old",
            by_path["docassemble/forms/data/questions/main.yml"]["diff"],
        )
        self.assertEqual(by_path["remote-only.txt"]["change"], "deleted")
        prepare.assert_called_once()
        self.assertFalse(prepare.call_args.kwargs["persist_manifest"])
        get_remote.assert_called_once_with(
            repository_url="https://github.com/ada/docassemble-forms",
            user_id=7,
            ref="main",
            allow_missing=True,
            include_all_files=True,
        )

    def test_github_publish_preview_accepts_empty_repository(self):
        missing = {
            "missing": True,
            "repository_exists": True,
            "default_branch": "main",
            "sha": "",
            "files": {},
        }
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "get_native_github_integration",
                return_value={"enabled": True, "connected": True},
            ),
            patch.object(
                api_editor,
                "get_github_publish_owners",
                return_value=[{"login": "ada", "type": "user"}],
            ),
            patch.object(
                api_editor,
                "prepare_project_github_package",
                return_value={"manifest": {}, "manifest_path": ""},
            ),
            patch.object(api_editor, "repository_dependency_names", return_value=[]),
            patch.object(
                api_editor,
                "repository_publish_files",
                return_value={"files": {}, "managed_paths": []},
            ),
            patch.object(
                api_editor,
                "build_github_package_snapshot",
                return_value={"README.md": {"content": b"new", "mode": "100644"}},
            ),
            patch.object(
                api_editor,
                "get_github_repository_snapshot",
                side_effect=[missing, missing],
            ) as get_remote,
            patch.object(api_editor, "find_project_github_sync", return_value=None),
            patch.object(
                api_editor, "_sign_github_publish_preview", return_value="token"
            ),
        ):
            api_editor.current_user.email = "ada@example.com"
            with api_editor.app.test_request_context(
                "/al/editor/api/github/publish/preview",
                method="POST",
                json={
                    "project": "Housing",
                    "owner": "ada",
                    "package": "forms",
                    "branch": "main",
                },
            ):
                response = api_editor.editor_api_github_publish_preview()
        self.assertEqual(response.status_code, 200)
        preview = response.get_json()["data"]
        self.assertEqual(preview["remote_sha"], "")
        self.assertTrue(preview["repository_missing"])
        self.assertEqual(preview["preview_token"], "token")
        self.assertEqual(get_remote.call_count, 2)
        self.assertTrue(get_remote.call_args.kwargs["allow_missing"])

    def test_github_publish_can_recreate_deleted_target_branch(self):
        sync = {
            "commit": "former-branch-sha",
            "package": "forms",
            "repository_url": "https://github.com/ada/docassemble-forms",
            "branch": "draft",
        }
        with (
            patch.object(api_editor, "find_project_github_sync", return_value=sync),
            patch.object(
                api_editor,
                "get_github_branch_head",
                return_value={"missing": True, "repository_exists": True, "sha": ""},
            ) as get_remote,
        ):
            api_editor._assert_github_publish_branch_is_current(
                uid=7,
                project="Housing",
                repository_url="https://github.com/ada/docassemble-forms",
                branch="draft",
            )
        get_remote.assert_called_once_with(
            repository_url="https://github.com/ada/docassemble-forms",
            ref="draft",
            user_id=7,
        )

    def test_publish_after_preview_checks_the_reviewed_head_once(self):
        url = "https://github.com/ada/docassemble-forms"
        sync = {"repository_url": url, "branch": "main", "commit": "synced-sha"}
        files = {"docassemble/forms/data/questions/main.yml": b""}
        new_branch = {
            "missing": True,
            "repository_exists": True,
            "default_branch": "main",
            "sha": "",
        }
        cases = [
            ("unchanged", [{"sha": "synced-sha", "files": files}], "synced-sha", None),
            (
                "moved",
                [{"sha": "synced-sha", "files": files}],
                "previewed-sha",
                "changed after the publish preview",
            ),
            (
                "new branch from the default",
                [new_branch, {"sha": "main-sha", "files": files}],
                "main-sha",
                None,
            ),
        ]
        for label, heads, expected_sha, error in cases:
            with self.subTest(label):
                with (
                    patch.object(
                        api_editor, "find_project_github_sync", return_value=sync
                    ),
                    patch.object(
                        api_editor, "get_github_branch_head", side_effect=heads
                    ) as head,
                ):
                    check = lambda: api_editor._assert_github_publish_branch_is_current(
                        uid=7,
                        project="Housing",
                        repository_url=url,
                        branch="draft" if len(heads) > 1 else "main",
                        expected_remote_sha=expected_sha,
                    )
                    if error:
                        with self.assertRaisesRegex(ValueError, error):
                            check()
                    else:
                        check()
                self.assertEqual(head.call_count, len(heads))

    def test_publishing_to_an_existing_branch_requires_its_head_to_be_pulled(self):
        url = "https://github.com/ada/docassemble-forms"
        sync = {"repository_url": url, "branch": "main", "commit": "synced-sha"}
        other_repo = {**sync, "repository_url": "https://github.com/ada/other"}
        files = {"docassemble/forms/data/questions/main.yml": b"---\n"}
        head = {"sha": "synced-sha", "files": files}
        newer = {"sha": "newer-sha", "files": files}
        readme_only = {"sha": "init-sha", "files": {"README.md": b"# forms\n"}}
        cases = [
            ("new branch", sync, "draft", {"missing": True, "sha": ""}, None),
            ("new repository", None, "main", {"missing": True, "sha": ""}, None),
            ("initialized with a README", None, "main", readme_only, None),
            ("synced head", sync, "main", head, None),
            ("other branch at the synced commit", sync, "draft", head, None),
            ("synced branch advanced", sync, "main", newer, "has advanced"),
            ("never synced", None, "main", head, "Create a project"),
            (
                "synced to another repository",
                other_repo,
                "main",
                head,
                "Create a project",
            ),
            ("other branch elsewhere", sync, "draft", newer, "Pull from GitHub"),
            (
                "no recorded commit",
                {**sync, "commit": ""},
                "main",
                head,
                "has advanced",
            ),
        ]
        for label, project_sync, branch, remote, expected in cases:
            with self.subTest(label):
                message = api_editor._github_branch_unpulled_message(
                    project_sync, repository_url=url, branch=branch, remote=remote
                )
                if expected is None:
                    self.assertIsNone(message)
                else:
                    self.assertIn(expected, message)

    def test_github_authorization_requires_editor_access(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=False),
            patch.object(api_editor, "github_authorization_url") as authorize,
        ):
            with api_editor.app.test_request_context("/al/editor/github/authorize"):
                response = api_editor.editor_github_authorize()
        self.assertEqual(response.status_code, 401)
        authorize.assert_not_called()

    def test_github_authorization_redirects_to_native_oauth_flow(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(
                api_editor,
                "github_authorization_url",
                return_value="https://github.com/login/oauth/authorize?scope=repo+workflow",
            ),
        ):
            with api_editor.app.test_request_context("/al/editor/github/authorize"):
                response = api_editor.editor_github_authorize()
        self.assertEqual(response.status_code, 302)
        self.assertIn("scope=repo+workflow", response.location)

    def test_expression_parser_requires_editor_authentication(self):
        with patch.object(api_editor, "_editor_auth_check", return_value=False):
            response = api_editor.app.test_client().post(
                "/al/editor/api/expression", json={"source": "x + 1"}
            )
        self.assertIn(response.status_code, (401, 403))

    def test_expression_parser_never_executes_source(self):
        with patch.object(api_editor, "_editor_auth_check", return_value=True):
            client = api_editor.app.test_client()
            response = client.post(
                "/al/editor/api/expression",
                json={"source": '__import__("os").system("false")'},
            )
            self.assertEqual(response.status_code, 200)
            self.assertTrue(response.json["data"]["valid"])
            self.assertFalse(response.json["data"]["supported"])
            self.assertEqual(
                client.post("/al/editor/api/expression", json=["bad"]).status_code, 400
            )

    def test_projects_only_marks_projects_with_a_github_manifest_as_synced(self):
        def find_sync(*, user_id, project_name):
            if project_name != "SyncedProject":
                return None
            return {
                "package": "SyncedProject",
                "repository_url": "https://github.com/LegalAid/docassemble-SyncedProject",
                "branch": "main",
                "commit": "base-sha",
            }

        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "playground_list_projects",
                return_value=["LocalOnly", "SyncedProject"],
            ),
            patch.object(api_editor, "find_project_github_sync", side_effect=find_sync),
        ):
            with api_editor.app.test_request_context("/al/editor/api/projects"):
                response = api_editor.editor_api_projects()

        data = response.get_json()["data"]
        self.assertEqual(data["projects"], ["LocalOnly", "SyncedProject"])
        self.assertNotIn("LocalOnly", data["github_syncs"])
        self.assertEqual(
            data["github_syncs"]["SyncedProject"]["repository_url"],
            "https://github.com/LegalAid/docassemble-SyncedProject",
        )

    def test_pull_uses_recorded_commit_as_a_three_way_merge_base(self):
        sync = {
            "package": "HousingForms",
            "repository_url": "https://github.com/LegalAid/docassemble-HousingForms",
            "branch": "main",
            "commit": "base-sha",
        }
        remote = {"sha": "remote-sha", "branch": "main", "files": {}}
        base = {"sha": "base-sha", "branch": "base-sha", "files": {}}
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "find_project_github_sync", return_value=sync),
            patch.object(
                api_editor, "get_github_repository_snapshot", side_effect=[remote, base]
            ) as snapshots,
            patch.object(
                api_editor,
                "merge_github_snapshot",
                return_value={"merged": True, "files": 3, "commit": "remote-sha"},
            ) as merge,
            patch.object(api_editor, "adopt_repository_snapshot") as adopt,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/github/pull", method="POST", json={"project": "Housing"}
            ):
                response = api_editor.editor_api_github_pull()

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.get_json()["success"])
        self.assertEqual(snapshots.call_args_list[0].kwargs["ref"], "main")
        self.assertEqual(snapshots.call_args_list[1].kwargs["ref"], "base-sha")
        self.assertIs(merge.call_args.kwargs["base_snapshot"], base)
        # Only workflow and pyproject.toml edits made on GitHub since the base
        # are taken, so the base files go along.
        adopt.assert_called_once_with(
            7, "Housing", "HousingForms", remote["files"], base["files"]
        )
        self.assertIs(merge.call_args.kwargs["remote_snapshot"], remote)

    def test_pull_merges_another_branch_from_its_common_commit(self):
        sync = {
            "package": "HousingForms",
            "repository_url": "https://github.com/LegalAid/docassemble-HousingForms",
            "branch": "main",
            "commit": "synced-sha",
        }
        remote = {"sha": "draft-sha", "branch": "draft", "files": {}}
        base = {"sha": "fork-sha", "branch": "fork-sha", "files": {}}
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "find_project_github_sync", return_value=sync),
            patch.object(
                api_editor, "get_github_repository_snapshot", side_effect=[remote, base]
            ) as snapshots,
            patch.object(
                api_editor, "get_github_merge_base", return_value="fork-sha"
            ) as merge_base,
            patch.object(
                api_editor,
                "merge_github_snapshot",
                return_value={"merged": True, "files": 3, "commit": "draft-sha"},
            ) as merge,
            patch.object(api_editor, "adopt_repository_snapshot"),
            patch.object(api_editor, "_reconcile_project_modules"),
            patch.object(api_editor, "_restart_state_payload", return_value={}),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/github/pull",
                method="POST",
                json={"project": "Housing", "branch": "draft"},
            ):
                response = api_editor.editor_api_github_pull()

        self.assertTrue(response.get_json()["success"], response.get_json())
        self.assertEqual(snapshots.call_args_list[0].kwargs["ref"], "draft")
        self.assertEqual(snapshots.call_args_list[1].kwargs["ref"], "fork-sha")
        merge_base.assert_called_once_with(
            repository_url=sync["repository_url"],
            base="synced-sha",
            head="draft-sha",
            user_id=7,
        )
        self.assertIs(merge.call_args.kwargs["base_snapshot"], base)
        self.assertIs(merge.call_args.kwargs["remote_snapshot"], remote)

    def test_pull_refuses_a_branch_with_no_shared_history(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "find_project_github_sync",
                return_value={
                    "package": "HousingForms",
                    "repository_url": "https://github.com/LegalAid/docassemble-HousingForms",
                    "branch": "main",
                    "commit": "synced-sha",
                },
            ),
            patch.object(
                api_editor,
                "get_github_repository_snapshot",
                return_value={"sha": "orphan-sha", "branch": "orphan", "files": {}},
            ),
            patch.object(api_editor, "get_github_merge_base", return_value=None),
            patch.object(api_editor, "merge_github_snapshot") as merge,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/github/pull",
                method="POST",
                json={"project": "Housing", "branch": "orphan"},
            ):
                response = api_editor.editor_api_github_pull()

        self.assertEqual(response.status_code, 400)
        self.assertIn("Create a project", response.get_json()["error"]["message"])
        merge.assert_not_called()

    def test_pull_reports_conflicts_without_claiming_success(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "find_project_github_sync",
                return_value={
                    "package": "HousingForms",
                    "repository_url": "https://github.com/LegalAid/docassemble-HousingForms",
                    "branch": "main",
                    "commit": "same-sha",
                },
            ),
            patch.object(
                api_editor,
                "get_github_repository_snapshot",
                return_value={"sha": "same-sha", "branch": "main", "files": {}},
            ),
            patch.object(
                api_editor,
                "merge_github_snapshot",
                return_value={"merged": False, "conflicts": ["questions/main.yml"]},
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/github/pull", method="POST", json={"project": "Housing"}
            ):
                response = api_editor.editor_api_github_pull()

        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()["error"]["type"], "merge_conflict")
        self.assertEqual(
            response.get_json()["error"]["details"]["conflicts"], ["questions/main.yml"]
        )

    def test_status_treats_stale_github_credentials_as_disconnected(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "get_native_github_integration",
                return_value={
                    "enabled": True,
                    "connected": True,
                    "organizations_enabled": True,
                    "configure_url": "/github",
                },
            ),
            patch.object(
                api_editor,
                "get_github_publish_owners",
                side_effect=api_editor.GithubCredentialError(
                    "The GitHub connection could not be read; reconnect it in Docassemble"
                ),
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/github/status?project=Housing"
            ):
                response = api_editor.editor_api_github_status()

        self.assertEqual(response.status_code, 200)
        data = response.get_json()["data"]
        self.assertFalse(data["connected"])
        self.assertFalse(data["organizations_enabled"])
        self.assertEqual(data["owners"], [])
        self.assertEqual(data["configure_url"], "/al/editor/github/authorize")

    def test_status_returns_last_published_target_and_commit(self):
        commit = "a" * 40
        sync = {
            "package": "HousingForms",
            "repository_url": "https://github.com/LegalAid/docassemble-HousingForms",
            "branch": "feature/housing",
            "commit": commit,
        }
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "find_project_github_sync", return_value=sync),
            patch.object(
                api_editor,
                "get_native_github_integration",
                return_value={"enabled": True, "connected": True},
            ),
            patch.object(
                api_editor,
                "get_github_publish_owners",
                return_value=[{"login": "LegalAid", "type": "organization"}],
            ),
            patch.object(
                api_editor, "get_github_workflow_access", return_value={"owners": {}}
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/github/status?project=Housing"
            ):
                response = api_editor.editor_api_github_status()

        self.assertEqual(response.status_code, 200)
        saved = response.get_json()["data"]["sync"]
        self.assertEqual(saved["package"], "HousingForms")
        self.assertEqual(saved["owner"], "LegalAid")
        self.assertEqual(saved["branch"], "feature/housing")
        self.assertTrue(saved["published"])
        self.assertEqual(
            saved["commit_url"],
            "https://github.com/LegalAid/docassemble-HousingForms/commit/" + commit,
        )

    def test_status_does_not_claim_an_unpublished_manifest_was_published(self):
        sync = {
            "package": "HousingForms",
            "repository_url": "https://github.com/LegalAid/docassemble-HousingForms",
            "branch": "feature/housing",
            "commit": "",
        }
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "find_project_github_sync", return_value=sync),
            patch.object(
                api_editor,
                "get_native_github_integration",
                return_value={"enabled": False, "connected": False},
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/github/status?project=Housing"
            ):
                response = api_editor.editor_api_github_status()

        self.assertEqual(response.status_code, 200)
        saved = response.get_json()["data"]["sync"]
        self.assertEqual(saved["package"], "HousingForms")
        self.assertEqual(saved["branch"], "feature/housing")
        self.assertFalse(saved["published"])
        self.assertIsNone(saved["commit_url"])

    def test_status_marks_each_owner_with_its_workflow_access(self):
        access = {
            "token_type": "github_app",
            "owners": {
                "ada": {"status": "granted", "message": "", "action": "", "url": ""},
                "LegalAid": {
                    "status": "app_not_installed",
                    "message": "Not installed.",
                    "action": "install",
                    "url": "https://github.com/apps/da/installations/new",
                },
            },
        }
        for check in ({"return_value": access}, {"side_effect": RuntimeError("x")}):
            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(
                    api_editor,
                    "get_native_github_integration",
                    return_value={
                        "enabled": True,
                        "connected": True,
                        "organizations_enabled": True,
                    },
                ),
                patch.object(
                    api_editor,
                    "get_github_publish_owners",
                    return_value=[
                        {"login": "ada", "type": "user"},
                        {"login": "LegalAid", "type": "organization"},
                    ],
                ),
                patch.object(api_editor, "get_github_workflow_access", **check),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/github/status?project=Housing"
                ):
                    response = api_editor.editor_api_github_status()
            self.assertEqual(response.status_code, 200)
            owners = {
                owner["login"]: owner for owner in response.get_json()["data"]["owners"]
            }
            if "return_value" in check:
                self.assertEqual(
                    owners["LegalAid"]["workflow_access"]["status"],
                    "app_not_installed",
                )
                self.assertEqual(owners["ada"]["workflow_access"]["status"], "granted")
            else:
                # A failed check must not hide the owners or fail the dialog.
                self.assertEqual(set(owners), {"ada", "LegalAid"})
                self.assertNotIn("workflow_access", owners["ada"])

    MANIFEST_PATH = "/playground/packages/Housing/docassemble.HousingForms"
    MANIFEST_INFO = {
        "interview_files": ["main.yml"],
        "template_files": [],
        "module_files": [],
        "static_files": [],
        "sources_files": ["main.feature"],
        "dependencies": [],
        "description": "Housing forms",
        "license": "MIT License",
        "readme": "# Housing forms",
        "url": "",
        "version": "0.0.1",
    }

    def test_pull_with_no_upstream_change_leaves_local_settings_alone(self):
        remote = {"sha": "same-sha", "branch": "main", "files": {"a": b"1"}}
        for commit, expected_base in (("same-sha", remote["files"]), ("", None)):
            sync = {
                "package": "HousingForms",
                "repository_url": "https://github.com/LegalAid/docassemble-HousingForms",
                "branch": "main",
                "commit": commit,
            }
            with (
                self.subTest(commit=commit),
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(api_editor, "find_project_github_sync", return_value=sync),
                patch.object(
                    api_editor, "get_github_repository_snapshot", return_value=remote
                ),
                patch.object(
                    api_editor,
                    "merge_github_snapshot",
                    return_value={"merged": True, "files": 1, "commit": "same-sha"},
                ),
                patch.object(api_editor, "adopt_repository_snapshot") as adopt,
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/github/pull",
                    method="POST",
                    json={"project": "Housing"},
                ):
                    response = api_editor.editor_api_github_pull()
                self.assertEqual(response.status_code, 200)
                # A recorded sync point is the base even when GitHub has not
                # moved, so nothing is taken; only a manifest with no sync
                # point falls back to taking the repository whole.
                self.assertEqual(adopt.call_args.args[4], expected_base)

    def test_publish_prepares_manifest_and_queues_a_background_commit(self):
        """The request must not hold open the per-file GitHub round trips."""
        prepared = {
            "package": "HousingForms",
            "repository": "docassemble-HousingForms",
            "manifest_path": "/playground/packages/Housing/docassemble.HousingForms",
        }
        sent = {}

        def fake_send_task(task_name, kwargs=None, **options):
            sent["task_name"] = task_name
            sent["kwargs"] = kwargs
            sent["options"] = options
            return types.SimpleNamespace(id="celery-task-1")

        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "_editor_async_is_configured", return_value=True),
            patch.object(api_editor, "r", _FakeRedis()),
            patch.object(
                api_editor,
                "workerapp",
                types.SimpleNamespace(send_task=fake_send_task),
            ),
            patch.object(
                api_editor,
                "get_native_github_integration",
                return_value={
                    "enabled": True,
                    "connected": True,
                    "organizations_enabled": True,
                },
            ),
            patch.object(
                api_editor,
                "get_github_publish_owners",
                return_value=[
                    {"login": "ada", "type": "user"},
                    {"login": "LegalAid", "type": "organization"},
                ],
            ),
            patch.object(
                api_editor,
                "prepare_project_github_package",
                return_value=prepared,
            ) as prepare,
            patch.object(
                api_editor,
                "_verify_github_publish_preview",
                return_value={
                    "user_id": 7,
                    "project": "Housing",
                    "package": "HousingForms",
                    "owner": "LegalAid",
                    "repository_url": "https://github.com/LegalAid/docassemble-HousingForms",
                    "branch": "feature/github",
                    "source_revision": "source-digest",
                    "remote_sha": "remote-base-sha",
                },
            ),
            patch.object(
                api_editor,
                "load_project_github_manifest",
                return_value=(
                    {},
                    "/playground/packages/Housing/docassemble.HousingForms",
                ),
            ),
            patch.object(
                api_editor,
                "repository_publish_files",
                return_value={"files": {}, "managed_paths": []},
            ),
            patch.object(
                api_editor,
                "build_github_package_snapshot",
                return_value={"main.yml": {"content": b"---\n", "mode": "100644"}},
            ),
            patch.object(
                api_editor,
                "github_package_snapshot_revision",
                return_value="source-digest",
            ),
            patch.object(api_editor, "ensure_github_repository") as ensure_repository,
            patch.object(api_editor, "publish_github_package") as publish,
            patch.object(api_editor, "_editor_user_designator", return_value="Ada"),
            patch.object(
                api_editor,
                "repository_dependency_names",
                return_value=["docassemble.AssemblyLine"],
            ),
        ):
            api_editor.current_user.email = "ada@example.com"
            with api_editor.app.test_request_context(
                "/al/editor/api/github/publish",
                method="POST",
                json={
                    "project": "Housing",
                    "owner": "LegalAid",
                    "package": "HousingForms",
                    "branch": "feature/github",
                    "commit_message": "Update interview",
                    "preview_token": "signed-preview",
                },
            ):
                response = api_editor.editor_api_github_publish()

        self.assertEqual(response.status_code, 202, response.get_json())
        payload = response.get_json()
        self.assertEqual(payload["status"], "queued")
        data = payload["data"]
        self.assertEqual(data["repository"], "docassemble-HousingForms")
        self.assertEqual(data["owner"], "LegalAid")
        self.assertEqual(data["branch"], "feature/github")
        self.assertEqual(
            data["job_url"],
            f"/al/editor/api/github/publish/jobs/{payload['job_id']}",
        )
        self.assertEqual(data["state"]["status"], "queued")

        # No GitHub traffic happens in the request itself.
        ensure_repository.assert_not_called()
        publish.assert_not_called()
        prepare.assert_called_once_with(
            user_id=7,
            project_name="Housing",
            package_name="HousingForms",
            author_name="Ada",
            author_email="ada@example.com",
            dependencies=["docassemble.AssemblyLine"],
        )
        self.assertEqual(
            sent["task_name"],
            "docassemble.ALWeaver.api_weaver_worker.weaver_editor_github_publish_task",
        )
        self.assertEqual(
            sent["kwargs"],
            {
                "job_id": payload["job_id"],
                "uid": 7,
                "project": "Housing",
                "package": "HousingForms",
                "repository": "docassemble-HousingForms",
                "owner": "LegalAid",
                "owner_type": "organization",
                "author_name": "Ada",
                "author_email": "ada@example.com",
                "branch": "feature/github",
                "commit_message": "Update interview",
                "repository_url": "https://github.com/LegalAid/docassemble-HousingForms",
                "expected_remote_sha": "remote-base-sha",
                "expected_source_revision": "source-digest",
            },
        )
        self.assertEqual(sent["options"]["task_id"], payload["job_id"])

    def test_publish_without_a_preview_queues_the_commit(self):
        """Previewing is optional; the worker still guards synced branches."""
        sent = {}

        def fake_send_task(task_name, kwargs=None, **options):
            sent["kwargs"] = kwargs
            return types.SimpleNamespace(id="celery-task-1")

        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "_editor_async_is_configured", return_value=True),
            patch.object(api_editor, "r", _FakeRedis()),
            patch.object(
                api_editor,
                "workerapp",
                types.SimpleNamespace(send_task=fake_send_task),
            ),
            patch.object(
                api_editor,
                "get_native_github_integration",
                return_value={"enabled": True, "connected": True},
            ),
            patch.object(
                api_editor,
                "get_github_publish_owners",
                return_value=[{"login": "ada", "type": "user"}],
            ),
            patch.object(
                api_editor,
                "prepare_project_github_package",
                return_value={
                    "package": "HousingForms",
                    "repository": "docassemble-HousingForms",
                },
            ),
            patch.object(api_editor, "build_github_package_snapshot") as build_snapshot,
            patch.object(api_editor, "_editor_user_designator", return_value="Ada"),
            patch.object(api_editor, "repository_dependency_names", return_value=[]),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/github/publish",
                method="POST",
                json={
                    "project": "Housing",
                    "owner": "ada",
                    "package": "HousingForms",
                    "branch": "main",
                    "commit_message": "Update interview",
                },
            ):
                response = api_editor.editor_api_github_publish()

        self.assertEqual(response.status_code, 202, response.get_json())
        build_snapshot.assert_not_called()
        self.assertIsNone(sent["kwargs"]["expected_remote_sha"])
        self.assertIsNone(sent["kwargs"]["expected_source_revision"])

    def test_publish_refuses_when_celery_is_not_configured(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "_editor_async_is_configured", return_value=False),
            patch.object(
                api_editor,
                "get_worker_configuration_status",
                return_value={"configured": False, "message": "Not configured."},
            ),
            patch.object(api_editor, "prepare_project_github_package") as prepare,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/github/publish",
                method="POST",
                json={
                    "project": "Housing",
                    "owner": "LegalAid",
                    "package": "HousingForms",
                    "branch": "main",
                    "commit_message": "Update interview",
                },
            ):
                response = api_editor.editor_api_github_publish()

        self.assertEqual(response.status_code, 503)
        error = response.get_json()["error"]
        self.assertEqual(error["code"], "editor_async_not_configured")
        prepare.assert_not_called()

    def test_publish_job_runs_the_github_calls_and_records_the_commit(self):
        redis = _FakeRedis()
        with (
            patch.object(api_editor, "r", redis),
            patch.object(
                api_editor,
                "get_github_branch_head",
                return_value={"missing": True, "repository_exists": False, "sha": ""},
            ),
            patch.object(
                api_editor,
                "load_project_github_manifest",
                return_value=(dict(self.MANIFEST_INFO), self.MANIFEST_PATH),
            ) as load_manifest,
            patch.object(
                api_editor,
                "ensure_github_repository",
                return_value={
                    "html_url": "https://github.com/LegalAid/docassemble-HousingForms",
                    "default_branch": "main",
                    "created_by_weaver": True,
                },
            ) as ensure_repository,
            patch.object(
                api_editor,
                "publish_github_package",
                return_value={
                    "sha": "commit-sha",
                    "files": 12,
                    "warnings": ["Workflow changes skipped."],
                    "skipped_workflows": [".github/workflows/test.yml"],
                },
            ) as publish,
            patch.object(
                api_editor,
                "repository_publish_files",
                return_value={
                    "files": {
                        ".github/workflows/run_interview_tests.yml": "name: ALKiln",
                        "pyproject.toml": "[project]\n",
                    },
                    "managed_paths": [
                        ".github/workflows/run_interview_tests.yml",
                        "pyproject.toml",
                    ],
                    "dependency_names": [],
                },
            ) as repository_files,
        ):
            api_editor._store_job_state(
                api_editor.GITHUB_PUBLISH_JOB,
                "job-1",
                {
                    "status": "queued",
                    "error": {
                        "type": "job_expired",
                        "message": "stale transient error",
                    },
                },
            )
            result = api_editor._complete_github_publish_job(
                job_id="job-1",
                uid=7,
                project="Housing",
                package="HousingForms",
                repository="docassemble-HousingForms",
                owner="LegalAid",
                owner_type="organization",
                author_name="Ada",
                author_email="ada@example.com",
                branch="feature/github",
                commit_message="Update interview",
                repository_url="https://github.com/LegalAid/docassemble-HousingForms",
            )
            state = api_editor._load_job_state(api_editor.GITHUB_PUBLISH_JOB, "job-1")
            publish_kwargs = publish.call_args.kwargs
            # The progress hook writes through to the job record.
            publish_kwargs["on_progress"]("Uploading main.yml (1 of 2).", 50)
            progressed = api_editor._load_job_state(
                api_editor.GITHUB_PUBLISH_JOB, "job-1"
            )

        self.assertEqual(result["commit_sha"], "commit-sha")
        self.assertEqual(result["files_committed"], 12)
        self.assertEqual(result["warnings"], ["Workflow changes skipped."])
        self.assertEqual(result["skipped_workflows"], [".github/workflows/test.yml"])
        self.assertTrue(result["repository_created"])
        self.assertEqual(
            result["commit_url"],
            "https://github.com/LegalAid/docassemble-HousingForms/commit/commit-sha",
        )
        self.assertEqual(state["status"], "succeeded")
        self.assertIsNone(state.get("error"))
        self.assertEqual(state["progress"], 100)
        self.assertEqual(state["result"], result)

        ensure_repository.assert_called_once_with(
            owner="LegalAid",
            repository="docassemble-HousingForms",
            description="A docassemble project for Housing.",
            owner_type="organization",
            user_id=7,
        )
        # The manifest is re-read in the worker rather than trusting a path
        # handed across the queue, so this works on a multi-server install.
        load_manifest.assert_called_once_with(
            user_id=7,
            project_name="Housing",
            package_name="HousingForms",
        )
        self.assertEqual(publish_kwargs["package_info"], self.MANIFEST_INFO)
        self.assertEqual(publish_kwargs["manifest_path"], self.MANIFEST_PATH)
        self.assertEqual(publish_kwargs["default_branch"], "main")
        self.assertEqual(publish_kwargs["branch"], "feature/github")
        self.assertEqual(publish_kwargs["author_email"], "ada@example.com")
        # The worker builds the workflows and pyproject.toml from the stored
        # settings and the manifest it just re-read.
        repository_files.assert_called_once_with(
            7, "Housing", "HousingForms", manifest=self.MANIFEST_INFO
        )
        self.assertEqual(
            publish_kwargs["extra_repository_files"],
            repository_files.return_value["files"],
        )
        self.assertEqual(publish_kwargs["preserved_path_prefixes"], (".github/",))
        self.assertEqual(
            publish_kwargs["managed_paths"],
            repository_files.return_value["managed_paths"],
        )
        self.assertEqual(progressed["message"], "Uploading main.yml (1 of 2).")
        self.assertEqual(progressed["progress"], 55)

    def test_publish_job_records_a_lost_github_connection_as_a_failure(self):
        with (
            patch.object(api_editor, "r", _FakeRedis()),
            patch.object(
                api_editor,
                "get_github_branch_head",
                return_value={"missing": True, "repository_exists": False, "sha": ""},
            ),
            patch.object(
                api_editor,
                "ensure_github_repository",
                side_effect=api_editor.GithubCredentialError(
                    "The GitHub connection has expired; reconnect it in Docassemble"
                ),
            ),
        ):
            with self.assertRaises(api_editor.GithubCredentialError):
                api_editor._complete_github_publish_job(
                    job_id="job-2",
                    uid=7,
                    project="Housing",
                    package="HousingForms",
                    repository="docassemble-HousingForms",
                    owner="LegalAid",
                    owner_type="organization",
                    author_name="Ada",
                    author_email="ada@example.com",
                    branch="main",
                    commit_message="Update interview",
                    repository_url="https://github.com/LegalAid/docassemble-HousingForms",
                )
            state = api_editor._load_job_state(api_editor.GITHUB_PUBLISH_JOB, "job-2")

        self.assertEqual(state["status"], "failed")
        self.assertEqual(state["stage"], "ensure_repository")
        self.assertEqual(state["error"]["type"], "github_not_connected")

    def test_publish_refuses_to_replace_a_branch_that_advanced_since_sync(self):
        sync = {
            "package": "HousingForms",
            "repository_url": "https://github.com/LegalAid/docassemble-HousingForms",
            "branch": "feature/github",
            "commit": "base-sha",
        }
        remote = {
            "sha": "remote-newer-sha",
            "files": {"docassemble/HousingForms/data/questions/main.yml": b"newer"},
        }
        redis = _FakeRedis()
        with (
            patch.object(api_editor, "r", redis),
            patch.object(api_editor, "find_project_github_sync", return_value=sync),
            patch.object(
                api_editor, "get_github_branch_head", return_value=remote
            ) as read_remote,
            patch.object(api_editor, "ensure_github_repository") as ensure_repository,
            patch.object(api_editor, "publish_github_package") as publish,
        ):
            with self.assertRaisesRegex(ValueError, "has advanced"):
                api_editor._complete_github_publish_job(
                    job_id="job-remote-advance",
                    uid=7,
                    project="Housing",
                    package="HousingForms",
                    repository="docassemble-HousingForms",
                    owner="LegalAid",
                    owner_type="organization",
                    author_name="Ada",
                    author_email="ada@example.com",
                    branch="feature/github",
                    commit_message="Update interview",
                    repository_url="https://github.com/LegalAid/docassemble-HousingForms",
                )
            state = api_editor._load_job_state(
                api_editor.GITHUB_PUBLISH_JOB, "job-remote-advance"
            )

        self.assertEqual(state["status"], "failed")
        self.assertIn("Pull the remote changes", state["error"]["message"])
        read_remote.assert_called_once_with(
            repository_url="https://github.com/LegalAid/docassemble-HousingForms",
            ref="feature/github",
            user_id=7,
        )
        ensure_repository.assert_not_called()
        publish.assert_not_called()

    def test_publish_job_status_is_scoped_to_its_owner(self):
        redis = _FakeRedis()
        with patch.object(api_editor, "r", redis):
            api_editor._store_job_state(
                api_editor.GITHUB_PUBLISH_JOB,
                "job-3",
                {"status": "succeeded", "owner_user_id": 7, "result": {"files": 3}},
            )
            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/github/publish/jobs/job-3"
                ):
                    owner_response = api_editor.editor_api_github_publish_job("job-3")
            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=99),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/github/publish/jobs/job-3"
                ):
                    other_response = api_editor.editor_api_github_publish_job("job-3")

        self.assertEqual(owner_response.status_code, 200)
        self.assertEqual(owner_response.get_json()["status"], "succeeded")
        self.assertEqual(other_response.status_code, 404)

    def test_publish_rejects_invalid_branch_before_writing_manifest(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "prepare_project_github_package") as prepare,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/github/publish",
                method="POST",
                json={
                    "project": "Housing",
                    "owner": "ada",
                    "package": "HousingForms",
                    "branch": "bad branch",
                    "commit_message": "Update interview",
                },
            ):
                response = api_editor.editor_api_github_publish()

        self.assertEqual(response.status_code, 400)
        self.assertIn("valid Git branch", response.get_json()["error"]["message"])
        prepare.assert_not_called()


class TestEditorFilesApi(unittest.TestCase):
    def test_unavailable_project_returns_path_free_not_found(self):
        private_path = "/var/lib/docassemble/playground/42/private-project"
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "playground_list_yaml_files",
                side_effect=FileNotFoundError(private_path),
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/files?project=private-project"
            ):
                response = api_editor.editor_api_files()

        payload = response.get_json()
        self.assertEqual(response.status_code, 404)
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error"]["type"], "not_found")
        self.assertNotIn(private_path, response.get_data(as_text=True))
        self.assertNotIn("private-project", response.get_data(as_text=True))

    def test_owner_can_list_project_files(self):
        files = [{"filename": "main.yml", "label": "main.yml"}]
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor, "playground_list_yaml_files", return_value=files
            ) as list_files,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/files?project=owned-project"
            ):
                response = api_editor.editor_api_files()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["data"]["files"], files)
        list_files.assert_called_once_with(7, "owned-project")


class TestEditorProjectSearchApi(unittest.TestCase):
    def test_search_names_binary_and_oversized_files_it_cannot_inspect(self):
        with (
            patch.object(api_editor, "playground_list_yaml_files", return_value=[]),
            patch.object(
                api_editor,
                "_list_editor_section_files",
                side_effect=lambda uid, project, section: (
                    [
                        {"filename": "image.png", "editable": False, "size": 50},
                        {
                            "filename": "large.txt",
                            "editable": True,
                            "size": api_editor.EDITOR_SEARCH_MAX_FILE_BYTES + 1,
                        },
                    ]
                    if section == "static"
                    else []
                ),
            ),
            patch.object(api_editor, "_read_project_text_file") as read,
        ):
            files, skipped = api_editor._project_text_files(7, "default")
        self.assertEqual(files, [])
        self.assertEqual(
            skipped,
            [
                {
                    "section": "static",
                    "filename": "image.png",
                    "reason": "binary_or_unsupported",
                },
                {"section": "static", "filename": "large.txt", "reason": "too_large"},
            ],
        )
        read.assert_not_called()

    def test_search_returns_context_group_metadata_and_revisions(self):
        project_files = [
            {
                "section": "interview",
                "file_type": "interview",
                "file_type_label": "Interviews",
                "filename": "main.yml",
                "content": "question: Alpha\n",
                "revision": "revision-main",
            },
            {
                "section": "modules",
                "file_type": "modules",
                "file_type_label": "Modules",
                "filename": "helper.py",
                "content": "alpha = 1\n",
                "revision": "revision-helper",
            },
        ]
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "_project_text_files",
                return_value=(project_files, []),
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/project/search",
                method="POST",
                json={"project": "default", "query": "alpha", "mode": "text"},
            ):
                response = api_editor.editor_api_project_search()

        self.assertEqual(response.status_code, 200)
        data = response.get_json()["data"]
        self.assertEqual(data["match_count"], 2)
        self.assertEqual(data["file_count"], 2)
        self.assertEqual(data["files"][0]["file_type_label"], "Interviews")
        self.assertEqual(data["files"][0]["matches"][0]["line"], 1)
        self.assertEqual(data["files"][1]["revision"], "revision-helper")

    def test_replace_preflights_exact_spans_before_committing(self):
        source = "alpha and alpha\n"
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "_read_project_text_file", return_value=source),
            patch.object(api_editor, "_commit_project_replacements") as commit,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/project/replace",
                method="POST",
                json={
                    "project": "default",
                    "query": "alpha",
                    "replacement": "beta",
                    "mode": "text",
                    "files": [
                        {
                            "section": "interview",
                            "filename": "main.yml",
                            "revision": "test-revision",
                            "matches": [{"start": 0, "end": 5}],
                        }
                    ],
                },
            ):
                response = api_editor.editor_api_project_replace()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["data"]["replacement_count"], 1)
        changes = commit.call_args.args[2]
        self.assertEqual(changes[0]["updated"], "beta and alpha\n")

    def test_replace_rejects_stale_search_without_writing(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "_read_project_text_file", return_value="alpha\n"),
            patch.object(api_editor, "_commit_project_replacements") as commit,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/project/replace",
                method="POST",
                json={
                    "project": "default",
                    "query": "alpha",
                    "replacement": "beta",
                    "mode": "text",
                    "files": [
                        {
                            "section": "interview",
                            "filename": "main.yml",
                            "revision": "older-revision",
                            "matches": [{"start": 0, "end": 5}],
                        }
                    ],
                },
            ):
                response = api_editor.editor_api_project_replace()

        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()["error"]["code"], "stale_search")
        commit.assert_not_called()

    def test_variable_refactor_reuses_safe_rename_planning_across_a_file(self):
        source = (
            "---\n"
            "id: ask_name\n"
            "question: Name\n"
            "fields:\n"
            "  - Name: old_name\n"
            "---\n"
            "id: use_name\n"
            "code: |\n"
            "  if old_name:\n"
            "    pass\n"
        )
        project_files = [
            {
                "section": "interview",
                "file_type": "interview",
                "file_type_label": "Interviews",
                "filename": "main.yml",
                "content": source,
                "revision": "source-revision",
            }
        ]
        validation = types.SimpleNamespace(blocking=False)
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "_project_text_files",
                return_value=(project_files, []),
            ),
            patch.object(
                api_editor, "_project_search_revision", return_value="manifest"
            ),
            patch.object(
                api_editor, "validate_candidate_source", return_value=validation
            ),
            patch.object(api_editor, "_commit_project_replacements") as commit,
            patch.object(api_editor, "_list_editor_section_files", return_value=[]),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/project/replace",
                method="POST",
                json={
                    "project": "default",
                    "query": "old_name",
                    "replacement": "client_name",
                    "mode": "variable",
                    "project_revision": "manifest",
                },
            ):
                response = api_editor.editor_api_project_replace()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["data"]["replacement_count"], 2)
        updated = commit.call_args.args[2][0]["updated"]
        self.assertNotIn("old_name", updated)
        self.assertEqual(updated.count("client_name"), 2)

    def test_final_batch_preflight_detects_a_race_before_any_write(self):
        changes = [
            {
                "section": "interview",
                "filename": "one.yml",
                "original": "one old",
                "updated": "one new",
            },
            {
                "section": "modules",
                "filename": "two.py",
                "original": "two old",
                "updated": "two new",
            },
        ]
        with (
            patch.object(
                api_editor,
                "_read_project_text_file",
                side_effect=["one old", "two changed elsewhere"],
            ),
            patch.object(api_editor, "_write_project_text_file") as write,
        ):
            with self.assertRaises(api_editor.StaleProjectSearchError) as raised:
                api_editor._commit_project_replacements(7, "default", changes)

        self.assertEqual(raised.exception.files[0]["filename"], "two.py")
        write.assert_not_called()


class TestEditorLiteralFileDependencies(unittest.TestCase):
    def test_interview_rename_is_refused_when_a_local_include_names_the_file(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor, "_project_yaml_filenames", return_value=["main.yml"]
            ),
            patch.object(
                api_editor,
                "playground_read_yaml",
                return_value="---\ninclude:\n  - questions.yml\n",
            ),
            patch.object(api_editor, "rename_saved_file") as rename,
            api_editor.app.test_request_context(
                "/al/editor/api/file/rename",
                method="POST",
                json={
                    "project": "default",
                    "filename": "questions.yml",
                    "new_filename": "renamed_questions.yml",
                },
            ),
        ):
            response = api_editor.editor_api_rename_file()

        self.assertEqual(response.status_code, 409)
        error = response.get_json()["error"]
        self.assertEqual(error["code"], "file_has_references")
        self.assertEqual(error["dependencies"][0]["filename"], "main.yml")
        self.assertIn("Update those references first", error["message"])
        rename.assert_not_called()

    def test_interview_delete_is_refused_when_a_local_include_names_the_file(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor, "_project_yaml_filenames", return_value=["main.yml"]
            ),
            patch.object(
                api_editor,
                "playground_read_yaml",
                return_value="---\ninclude:\n  - questions.yml\n",
            ),
            patch.object(api_editor, "delete_saved_file") as delete,
            api_editor.app.test_request_context(
                "/al/editor/api/file/delete",
                method="POST",
                json={"project": "default", "filename": "questions.yml"},
            ),
        ):
            response = api_editor.editor_api_delete_file()

        self.assertEqual(response.status_code, 409)
        error = response.get_json()["error"]
        self.assertEqual(error["code"], "file_has_references")
        self.assertEqual(error["dependencies"][0]["filename"], "main.yml")
        delete.assert_not_called()

    def test_template_and_module_literal_references_are_reported(self):
        yaml_files = ["main.yml"]
        sources = {
            "main.yml": (
                "---\nmodules:\n  - helper_module\n"
                "---\nattachment:\n  docx template file: forms/answer.docx\n"
            )
        }
        with (
            patch.object(
                api_editor, "_project_yaml_filenames", return_value=yaml_files
            ),
            patch.object(
                api_editor,
                "playground_read_yaml",
                side_effect=lambda _uid, _project, filename: sources[filename],
            ),
            patch.object(
                api_editor,
                "_list_editor_section_files",
                return_value=[],
            ),
        ):
            module_refs = api_editor._literal_file_dependencies(
                7, "default", "modules", "helper_module.py"
            )
            template_refs = api_editor._literal_file_dependencies(
                7, "default", "templates", "answer.docx"
            )

        self.assertEqual(module_refs[0]["kind"], "module")
        self.assertEqual(template_refs[0]["kind"], "template")
        self.assertEqual(template_refs[0]["reference"], "forms/answer.docx")


class TestEditorApiFileCreation(unittest.TestCase):
    def test_filename_normalizers_reject_path_separators(self):
        normalizers = (
            api_editor._normalize_filename,
            api_editor._normalize_new_filename,
            api_editor._normalize_storage_filename,
        )
        for normalize in normalizers:
            for value in (
                "../outside.yml",
                "..\\outside.yml",
                "nested/interview.yml",
                "nested\\interview.yml",
                "/tmp/outside.yml",
            ):
                with self.subTest(normalizer=normalize.__name__, value=value):
                    with self.assertRaisesRegex(ValueError, "path separator"):
                        normalize(value)

        self.assertEqual(api_editor._normalize_filename("main.yml"), "main.yml")
        self.assertEqual(
            api_editor._normalize_storage_filename("answer.docx"), "answer.docx"
        )

    def test_github_import_derives_project_name_from_repository(self):
        snapshot = {
            "url": "https://github.com/OtherOrg/docassemble-PublicForms",
            "repository": "docassemble-PublicForms",
            "branch": "feature/x",
            "sha": "remote-sha",
            "files": {},
        }
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "get_list_of_projects", return_value=[]),
            patch.object(
                api_editor,
                "next_available_project_name",
                side_effect=lambda base, existing: base,
            ),
            patch.object(api_editor, "create_project") as create,
            patch.object(
                api_editor, "github_project_name", return_value="PublicFormsX"
            ) as name,
            patch.object(
                api_editor, "get_github_repository_snapshot", return_value=snapshot
            ),
            patch.object(
                api_editor,
                "import_github_snapshot",
                return_value={
                    "package": "PublicForms",
                    "filename": "main.yml",
                    "files_imported": 1,
                },
            ),
            patch.object(api_editor, "adopt_repository_snapshot"),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/new-project",
                method="POST",
                json={
                    "project_name": "",
                    "github_url": "https://github.com/OtherOrg/docassemble-PublicForms",
                    "create_test": False,
                },
            ):
                response = api_editor.editor_api_new_project()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["data"]["project"], "PublicFormsX")
        create.assert_called_once_with(7, "PublicFormsX")
        name.assert_called_once_with("docassemble-PublicForms", "feature/x")

    def test_new_project_can_import_any_github_repository_url(self):
        snapshot = {
            "url": "https://github.com/OtherOrg/docassemble-PublicForms",
            "repository": "docassemble-PublicForms",
            "branch": "main",
            "sha": "remote-sha",
            "files": {},
        }
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "get_list_of_projects", return_value=[]),
            patch.object(
                api_editor, "next_available_project_name", return_value="PublicForms"
            ),
            patch.object(api_editor, "create_project") as create,
            patch.object(
                api_editor, "get_github_repository_snapshot", return_value=snapshot
            ) as fetch,
            patch.object(
                api_editor,
                "import_github_snapshot",
                return_value={
                    "package": "PublicForms",
                    "filename": "main.yml",
                    "files_imported": 4,
                },
            ) as import_snapshot,
            patch.object(api_editor, "adopt_repository_snapshot") as adopt,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/new-project",
                method="POST",
                json={
                    "project_name": "PublicForms",
                    "github_url": "https://github.com/OtherOrg/docassemble-PublicForms",
                    "create_test": False,
                },
            ):
                response = api_editor.editor_api_new_project()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["data"]["filename"], "main.yml")
        create.assert_called_once_with(7, "PublicForms")
        self.assertEqual(
            fetch.call_args.kwargs["repository_url"],
            "https://github.com/OtherOrg/docassemble-PublicForms",
        )
        import_snapshot.assert_called_once_with(
            user_id=7, project_name="PublicForms", snapshot=snapshot
        )
        adopt.assert_called_once_with(
            7, "PublicForms", "PublicForms", snapshot["files"]
        )

    def test_failed_github_read_creates_no_project(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "create_project") as create,
            patch.object(
                api_editor,
                "get_github_repository_snapshot",
                side_effect=ValueError("GitHub repository was not found"),
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/new-project",
                method="POST",
                json={
                    "project_name": "",
                    "github_url": "https://github.com/OtherOrg/docassemble-Missing",
                },
            ):
                response = api_editor.editor_api_new_project()

        self.assertGreaterEqual(response.status_code, 400)
        create.assert_not_called()

    def test_save_file_accepts_intentionally_empty_source(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "playground_write_yaml") as mock_write,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/file",
                method="POST",
                json={
                    "project": "default",
                    "filename": "test.yml",
                    "content": "",
                },
            ):
                response = api_editor.editor_api_save_file()

        self.assertEqual(response.status_code, 200)
        mock_write.assert_called_once_with(7, "default", "test.yml", "")

    def test_save_file_rejects_missing_content(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "playground_write_yaml") as mock_write,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/file",
                method="POST",
                json={"project": "default", "filename": "test.yml"},
            ):
                response = api_editor.editor_api_save_file()

        self.assertEqual(response.status_code, 400)
        self.assertIn("content must be", response.get_json()["error"]["message"])
        mock_write.assert_not_called()

    def test_save_file_missing_owner_project_is_structured_not_found(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "playground_read_yaml",
                side_effect=FileNotFoundError(
                    "/usr/share/docassemble/files/playground/99/private.yml"
                ),
            ),
            patch.object(api_editor, "playground_write_yaml") as mock_write,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/file",
                method="POST",
                json={
                    "project": "another-users-project",
                    "filename": "private.yml",
                    "content": "---\nid: private\nquestion: Changed\n",
                },
            ):
                response = api_editor.editor_api_save_file()

        self.assertEqual(response.status_code, 404)
        payload = response.get_json()
        self.assertEqual(payload["error"]["type"], "not_found")
        self.assertEqual(payload["error"]["code"], "interview_file_not_found")
        self.assertNotIn("/usr/share/docassemble", response.get_data(as_text=True))
        mock_write.assert_not_called()

    def test_get_file_missing_owner_project_does_not_disclose_storage_path(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "playground_read_yaml",
                side_effect=FileNotFoundError(
                    "/usr/share/docassemble/files/playground/99/private.yml"
                ),
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/file?project=another-users-project&filename=private.yml",
                method="GET",
            ):
                response = api_editor.editor_api_get_file()

        self.assertEqual(response.status_code, 404)
        payload = response.get_json()
        self.assertEqual(payload["error"]["type"], "not_found")
        self.assertEqual(payload["error"]["code"], "interview_file_not_found")
        self.assertNotIn("/usr/share/docassemble", response.get_data(as_text=True))

    def test_invalid_source_requires_explicit_draft_confirmation(self):
        saved = "---\nid: example\nquestion: Saved\n"
        invalid = "---\nid: example\nquestion: Draft\ncode: |\n  def broken(:\n"
        finding = {
            "level": "error",
            "severity": "error",
            "message": "Python syntax error",
            "filename": "test.yml",
        }
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "playground_read_yaml", return_value=saved),
            patch.object(api_editor, "playground_write_yaml") as mock_write,
            patch.object(api_editor, "_validate_source_text", return_value=[finding]),
            patch.object(
                api_editor,
                "_lint_summary_for_findings",
                return_value={"error": 1, "warning": 0, "info": 0},
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/file",
                method="POST",
                json={
                    "project": "default",
                    "filename": "test.yml",
                    "content": invalid,
                    "expected_revision": api_editor.source_revision(saved),
                },
            ):
                response = api_editor.editor_api_save_file()

        self.assertEqual(response.status_code, 422)
        payload = response.get_json()
        self.assertEqual(payload["error"]["code"], "draft_confirmation_required")
        self.assertEqual(payload["error"]["details"]["blocking_count"], 1)
        self.assertEqual(payload["error"]["details"]["diagnostics"], [finding])
        mock_write.assert_not_called()

    def test_invalid_source_saves_only_with_explicit_draft_flag(self):
        saved = "---\nid: example\nquestion: Saved\n"
        invalid = "---\nid: example\nquestion: Draft\ncode: |\n  def broken(:\n"
        finding = {"level": "error", "message": "Python syntax error"}
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "playground_read_yaml", return_value=saved),
            patch.object(api_editor, "playground_write_yaml") as mock_write,
            patch.object(api_editor, "_validate_source_text", return_value=[finding]),
            patch.object(
                api_editor,
                "_lint_summary_for_findings",
                return_value={"error": 1, "warning": 0, "info": 0},
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/file",
                method="POST",
                json={
                    "project": "default",
                    "filename": "test.yml",
                    "content": invalid,
                    "expected_revision": api_editor.source_revision(saved),
                    "save_as_draft": True,
                },
            ):
                response = api_editor.editor_api_save_file()

        self.assertEqual(response.status_code, 200)
        mock_write.assert_called_once_with(7, "default", "test.yml", invalid)

    def test_optional_courtforms_metadata_does_not_require_draft_confirmation(self):
        saved = "---\nid: petition\nquestion: Petition\n"
        updated = saved + "# template update\n"
        finding = {
            "level": "error",
            "severity": "error",
            "source": "dayamlchecker",
            "message": "metadata block is missing common CourtFormsOnline publishing fields: can_I_use_this_form",
        }
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "playground_read_yaml", return_value=saved),
            patch.object(api_editor, "playground_write_yaml") as mock_write,
            patch.object(api_editor, "_validate_source_text", return_value=[finding]),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/file",
                method="POST",
                json={
                    "project": "default",
                    "filename": "test.yml",
                    "content": updated,
                    "expected_revision": api_editor.source_revision(saved),
                },
            ):
                response = api_editor.editor_api_save_file()

        self.assertEqual(response.status_code, 200, response.get_json())
        mock_write.assert_called_once_with(7, "default", "test.yml", updated)

    def test_validate_source_uses_submitted_buffer(self):
        submitted = "---\nid: unsaved\nquestion: Unsaved title\n"
        saved = "---\nid: saved\nquestion: Saved title\n"
        findings = [
            {
                "severity": "warning",
                "level": "warning",
                "message": "Unsaved diagnostic",
                "filename": "test.yml",
                "file_name": "test.yml",
                "block_id": "unsaved",
                "source_range": None,
                "yaml_path": None,
            }
        ]
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "playground_read_yaml", return_value=saved),
            patch.object(
                api_editor, "_validate_source_text", return_value=findings
            ) as mock_validate,
            patch.object(
                api_editor, "playground_get_variables"
            ) as mock_saved_variable_check,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/validate-source",
                method="POST",
                json={
                    "project": "default",
                    "filename": "test.yml",
                    "raw_yaml": submitted,
                    "revision": "test-revision",
                },
            ):
                response = api_editor.editor_api_validate_source()

        self.assertEqual(response.status_code, 200)
        payload = response.get_json()["data"]
        self.assertEqual(payload["scope"], "unsaved_source")
        self.assertEqual(payload["filename"], "test.yml")
        self.assertEqual(payload["diagnostics"], findings)
        self.assertEqual(payload["errors"], findings)
        self.assertTrue(payload["base_revision_matches"])
        mock_validate.assert_called_once_with(submitted, "test.yml")
        mock_saved_variable_check.assert_not_called()

    def test_validate_source_handles_yaml_beyond_recursion_depth(self):
        from .editor_utils import parse_interview_yaml

        saved = "---\nid: saved\nquestion: Saved title\n"
        deep_yaml = "value: " + "{item: " * 400 + "nested" + "}" * 400 + "\n"
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "playground_read_yaml", return_value=saved),
            patch.object(
                api_editor, "parse_interview_yaml", wraps=parse_interview_yaml
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/validate-source",
                method="POST",
                json={
                    "project": "default",
                    "filename": "deep.yml",
                    "raw_yaml": deep_yaml,
                },
            ):
                response = api_editor.editor_api_validate_source()

        self.assertEqual(response.status_code, 200, response.get_json())
        diagnostic = response.get_json()["data"]["diagnostics"][0]
        self.assertEqual(diagnostic["level"], "error")
        self.assertEqual(diagnostic["source"], "yaml-parser")
        self.assertIn("supported validation depth", diagnostic["message"])

    def test_get_file_loads_deep_yaml_as_recoverable_unparseable_block(self):
        from .editor_utils import parse_interview_yaml

        deep_yaml = "value: " + "{item: " * 400 + "nested" + "}" * 400 + "\n"
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "playground_read_yaml", return_value=deep_yaml),
            patch.object(
                api_editor, "parse_interview_yaml", wraps=parse_interview_yaml
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/file?project=default&filename=deep.yml",
                method="GET",
            ):
                response = api_editor.editor_api_get_file()

        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(
            response.get_json()["data"]["blocks"][0]["title"], "Unparseable block"
        )

    def test_block_save_rejects_deep_yaml_with_bounded_error(self):
        from .editor_utils import _safe_load_interview_document

        deep_yaml = "value: " + "{item: " * 400 + "nested" + "}" * 400 + "\n"
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "_safe_load_interview_document",
                wraps=_safe_load_interview_document,
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/block",
                method="POST",
                json={
                    "project": "default",
                    "filename": "deep.yml",
                    "block_id": "deep",
                    "block_yaml": deep_yaml,
                },
            ):
                response = api_editor.editor_api_save_block()

        self.assertEqual(response.status_code, 400, response.get_json())
        self.assertIn(
            "supported validation depth",
            response.get_json()["error"]["message"],
        )

    def test_deep_yaml_cannot_be_saved_even_as_a_draft(self):
        from .editor_agent_validation import validate_source_text

        saved = "---\nid: saved\nquestion: Saved title\n"
        deep_yaml = "value: " + "{item: " * 400 + "nested" + "}" * 400 + "\n"
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "playground_read_yaml", return_value=saved),
            patch.object(
                api_editor, "_validate_source_text", wraps=validate_source_text
            ),
            patch.object(api_editor, "playground_write_yaml") as mock_write,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/file",
                method="POST",
                json={
                    "project": "default",
                    "filename": "deep.yml",
                    "content": deep_yaml,
                    "expected_revision": api_editor.source_revision(saved),
                    "save_as_draft": True,
                },
            ):
                response = api_editor.editor_api_save_file()

        self.assertEqual(response.status_code, 422, response.get_json())
        self.assertEqual(response.get_json()["error"]["code"], "yaml_nesting_too_deep")
        mock_write.assert_not_called()

    def test_get_file_returns_exact_raw_yaml_for_populated_and_empty_files(self):
        model = {
            "blocks": [],
            "metadata_blocks": [],
            "include_blocks": [],
            "default_screen_parts_blocks": [],
            "order_blocks": [],
        }
        for source in ("metadata:\n  title: 'Exact'\n", ""):
            with self.subTest(source=source):
                with (
                    patch.object(api_editor, "_editor_auth_check", return_value=True),
                    patch.object(api_editor, "_current_user_id", return_value=7),
                    patch.object(
                        api_editor, "playground_read_yaml", return_value=source
                    ),
                    patch.object(
                        api_editor, "parse_interview_yaml", return_value=model
                    ),
                ):
                    with api_editor.app.test_request_context(
                        "/al/editor/api/file?project=default&filename=test.yml",
                        method="GET",
                    ):
                        response = api_editor.editor_api_get_file()

                payload = response.get_json()["data"]
                self.assertEqual(payload["filename"], "test.yml")
                self.assertEqual(payload["raw_yaml"], source)
                self.assertIn("revision", payload)

    def test_new_project_route_uploads_docx_queues_background_job(self):
        docx_path = Path(__file__).parent / "test/test_docx_no_pdf_field_names.docx"

        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "get_list_of_projects", return_value=[]),
            patch.object(api_editor, "_editor_async_is_configured", return_value=True),
            patch.object(
                api_editor, "next_available_project_name", return_value="DocxSmoke"
            ),
            patch.object(api_editor, "get_default_github_owner", return_value="myorg"),
            patch.object(api_editor, "create_project") as mock_create_project,
            patch.object(api_editor, "_start_new_project_upload_job") as mock_start_job,
        ):
            mock_start_job.return_value = {
                "job_id": "job-123",
                "job_url": "/al/editor/api/new-project/jobs/job-123",
                "state": {
                    "status": "queued",
                    "project": "DocxSmoke",
                    "generated_from": docx_path.name,
                    "uploaded_count": 1,
                },
            }
            with api_editor.app.test_client() as client:
                with docx_path.open("rb") as docx_handle:
                    response = client.post(
                        "/al/editor/api/new-project",
                        data={
                            "project_name": "DocxSmoke",
                            "generation_notes": "Demand Letter",
                            "help_source_text": "Demand letter context",
                            "help_page_url": "https://example.com/help",
                            "help_page_title": "Help page title",
                            "use_llm_assist": "true",
                            "files": (
                                BytesIO(docx_handle.read()),
                                docx_path.name,
                            ),
                        },
                        content_type="multipart/form-data",
                    )

        payload = response.get_json()
        self.assertTrue(payload["success"])
        self.assertEqual(response.status_code, 202, response.get_json())
        self.assertEqual(payload["status"], "queued")
        self.assertEqual(payload["job_id"], "job-123")
        self.assertEqual(payload["job_url"], "/al/editor/api/new-project/jobs/job-123")
        self.assertEqual(payload["data"]["project"], "DocxSmoke")
        self.assertEqual(payload["data"]["generated_from"], docx_path.name)
        self.assertEqual(payload["data"]["uploaded_count"], 1)
        mock_create_project.assert_called_once_with(7, "DocxSmoke")
        mock_start_job.assert_called_once()
        start_kwargs = mock_start_job.call_args.kwargs
        self.assertEqual(start_kwargs["uid"], 7)
        self.assertEqual(start_kwargs["request_id"], payload["request_id"])
        self.assertEqual(start_kwargs["project_name"], "DocxSmoke")
        self.assertEqual(
            start_kwargs["generation_options"]["exact_name"], docx_path.name
        )
        self.assertEqual(
            start_kwargs["generation_options"]["help_source_text"],
            "Demand letter context",
        )
        self.assertEqual(
            start_kwargs["generation_options"]["help_page_url"],
            "https://example.com/help",
        )
        self.assertEqual(
            start_kwargs["generation_options"]["help_page_title"], "Help page title"
        )
        self.assertTrue(start_kwargs["generation_options"]["use_llm_assist"])
        self.assertFalse(start_kwargs["generation_options"]["create_package_zip"])
        self.assertFalse(start_kwargs["generation_options"]["separate_main_order"])
        self.assertTrue(start_kwargs["generation_options"]["include_next_steps"])
        self.assertTrue(start_kwargs["generation_options"]["include_download_screen"])
        self.assertTrue(
            start_kwargs["generation_options"]["interview_overrides"][
                "next_steps_enabled"
            ]
        )
        self.assertEqual(
            start_kwargs["generation_options"]["interview_overrides"]["github_user"],
            "myorg",
        )
        self.assertEqual(len(start_kwargs["uploaded_files"]), 1)
        self.assertEqual(start_kwargs["uploaded_files"][0]["filename"], docx_path.name)
        self.assertIsInstance(start_kwargs["uploaded_files"][0]["content_bytes"], bytes)
        self.assertEqual(
            start_kwargs["uploaded_files"][0]["mimetype"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    def test_new_project_rejects_malformed_upload_before_project_creation(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "get_list_of_projects", return_value=[]),
            patch.object(api_editor, "_editor_async_is_configured", return_value=True),
            patch.object(
                api_editor, "next_available_project_name", return_value="BadUpload"
            ),
            patch.object(
                api_editor,
                "validate_document_content",
                side_effect=ValueError("The PDF file is unreadable or malformed."),
            ),
            patch.object(api_editor, "create_project") as create_project,
            patch.object(api_editor, "_start_new_project_upload_job") as start_job,
        ):
            with api_editor.app.test_client() as client:
                response = client.post(
                    "/al/editor/api/new-project",
                    data={
                        "project_name": "BadUpload",
                        "files": (BytesIO(b"not a PDF"), "malformed.pdf"),
                    },
                    content_type="multipart/form-data",
                )

        payload = response.get_json()
        self.assertEqual(response.status_code, 400, payload)
        self.assertIn(
            "PDF file is unreadable or malformed", payload["error"]["message"]
        )
        create_project.assert_not_called()
        start_job.assert_not_called()

    def test_new_project_upload_refuses_when_celery_is_not_configured(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "_editor_async_is_configured", return_value=False),
            patch.object(api_editor, "create_project") as mock_create_project,
        ):
            with api_editor.app.test_client() as client:
                response = client.post(
                    "/al/editor/api/new-project",
                    data={
                        "project_name": "DocxSmoke",
                        "files": (BytesIO(b"not read"), "source.docx"),
                    },
                    content_type="multipart/form-data",
                )

        self.assertEqual(response.status_code, 503)
        self.assertEqual(
            response.get_json()["error"]["code"], "editor_async_not_configured"
        )
        details = response.get_json()["error"]["details"]
        self.assertFalse(details["configured"])
        self.assertIn("#celery-worker-configuration", details["docs_url"])
        mock_create_project.assert_not_called()

    def test_start_new_project_job_enqueues_celery_without_daemon_thread(self):
        async_result = types.SimpleNamespace(id="celery-task-1")
        with (
            patch.object(api_editor, "_store_new_project_job_state") as mock_store,
            patch.object(api_editor, "_update_new_project_job_state") as mock_update,
            patch.object(
                api_editor.workerapp, "send_task", return_value=async_result
            ) as mock_send,
        ):
            result = api_editor._start_new_project_upload_job(
                uid=7,
                request_id="req-1",
                project_name="DocxSmoke",
                uploaded_files=[
                    {
                        "filename": "source.docx",
                        "content_bytes": b"content",
                        "mimetype": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    }
                ],
                generation_options={"include_next_steps": False},
                debug_requested=False,
            )

        self.assertEqual(result["state"]["status"], "queued")
        self.assertEqual(result["state"]["owner_user_id"], 7)
        self.assertEqual(result["state"]["operation_type"], "new_project_upload")
        self.assertEqual(result["state"]["celery_task_id"], result["job_id"])
        mock_store.assert_called_once()
        mock_send.assert_called_once()
        self.assertEqual(mock_send.call_args.kwargs["task_id"], result["job_id"])
        self.assertEqual(
            mock_send.call_args.kwargs["kwargs"]["job_id"], result["job_id"]
        )
        mock_update.assert_not_called()

    def test_metadata_save_preserves_unrelated_source_exactly(self):
        from . import editor_utils as real_editor_utils

        source = (
            "# header\n"
            "metadata:\n"
            "  title: 'Original' # title comment\n"
            "---\n"
            "# unrelated comment\n"
            "id: intro\n"
            "question: |\n"
            "  Keep this exactly.\n"
        )
        edited = "# header\nmetadata:\n  title: 'Edited' # title comment"
        revision = real_editor_utils.source_revision(source)

        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "playground_read_yaml", return_value=source),
            patch.object(api_editor, "playground_write_yaml") as mock_write,
            patch.object(
                api_editor,
                "source_revision",
                side_effect=real_editor_utils.source_revision,
            ),
            patch.object(
                api_editor,
                "update_metadata_documents_in_yaml",
                side_effect=real_editor_utils.update_metadata_documents_in_yaml,
            ),
            patch.object(
                api_editor,
                "parse_interview_yaml",
                side_effect=real_editor_utils.parse_interview_yaml,
            ),
            patch.object(
                api_editor,
                "metadata_source_slice",
                side_effect=real_editor_utils.metadata_source_slice,
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/file/metadata",
                method="POST",
                json={
                    "project": "default",
                    "filename": "test.yml",
                    "raw_yaml": edited,
                    "expected_revision": revision,
                },
            ):
                response = api_editor.editor_api_save_metadata()

        expected = source.replace("'Original'", "'Edited'")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["data"]["raw_yaml"], expected)
        mock_write.assert_called_once_with(7, "default", "test.yml", expected)

    def test_assemblyline_settings_get_returns_schema_and_revision(self):
        source = "metadata:\n  title: Example\n"
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "playground_read_yaml", return_value=source),
            patch.object(
                api_editor,
                "read_settings",
                return_value={"schema": [], "values": {"title": "Example"}},
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/assemblyline-settings?project=default&filename=test.yml"
            ):
                response = api_editor.editor_api_get_assemblyline_settings()

        payload = response.get_json()["data"]
        self.assertEqual(payload["values"]["title"], "Example")
        self.assertEqual(payload["revision"], "test-revision")

    def test_assemblyline_settings_save_rejects_stale_revision(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "playground_read_yaml", return_value="source"),
            patch.object(api_editor, "update_settings", return_value="updated source"),
            patch.object(
                api_editor,
                "validate_candidate_source",
                return_value=types.SimpleNamespace(
                    blocking=False, diagnostics=[], model=None
                ),
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/assemblyline-settings",
                method="POST",
                json={
                    "project": "default",
                    "filename": "test.yml",
                    "expected_revision": "old-revision",
                    "settings": {"title": "Changed"},
                },
            ):
                response = api_editor.editor_api_save_assemblyline_settings()

        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()["error"]["code"], "revision_conflict")


class TestEditorJobReconciliation(unittest.TestCase):
    def reconcile(self, celery_state, *, task_result=None, initial=None):
        state = initial or {
            "status": "queued",
            "stage": "queued",
            "celery_task_id": "synthetic-task-id",
            "queued_at": 100,
        }
        with (
            patch.object(
                api_editor.workerapp,
                "AsyncResult",
                return_value=types.SimpleNamespace(
                    state=celery_state, result=task_result
                ),
            ),
            patch.object(
                api_editor,
                "_update_job_state",
                side_effect=lambda _kind, _job_id, **updates: {
                    **state,
                    **updates,
                },
            ) as update,
        ):
            result = api_editor._reconcile_new_project_job_state("synthetic-job", state)
        return result, update

    def test_success_persists_task_result_and_terminal_progress(self):
        state = {
            "status": "queued",
            "stage": "queued",
            "celery_task_id": "synthetic-task-id",
            "queued_at": 100,
            "error": {"type": "job_expired", "message": "stale transient error"},
        }
        result, update = self.reconcile(
            "SUCCESS", task_result={"project": "Synthetic"}, initial=state
        )
        self.assertEqual(result["status"], "succeeded")
        self.assertEqual(result["stage"], "done")
        self.assertEqual(result["progress"], 100)
        self.assertEqual(result["result"], {"project": "Synthetic"})
        self.assertIsNotNone(result["finished_at"])
        self.assertIsNone(result["error"])
        update.assert_called_once()

    def test_success_with_none_task_result_keeps_durable_job_result(self):
        state = {
            "status": "running",
            "stage": "copy_templates",
            "celery_task_id": "synthetic-task-id",
            "result": {
                "project": "Synthetic",
                "filename": "main.yml",
                "woven_templates": ["form.pdf"],
            },
            "partial_artifacts": ["main.yml", "form.pdf"],
            "incomplete_artifacts": [],
        }
        result, _update = self.reconcile("SUCCESS", task_result=None, initial=state)
        self.assertEqual(result["status"], "succeeded")
        self.assertEqual(result["result"], state["result"])
        self.assertEqual(result["partial_artifacts"], state["partial_artifacts"])
        self.assertEqual(result["incomplete_artifacts"], [])

    def test_job_status_api_keeps_durable_result_when_celery_returns_none(self):
        durable_result = {
            "project": "Synthetic",
            "filename": "main.yml",
            "woven_templates": ["form.pdf"],
        }
        state = {
            "status": "running",
            "stage": "copy_templates",
            "owner_user_id": 7,
            "celery_task_id": "synthetic-task-id",
            "result": durable_result,
            "partial_artifacts": ["main.yml", "form.pdf"],
            "incomplete_artifacts": [],
        }
        current_state = dict(state)
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "_load_new_project_job_state", return_value=state),
            patch.object(
                api_editor.workerapp,
                "AsyncResult",
                return_value=types.SimpleNamespace(state="SUCCESS", result=None),
            ),
            patch.object(
                api_editor,
                "_update_job_state",
                side_effect=lambda _kind, _job_id, **updates: current_state.update(
                    updates
                )
                or current_state,
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/new-project/jobs/synthetic-job"
            ):
                response = api_editor.editor_api_new_project_job("synthetic-job")

        payload = response.get_json()
        self.assertEqual(payload["data"]["status"], "succeeded")
        self.assertEqual(payload["data"]["result"], durable_result)
        self.assertEqual(payload["data"]["partial_artifacts"], ["main.yml", "form.pdf"])

    def test_worker_crash_becomes_structured_failure(self):
        result, _update = self.reconcile(
            "FAILURE", task_result=RuntimeError("worker died")
        )
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["error"]["type"], "celery_failure")
        self.assertIn("worker died", result["error"]["message"])
        self.assertIsNotNone(result["finished_at"])

    def test_revoked_and_unknown_tasks_become_terminal(self):
        cancelled, _ = self.reconcile("REVOKED")
        self.assertEqual(cancelled["status"], "cancelled")
        expired, _ = self.reconcile("UNKNOWN")
        self.assertEqual(expired["status"], "expired")
        self.assertEqual(expired["error"]["type"], "job_expired")

    def test_revoked_upload_job_status_is_visible_only_to_its_owner(self):
        redis = _FakeRedis()
        initial = {
            "status": "running",
            "stage": "generate_interview",
            "owner_user_id": 7,
            "celery_task_id": "synthetic-revoked-upload-task",
            "queued_at": 100,
            "project": "MatrixSynthetic",
            "generated_from": "matrix-private-canary.docx",
            "result": None,
        }
        with patch.object(api_editor, "r", redis):
            api_editor._store_job_state(
                api_editor.NEW_PROJECT_JOB, "synthetic-revoked-upload", initial
            )
            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(
                    api_editor.workerapp,
                    "AsyncResult",
                    return_value=types.SimpleNamespace(state="REVOKED", result=None),
                ),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/new-project/jobs/synthetic-revoked-upload"
                ):
                    owner_response = api_editor.editor_api_new_project_job(
                        "synthetic-revoked-upload"
                    )
            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=99),
                patch.object(api_editor.workerapp, "AsyncResult") as async_result,
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/new-project/jobs/synthetic-revoked-upload"
                ):
                    other_response = api_editor.editor_api_new_project_job(
                        "synthetic-revoked-upload"
                    )

        self.assertEqual(owner_response.status_code, 200)
        self.assertEqual(owner_response.get_json()["status"], "cancelled")
        self.assertEqual(
            owner_response.get_json()["data"]["message"], "The job was cancelled."
        )
        self.assertEqual(other_response.status_code, 404)
        self.assertNotIn(
            "matrix-private-canary.docx", other_response.get_data(as_text=True)
        )
        async_result.assert_not_called()

    def test_missing_celery_task_id_expires_instead_of_staying_queued(self):
        state = {"status": "queued", "stage": "queued"}
        with patch.object(
            api_editor,
            "_update_job_state",
            side_effect=lambda _kind, _job, **u: {**state, **u},
        ):
            result = api_editor._reconcile_new_project_job_state("synthetic-job", state)
        self.assertEqual(result["status"], "expired")
        self.assertEqual(result["error"]["type"], "job_expired")

    def test_missing_celery_task_id_waits_through_the_enqueue_race_window(self):
        state = {
            "status": "queued",
            "stage": "queued",
            "queued_at": api_editor.time.time(),
        }
        with patch.object(api_editor, "_update_job_state") as update:
            result = api_editor._reconcile_new_project_job_state("synthetic-job", state)
        self.assertEqual(result, state)
        update.assert_not_called()

    def test_active_and_delayed_states_remain_nonterminal(self):
        running, _ = self.reconcile("STARTED")
        self.assertEqual(running["status"], "running")
        self.assertIsNotNone(running["started_at"])
        queued, _ = self.reconcile("PENDING")
        self.assertEqual(queued["status"], "queued")

    def test_terminal_job_is_not_reconciled_again(self):
        state = {"status": "failed", "error": {"type": "validation_error"}}
        with (
            patch.object(api_editor.workerapp, "AsyncResult") as async_result,
            patch.object(api_editor, "_update_job_state") as update,
        ):
            result = api_editor._reconcile_new_project_job_state("synthetic-job", state)
        self.assertEqual(result, state)
        async_result.assert_not_called()
        update.assert_not_called()


class TestEditorNewProjectNaming(unittest.TestCase):
    """A generated project is named after its document, not "interview.yml"."""

    def _run_upload_job(self, yaml_filename, interview_filename=None):
        with (
            patch.object(api_editor, "_update_new_project_job_state"),
            patch.object(api_editor, "playground_write_yaml") as mock_write,
            patch.object(api_editor, "_copy_files_to_section"),
            patch.object(
                api_editor,
                "generate_interview_from_bytes",
                return_value={
                    "yaml_text": "metadata:\n  title: Eviction\n",
                    "yaml_filename": yaml_filename,
                    "input_filename": "eviction.pdf",
                    "generated_template_files": [],
                },
            ),
        ):
            result = api_editor._complete_new_project_upload_job(
                job_id="job-1",
                uid=7,
                project_name="Eviction",
                request_id="req-1",
                uploaded_files=[
                    {
                        "filename": "eviction.pdf",
                        "content_bytes": b"%PDF-1.4",
                        "mimetype": "application/pdf",
                    }
                ],
                generation_options={},
                debug_requested=False,
                interview_filename=interview_filename,
            )
        return result, mock_write

    def test_generated_project_keeps_the_descriptive_name_weaver_derived(self):
        result, mock_write = self._run_upload_job("eviction.yml")
        self.assertEqual(result["filename"], "eviction.yml")
        self.assertEqual(mock_write.call_args.args[2], "eviction.yml")

    def test_author_supplied_filename_wins(self):
        result, mock_write = self._run_upload_job(
            "eviction.yml", interview_filename="main.yml"
        )
        self.assertEqual(result["filename"], "main.yml")
        self.assertEqual(mock_write.call_args.args[2], "main.yml")

    def test_an_unusable_derived_name_falls_back_to_the_assemblyline_default(self):
        result, _mock_write = self._run_upload_job("")
        self.assertEqual(result["filename"], "main.yml")


class TestEditorNewProjectPartialArtifacts(unittest.TestCase):
    def test_new_project_failure_log_omits_uploaded_filename_and_exception_path(self):
        private_marker = "/private/matrix-docs/secret-name.pdf"
        source_template = {"filename": "secret-name.pdf", "content_bytes": b"%PDF"}
        with (
            patch.object(api_editor, "_update_new_project_job_state") as update,
            patch.object(
                api_editor,
                "generate_interview_from_bytes",
                side_effect=OSError(f"failed to read {private_marker}"),
            ),
            patch.object(api_editor, "log") as log_call,
        ):
            with self.assertRaisesRegex(OSError, "failed to read"):
                api_editor._complete_new_project_upload_job(
                    job_id="synthetic-job",
                    uid=7,
                    project_name="SyntheticPetition",
                    request_id="synthetic-request",
                    uploaded_files=[source_template],
                    generation_options={},
                    debug_requested=False,
                    create_test=False,
                )

        failed_update = update.call_args.kwargs
        self.assertEqual(failed_update["status"], "failed")
        self.assertEqual(
            failed_update["error"],
            {"type": "server_error", "message": "ALWeaver generation failed."},
        )
        logged_text = " ".join(str(call.args[0]) for call in log_call.call_args_list)
        self.assertNotIn("secret-name.pdf", logged_text)
        self.assertNotIn(private_marker, logged_text)
        self.assertIn("exception_type=OSError", logged_text)

    def test_failed_managed_test_write_reports_known_incomplete_feature(self):
        source_template = {"filename": "petition.pdf", "content_bytes": b"%PDF"}
        generator_result = {
            "yaml_text": "metadata:\n  title: Petition\n",
            "yaml_filename": "petition.yml",
            "input_filename": "petition.pdf",
            "template_filenames": ["petition.pdf"],
        }
        with (
            patch.object(api_editor, "_update_new_project_job_state") as update,
            patch.object(api_editor, "playground_write_yaml"),
            patch.object(
                api_editor,
                "generate_interview_from_bytes",
                return_value=generator_result,
            ),
            patch.object(api_editor, "_write_default_kiln_test") as write_test,
        ):
            write_test.side_effect = OSError("synthetic feature write failure")
            with self.assertRaisesRegex(OSError, "synthetic feature write failure"):
                api_editor._complete_new_project_upload_job(
                    job_id="synthetic-job",
                    uid=7,
                    project_name="SyntheticPetition",
                    request_id="synthetic-request",
                    uploaded_files=[source_template],
                    generation_options={},
                    debug_requested=False,
                    create_test=True,
                )

        failed_update = update.call_args.kwargs
        self.assertEqual(failed_update["status"], "failed")
        self.assertEqual(failed_update["stage"], "write_yaml")
        self.assertEqual(failed_update["partial_artifacts"], ["petition.yml"])
        self.assertEqual(
            failed_update["incomplete_artifacts"], ["weaver_it_runs.feature"]
        )
        self.assertEqual(
            failed_update["result"],
            {
                "project": "SyntheticPetition",
                "partial_artifacts": ["petition.yml"],
                "incomplete_artifacts": ["weaver_it_runs.feature"],
                "incomplete_stage": "write_yaml",
            },
        )
        self.assertTrue(
            any(
                call.kwargs.get("partial_artifacts") == ["petition.yml"]
                and call.kwargs.get("incomplete_artifacts")
                == ["weaver_it_runs.feature"]
                for call in update.call_args_list
            )
        )

    def test_lost_task_reconciliation_preserves_artifact_checkpoint(self):
        checkpoint = {
            "status": "running",
            "stage": "copy_templates",
            "celery_task_id": "synthetic-task",
            "partial_artifacts": ["petition.yml", "weaver_it_runs.feature"],
            "incomplete_artifacts": ["petition.pdf"],
            "result": {
                "project": "SyntheticPetition",
                "partial_artifacts": ["petition.yml", "weaver_it_runs.feature"],
                "incomplete_artifacts": ["petition.pdf"],
                "incomplete_stage": "copy_templates",
            },
        }
        reconciled_state = {**checkpoint}
        with (
            patch.object(
                api_editor.workerapp,
                "AsyncResult",
                return_value=SimpleNamespace(state="FAILURE", result="worker lost"),
            ),
            patch.object(
                api_editor,
                "_update_job_state",
                side_effect=lambda _kind, _job_id, **updates: reconciled_state.update(
                    updates
                )
                or reconciled_state,
            ),
        ):
            result = api_editor._reconcile_new_project_job_state(
                "synthetic-job", checkpoint
            )

        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["partial_artifacts"], checkpoint["partial_artifacts"])
        self.assertEqual(
            result["incomplete_artifacts"], checkpoint["incomplete_artifacts"]
        )
        self.assertEqual(result["result"], checkpoint["result"])

    def test_failed_template_copy_reports_saved_and_incomplete_logical_files(self):
        source_template = {"filename": "petition.pdf", "content_bytes": b"%PDF"}
        generated_template = {
            "filename": "petition_attachment.pdf",
            "content_bytes": b"%PDF-generated",
        }
        generator_result = {
            "yaml_text": "metadata:\n  title: Petition\n",
            "yaml_filename": "petition.yml",
            "input_filename": "petition.pdf",
            "template_filenames": ["petition.pdf"],
            "generated_template_files": [generated_template],
        }
        with (
            patch.object(api_editor, "_update_new_project_job_state") as update,
            patch.object(api_editor, "playground_write_yaml"),
            patch.object(
                api_editor,
                "generate_interview_from_bytes",
                return_value=generator_result,
            ),
            patch.object(
                api_editor,
                "_copy_files_to_section",
                side_effect=OSError("synthetic copy failure"),
            ),
        ):
            with self.assertRaisesRegex(OSError, "synthetic copy failure"):
                api_editor._complete_new_project_upload_job(
                    job_id="synthetic-job",
                    uid=7,
                    project_name="SyntheticPetition",
                    request_id="synthetic-request",
                    uploaded_files=[source_template],
                    generation_options={},
                    debug_requested=False,
                    create_test=False,
                )

        failed_update = update.call_args.kwargs
        self.assertEqual(failed_update["status"], "failed")
        self.assertEqual(failed_update["stage"], "copy_templates")
        self.assertEqual(failed_update["partial_artifacts"], ["petition.yml"])
        self.assertEqual(failed_update["incomplete_artifacts"], ["petition.pdf"])
        self.assertEqual(
            failed_update["result"],
            {
                "project": "SyntheticPetition",
                "partial_artifacts": ["petition.yml"],
                "incomplete_artifacts": ["petition.pdf"],
                "incomplete_stage": "copy_templates",
            },
        )

    def test_blank_project_uses_main_yml(self):
        with (
            patch.object(api_editor, "playground_write_yaml") as mock_write,
            patch.object(api_editor, "get_list_of_projects", return_value=[]),
            patch.object(
                api_editor, "next_available_project_name", return_value="Blank"
            ),
            patch.object(api_editor, "create_project"),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/new-project",
                json={"project_name": "Blank", "create_test": False},
            ):
                response = api_editor._new_project_from_template(7, "req-1")
        self.assertEqual(response.get_json()["data"]["filename"], "main.yml")
        self.assertEqual(mock_write.call_args.args[2], "main.yml")
        starter_yaml = mock_write.call_args.args[3]
        self.assertIn("blank_project_complete", starter_yaml)
        self.assertIn("event: blank_project_complete", starter_yaml)
        self.assertIn("question: Interview ready", starter_yaml)

    def test_blank_project_creates_a_default_test_unless_disabled(self):
        with (
            patch.object(api_editor, "playground_write_yaml"),
            patch.object(api_editor, "get_list_of_projects", return_value=[]),
            patch.object(
                api_editor, "next_available_project_name", return_value="Blank"
            ),
            patch.object(api_editor, "create_project"),
            patch.object(
                api_editor,
                "_write_default_kiln_test",
                return_value={"filename": "weaver_it_runs.feature"},
            ) as write_test,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/new-project", json={"project_name": "Blank"}
            ):
                response = api_editor._new_project_from_template(7, "req-1")

        self.assertEqual(
            response.get_json()["data"]["test_filename"],
            "weaver_it_runs.feature",
        )
        write_test.assert_called_once()

    def test_publishing_metadata_reaches_the_generator(self):
        pdf_path = Path(__file__).parent / "test/test_dropdown_fields.pdf"
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "get_list_of_projects", return_value=[]),
            patch.object(api_editor, "_editor_async_is_configured", return_value=True),
            patch.object(
                api_editor, "next_available_project_name", return_value="Meta"
            ),
            patch.object(api_editor, "get_default_github_owner", return_value=""),
            patch.object(api_editor, "create_project"),
            patch.object(api_editor, "_start_new_project_upload_job") as mock_start_job,
        ):
            mock_start_job.return_value = {
                "job_id": "job-1",
                "job_url": "/al/editor/api/new-project/jobs/job-1",
                "state": {"status": "queued"},
            }
            with api_editor.app.test_client() as client:
                with pdf_path.open("rb") as handle:
                    client.post(
                        "/al/editor/api/new-project",
                        data={
                            "project_name": "Meta",
                            "separate_main_order": "true",
                            "interview_title": "Petition to enforce the sanitary code",
                            "interview_short_title": "Sanitary code",
                            "interview_description": "Ask the court to inspect.",
                            "jurisdiction": "NAM-US-US+MA",
                            "landing_page_url": "https://example.org/sanitary",
                            "list_topics": "HO-00-00-00-00, HO-05-00-00-00",
                            "interview_filename": "sanitary_code.yml",
                            "default_state": "MA",
                            "files": (BytesIO(handle.read()), pdf_path.name),
                        },
                        content_type="multipart/form-data",
                    )

        kwargs = mock_start_job.call_args.kwargs
        overrides = kwargs["generation_options"]["interview_overrides"]
        self.assertTrue(kwargs["generation_options"]["separate_main_order"])
        self.assertEqual(kwargs["interview_filename"], "sanitary_code.yml")
        self.assertEqual(overrides["title"], "Petition to enforce the sanitary code")
        self.assertEqual(overrides["short_title"], "Sanitary code")
        self.assertEqual(overrides["description"], "Ask the court to inspect.")
        self.assertEqual(overrides["landing_page_url"], "https://example.org/sanitary")
        self.assertTrue(overrides["has_other_categories"])
        self.assertEqual(
            overrides["other_categories"], "HO-00-00-00-00, HO-05-00-00-00"
        )
        # An explicit jurisdiction is not overwritten by the default state.
        self.assertEqual(overrides["jurisdiction"], "NAM-US-US+MA")
        self.assertEqual(overrides["state"], "MA")


class TestEditorKilnTestApi(unittest.TestCase):
    def test_default_scope_follows_only_selected_entrypoint_and_keeps_it_last(self):
        contents = {
            "entry_a.yml": "include:\n  - shared.yml\n---\nid: end a\nquestion: Done A\n",
            "entry_b.yml": "include:\n  - shared.yml\n---\nid: end b\nquestion: Done B\n",
            "shared.yml": "include:\n  - leaf.yml\n---\nid: shared\nquestion: Shared\n",
            "leaf.yml": "id: leaf\nquestion: Leaf\n",
        }
        with patch.object(
            api_editor,
            "playground_read_yaml",
            side_effect=lambda uid, project, name: contents[name],
        ):
            selected = api_editor._kiln_entrypoint_files(7, "Housing", "entry_a.yml")
        self.assertEqual(selected, ["shared.yml", "leaf.yml", "entry_a.yml"])

    def test_list_returns_selectable_feature_files(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "_project_kiln_test_filenames",
                return_value=["main.feature", "short.feature"],
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/kiln-tests?project=Housing"
            ):
                response = api_editor.editor_api_kiln_tests()

        self.assertEqual(
            response.get_json()["data"]["tests"],
            ["main.feature", "short.feature"],
        )
        self.assertEqual(
            response.get_json()["data"]["managed_test_filename"],
            "weaver_it_runs.feature",
        )
        self.assertIsNone(response.get_json()["data"]["managed_accessibility_enabled"])

    def test_list_reports_the_managed_tests_accessibility_mode(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "_project_kiln_test_filenames",
                return_value=["weaver_it_runs.feature"],
            ),
            patch.object(
                api_editor,
                "_read_project_text_file",
                return_value="And I check all pages for accessibility issues",
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/kiln-tests?project=Housing"
            ):
                response = api_editor.editor_api_kiln_tests()
        self.assertTrue(response.get_json()["data"]["managed_accessibility_enabled"])

    def test_draft_syncs_the_selected_test_against_project_yaml(self):
        synced = {
            "proposed_feature_text": "Feature: synced",
            "diff": "+Feature: synced",
            "added_screens": ["new screen"],
            "removed_screens": [],
            "added_functionality": ["new_value"],
            "removed_functionality": [],
        }
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "_read_project_text_file",
                return_value="Feature: existing",
            ),
            patch.object(
                api_editor,
                "_project_kiln_test_filenames",
                return_value=["weaver_it_runs.feature"],
            ),
            patch.object(
                api_editor, "_project_interview_yaml", return_value="question: New"
            ) as project_yaml,
            patch.object(api_editor, "sync_kiln_feature", return_value=synced) as sync,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/kiln-test/draft",
                method="POST",
                json={
                    "project": "Housing",
                    "interview_filename": "main.yml",
                    "test_filename": "weaver_it_runs.feature",
                    "accessibility": False,
                    "yaml_filenames": ["shared.yml"],
                },
            ):
                response = api_editor.editor_api_draft_kiln_test()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["data"]["removed_screens"], [])
        project_yaml.assert_called_once_with(7, "Housing", ["shared.yml", "main.yml"])
        sync.assert_called_once_with(
            "Feature: existing",
            "question: New",
            interview_filename="main.yml",
            accessibility_enabled=False,
        )

    def test_fixture_analysis_reads_only_selected_yaml_files(self):
        contents = {
            "main.yml": "id: main",
            "shared.yml": "id: shared",
            "other.yml": "id: other",
        }
        with (
            patch.object(
                api_editor,
                "_project_yaml_filenames",
                return_value=list(contents),
            ),
            patch.object(
                api_editor,
                "playground_read_yaml",
                side_effect=lambda _uid, _project, filename: contents[filename],
            ) as read,
        ):
            combined = api_editor._project_interview_yaml(
                7, "Housing", ["main.yml", "shared.yml"]
            )
        self.assertEqual(combined, "id: main\n---\nid: shared")
        self.assertEqual(
            [call.args[2] for call in read.call_args_list],
            ["main.yml", "shared.yml"],
        )

    def test_apply_saves_to_the_sources_area(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "_project_kiln_test_filenames",
                return_value=["weaver_it_runs.feature"],
            ),
            patch.object(
                api_editor, "_read_project_text_file", return_value="Feature: old\n"
            ),
            patch.object(api_editor, "_write_project_text_file") as write,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/kiln-test/apply",
                method="POST",
                json={
                    "project": "Housing",
                    "test_filename": "weaver_it_runs.feature",
                    "mode": "it_runs",
                    "content": "Feature: synced\n",
                    "expected_revision": api_editor.source_revision("Feature: old\n"),
                },
            ):
                response = api_editor.editor_api_apply_kiln_test()

        self.assertEqual(response.status_code, 200)
        write.assert_called_once_with(
            7,
            "Housing",
            "data",
            "weaver_it_runs.feature",
            "Feature: synced\n",
        )

    def test_json_draft_creates_a_new_recorded_path(self):
        generated = {"feature_text": "Feature: recorded", "rows": ["| answer | 42 |"]}
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "_project_kiln_test_filenames", return_value=[]),
            patch.object(
                api_editor,
                "create_kiln_feature_from_json",
                return_value=generated,
            ) as create,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/kiln-test/draft",
                method="POST",
                json={
                    "project": "Housing",
                    "interview_filename": "main.yml",
                    "mode": "json",
                    "test_filename": "happy_path.feature",
                    "question_id": "done",
                    "json_text": '{"variables":{"answer":42}}',
                },
            ):
                response = api_editor.editor_api_draft_kiln_test()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["data"]["mode"], "json")
        create.assert_called_once_with(
            '{"variables":{"answer":42}}',
            interview_filename="main.yml",
            question_id="done",
            accessibility_enabled=True,
        )

    def test_json_apply_refuses_to_overwrite_an_existing_test(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "_project_kiln_test_filenames",
                return_value=["happy_path.feature"],
            ),
            patch.object(
                api_editor, "_read_project_text_file", return_value="Feature: old\n"
            ),
            patch.object(api_editor, "_write_project_text_file") as write,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/kiln-test/apply",
                method="POST",
                json={
                    "project": "Housing",
                    "test_filename": "happy_path.feature",
                    "mode": "json",
                    "content": "Feature: recorded\n",
                },
            ):
                response = api_editor.editor_api_apply_kiln_test()

        self.assertEqual(response.status_code, 400)
        write.assert_not_called()

    def test_managed_sync_rejects_stale_draft_and_allows_exact_retry(self):
        from . import editor_utils as real_editor_utils

        for current, expected_status in [
            ("Feature: another author\n", 409),
            ("Feature: candidate\n", 200),
        ]:
            with (
                patch.object(
                    api_editor,
                    "source_revision",
                    side_effect=real_editor_utils.source_revision,
                ),
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(
                    api_editor,
                    "_project_kiln_test_filenames",
                    return_value=["weaver_it_runs.feature"],
                ),
                patch.object(
                    api_editor, "_read_project_text_file", return_value=current
                ),
                patch.object(api_editor, "_write_project_text_file") as write,
                api_editor.app.test_request_context(
                    "/al/editor/api/kiln-test/apply",
                    method="POST",
                    json={
                        "project": "Housing",
                        "test_filename": "weaver_it_runs.feature",
                        "mode": "it_runs",
                        "content": "Feature: candidate\n",
                        "expected_revision": api_editor.source_revision(
                            "Feature: original\n"
                        ),
                    },
                ),
            ):
                response = api_editor.editor_api_apply_kiln_test()
            self.assertEqual(response.status_code, expected_status)
            write.assert_not_called()


class TestEditorNewProjectMultipleUploads(unittest.TestCase):
    """Every uploaded document is woven into the one generated interview."""

    def _run(self, uploaded_files, generator_result=None, **job_kwargs):
        payload = {
            "yaml_text": "metadata:\n  title: Filing\n",
            "yaml_filename": "filing.yml",
            "input_filename": uploaded_files[0]["filename"],
            "generated_template_files": [],
        }
        payload.update(generator_result or {})
        written = {}
        with (
            patch.object(api_editor, "_update_new_project_job_state"),
            patch.object(api_editor, "playground_write_yaml"),
            patch.object(api_editor, "_copy_files_to_section") as mock_copy,
            patch.object(
                api_editor, "generate_interview_from_bytes", return_value=payload
            ) as mock_generate,
        ):
            mock_copy.side_effect = lambda **kwargs: written.update(
                {
                    os.path.basename(path): Path(path).read_bytes()
                    for path in kwargs["files"]
                }
            )
            result = api_editor._complete_new_project_upload_job(
                job_id="job-1",
                uid=7,
                project_name="Filing",
                request_id="req-1",
                uploaded_files=uploaded_files,
                generation_options={},
                debug_requested=False,
                **job_kwargs,
            )
        return result, mock_generate.call_args.kwargs, written

    def test_the_companion_documents_reach_the_generator(self):
        result, generate_kwargs, written = self._run(
            [
                {
                    "filename": "petition.pdf",
                    "content_bytes": b"%PDF-petition",
                    "mimetype": "application/pdf",
                },
                {
                    "filename": "affidavit.pdf",
                    "content_bytes": b"%PDF-affidavit",
                    "mimetype": "application/pdf",
                },
            ],
            generator_result={"template_filenames": ["petition.pdf", "affidavit.pdf"]},
        )

        self.assertEqual(generate_kwargs["filename"], "petition.pdf")
        self.assertEqual(
            [
                document["filename"]
                for document in generate_kwargs["additional_documents"]
            ],
            ["affidavit.pdf"],
        )
        self.assertEqual(
            generate_kwargs["additional_documents"][0]["content_bytes"],
            b"%PDF-affidavit",
        )
        self.assertEqual(result["woven_templates"], ["petition.pdf", "affidavit.pdf"])
        self.assertEqual(sorted(written), ["affidavit.pdf", "petition.pdf"])

    def test_generation_warnings_are_kept_in_the_completed_job_result(self):
        warning = (
            "The field `shared_answer` appears in multiple templates with "
            "different inferred types. Review the field type."
        )
        result, _generate_kwargs, _written = self._run(
            [
                {
                    "filename": "petition.pdf",
                    "content_bytes": b"%PDF-petition",
                    "mimetype": "application/pdf",
                },
                {
                    "filename": "affidavit.pdf",
                    "content_bytes": b"%PDF-affidavit",
                    "mimetype": "application/pdf",
                },
            ],
            generator_result={"warnings": [warning]},
        )

        self.assertEqual(result["warnings"], [warning])

    def test_the_project_stores_the_names_the_yaml_refers_to(self):
        """Two uploads sharing a name are told apart by the generator."""
        _result, _generate_kwargs, written = self._run(
            [
                {
                    "filename": "form.pdf",
                    "content_bytes": b"%PDF-first",
                    "mimetype": "application/pdf",
                },
                {
                    "filename": "form.pdf",
                    "content_bytes": b"%PDF-second",
                    "mimetype": "application/pdf",
                },
            ],
            generator_result={"template_filenames": ["form.pdf", "form_2.pdf"]},
        )
        self.assertEqual(written["form.pdf"], b"%PDF-first")
        self.assertEqual(written["form_2.pdf"], b"%PDF-second")

    def test_a_renamed_template_replaces_the_original_in_the_project(self):
        """The YAML names fields that only exist in the rewritten file."""
        result, _generate_kwargs, written = self._run(
            [
                {
                    "filename": "petition.pdf",
                    "content_bytes": b"%PDF-original",
                    "mimetype": "application/pdf",
                },
                {
                    "filename": "affidavit.pdf",
                    "content_bytes": b"%PDF-untouched",
                    "mimetype": "application/pdf",
                },
            ],
            generator_result={
                "template_filenames": ["petition.pdf", "affidavit.pdf"],
                "normalized_template_files": [
                    {"filename": "petition.pdf", "content_bytes": b"%PDF-renamed"}
                ],
            },
        )
        self.assertEqual(written["petition.pdf"], b"%PDF-renamed")
        self.assertEqual(written["affidavit.pdf"], b"%PDF-untouched")
        self.assertEqual(result["renamed_template_count"], 1)


INTERVIEW_WITH_TWO_DOCUMENTS = """---
objects:
  - petition: ALDocument.using(filename="petition", enabled=True)
  - affidavit: ALDocument.using(filename="affidavit", enabled=True)
---
objects:
  - al_user_bundle: ALDocumentBundle.using(elements=[petition, affidavit], filename="p", enabled=True)
---
attachment:
  name: Petition
  variable name: petition[i]
  pdf template file: petition.pdf
---
attachment:
  name: Affidavit
  variable name: affidavit[i]
  pdf template file: affidavit.pdf
"""


class TestEditorTemplateAnalysisApi(unittest.TestCase):
    """Analyzing a template that was added to a project after it was created."""

    def _post(self, payload, handler):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/template/import", method="POST", json=payload
            ):
                return handler()

    def test_a_template_that_is_not_in_the_project_is_a_404(self):
        with patch.object(
            api_editor,
            "_template_import_target",
            side_effect=FileNotFoundError("nope.pdf is not in this project"),
        ):
            response = self._post(
                {"project": "Eviction", "filename": "main.yml", "template": "nope.pdf"},
                api_editor.editor_api_import_template,
            )
        self.assertEqual(response.status_code, 404)

    def test_a_queued_analysis_reports_where_to_poll(self):
        with (
            patch.object(
                api_editor,
                "_template_import_target",
                return_value=("/tmp/a.pdf", "affidavit.pdf"),
            ),
            patch.object(api_editor, "playground_read_yaml", return_value="---\n"),
            patch.object(api_editor, "_editor_async_is_configured", return_value=True),
            patch.object(api_editor, "_store_job_state"),
            patch.object(api_editor, "_update_job_state"),
            patch.object(
                api_editor.workerapp,
                "send_task",
                return_value=types.SimpleNamespace(id="celery-1"),
            ) as mock_send,
        ):
            response = self._post(
                {
                    "project": "Eviction",
                    "filename": "main.yml",
                    "template": "affidavit.pdf",
                },
                api_editor.editor_api_import_template,
            )
        self.assertEqual(response.status_code, 202)
        body = response.get_json()
        self.assertIn("/al/editor/api/template/import/jobs/", body["job_url"])
        self.assertEqual(
            mock_send.call_args.kwargs["kwargs"]["template_filename"], "affidavit.pdf"
        )
        self.assertEqual(mock_send.call_args.kwargs["task_id"], body["job_id"])

    def test_applying_against_a_changed_interview_is_a_conflict(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor, "playground_read_yaml", return_value="---\nobjects: {}\n"
            ),
            patch.object(api_editor, "source_revision", return_value="now"),
            patch.object(
                api_editor,
                "insert_block_in_yaml",
                side_effect=lambda content, *_args, **_kwargs: content
                + "\n# candidate",
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/template/apply",
                method="POST",
                json={
                    "project": "Eviction",
                    "filename": "main.yml",
                    "expected_revision": "then",
                    "blocks": ["objects:\n  - affidavit: ALDocument.using()"],
                },
            ):
                response = api_editor.editor_api_apply_template_analysis()
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()["error"]["code"], "revision_conflict")


class TestEditorApplyBlockIds(unittest.TestCase):
    """One colliding screen id must not throw away the whole apply."""

    def test_a_colliding_id_is_numbered_rather_than_refused(self):
        taken = {"user name", "user name 2"}
        block, block_id = api_editor._block_id_without_collision(
            "id: user name\nquestion: |\n  What is your name?\n", taken
        )
        self.assertEqual(block_id, "user name 3")
        self.assertIn("id: user name 3\n", block)
        self.assertIn("question: |\n", block)

    def test_an_id_nothing_else_uses_is_left_exactly_as_it_is(self):
        block, block_id = api_editor._block_id_without_collision(
            "id: rent\nquestion: |\n  Rent\n", {"user name"}
        )
        self.assertEqual(block_id, "rent")
        self.assertEqual(block, "id: rent\nquestion: |\n  Rent\n")

    def test_a_block_with_no_id_is_left_alone(self):
        block, block_id = api_editor._block_id_without_collision(
            "objects:\n  - affidavit: ALDocument.using()\n", {"user name"}
        )
        self.assertIsNone(block_id)
        self.assertEqual(block, "objects:\n  - affidavit: ALDocument.using()\n")


class TestEditorApplyBlockReplacement(unittest.TestCase):
    """Re-reading a revised form rewrites its attachment block in place."""

    def _apply(self, blocks):
        written = {}
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "playground_read_yaml",
                return_value=INTERVIEW_WITH_TWO_DOCUMENTS,
            ),
            patch.object(api_editor, "playground_write_yaml") as mock_write,
            patch.object(api_editor, "update_block_in_yaml") as mock_update,
            patch.object(
                api_editor,
                "parse_interview_yaml",
                return_value={
                    "blocks": [],
                    "metadata_blocks": [],
                    "include_blocks": [],
                    "default_screen_parts_blocks": [],
                    "order_blocks": [],
                },
            ),
        ):
            mock_update.side_effect = lambda content, block_id, new_yaml: content
            mock_write.side_effect = (
                lambda uid, project, filename, content: written.update(
                    {"content": content}
                )
            )
            with api_editor.app.test_request_context(
                "/al/editor/api/template/apply",
                method="POST",
                json={
                    "project": "Eviction",
                    "filename": "main.yml",
                    "expected_revision": "test-revision",
                    "blocks": blocks,
                },
            ):
                response = api_editor.editor_api_apply_template_analysis()
        return response, mock_update

    def test_a_block_with_a_replace_target_rewrites_rather_than_adds(self):
        response, mock_update = self._apply(
            [
                {
                    "yaml": "attachment:\n  name: Petition\n",
                    "replace_block_id": "block-3",
                }
            ]
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(mock_update.call_args.args[1], "block-3")
        self.assertEqual(response.get_json()["data"]["replaced_block_ids"], ["block-3"])

    def test_a_plain_string_is_still_an_addition(self):
        response, mock_update = self._apply(
            ["objects:\n  - cover_sheet: ALDocument.using()\n"]
        )
        self.assertEqual(response.status_code, 200)
        mock_update.assert_not_called()


class TestEditorAttachmentMappingsApi(unittest.TestCase):
    SOURCE = 'id: output\nattachment:\n  pdf template file: form.pdf\n  fields:\n    name: "${ old }" # keep\n'

    def _request(
        self,
        payload,
        authenticated=True,
        template_directory="/tmp/missing-template-folder",
    ):
        from . import editor_utils as real_utils

        with (
            patch.object(api_editor, "_editor_auth_check", return_value=authenticated),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "playground_read_yaml", return_value=self.SOURCE),
            patch.object(
                api_editor,
                "parse_interview_yaml",
                side_effect=real_utils.parse_interview_yaml,
            ),
            patch.object(
                api_editor,
                "update_block_in_yaml",
                side_effect=real_utils.update_block_in_yaml,
            ),
            patch.object(api_editor, "playground_write_yaml") as write,
            patch.object(
                api_editor,
                "_editor_storage_directory",
                return_value=(None, template_directory),
            ),
            api_editor.app.test_request_context(
                "/al/editor/api/attachment-mappings",
                method="POST",
                json={
                    "project": "test",
                    "filename": "main.yml",
                    "block_id": "output",
                    **payload,
                },
            ),
        ):
            response = api_editor.editor_api_attachment_mappings()
        return response, write

    def test_save_patches_value_and_preserves_comment(self):
        response, write = self._request(
            {
                "expected_revision": "test-revision",
                "updates": [
                    {"index": 0, "values": {"name": "${ new if ready else '' }"}}
                ],
            }
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(write.call_count, 1)
        self.assertIn(" # keep", write.call_args.args[3])
        self.assertIn("new if ready", write.call_args.args[3])

    def test_stale_revision_and_invalid_updates_do_not_write(self):
        for payload, status in (
            ({"updates": [{"index": 0, "values": {"name": "new_value"}}]}, 400),
            (
                {
                    "expected_revision": "stale",
                    "updates": [{"index": 0, "values": {"name": "new_value"}}],
                },
                409,
            ),
            ({"expected_revision": "test-revision", "updates": ["invalid"]}, 400),
        ):
            response, write = self._request(payload)
            self.assertEqual(response.status_code, status)
            write.assert_not_called()

    def test_missing_template_still_allows_existing_mapping_edits(self):
        response, write = self._request({})
        self.assertEqual(response.status_code, 200)
        attachment = response.get_json()["data"]["attachments"][0]
        self.assertEqual(attachment["rows"][0]["name"], "name")
        self.assertIn("Could not check", attachment["warning"])
        write.assert_not_called()

    def test_requires_authentication(self):
        response, write = self._request({}, authenticated=False)
        self.assertIn(response.status_code, (401, 403))
        write.assert_not_called()


class TestEditorDocumentsApi(unittest.TestCase):
    """Rearranging the documents an interview assembles."""

    def test_it_lists_the_documents_and_their_bundle_order(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "playground_read_yaml",
                return_value=INTERVIEW_WITH_TWO_DOCUMENTS,
            ),
            patch.object(
                api_editor,
                "_list_editor_section_files",
                return_value=[
                    {"filename": "petition.pdf"},
                    {"filename": "affidavit.pdf"},
                    {"filename": "leftover.pdf"},
                ],
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/documents?project=Eviction&filename=main.yml"
            ):
                response = api_editor.editor_api_documents()
        data = response.get_json()["data"]
        self.assertEqual(
            [document["name"] for document in data["documents"]],
            ["petition", "affidavit"],
        )
        self.assertEqual(data["bundles"][0]["elements"], ["petition", "affidavit"])

    def test_it_says_which_template_files_are_not_imported_yet(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "playground_read_yaml",
                return_value=INTERVIEW_WITH_TWO_DOCUMENTS
                + "---\nquestion: |\n  Hi\nsubquestion: |\n  See logo.png\n",
            ),
            patch.object(
                api_editor,
                "_list_editor_section_files",
                return_value=[
                    {"filename": "petition.pdf"},
                    {"filename": "logo.png"},
                    {"filename": "leftover.pdf"},
                ],
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/documents?project=Eviction&filename=main.yml"
            ):
                response = api_editor.editor_api_documents()
        templates = response.get_json()["data"]["templates"]
        self.assertEqual(templates["petition.pdf"]["status"], "attached")
        self.assertEqual(templates["petition.pdf"]["document"], "petition")
        self.assertEqual(templates["logo.png"]["status"], "referenced")
        self.assertEqual(templates["leftover.pdf"]["status"], "not_imported")

    def _save(self, payload):
        written = {}
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "playground_read_yaml",
                return_value=INTERVIEW_WITH_TWO_DOCUMENTS,
            ),
            patch.object(api_editor, "playground_write_yaml") as mock_write,
        ):
            mock_write.side_effect = (
                lambda uid, project, filename, content: written.update(
                    {"content": content}
                )
            )
            with api_editor.app.test_request_context(
                "/al/editor/api/documents", method="POST", json=payload
            ):
                response = api_editor.editor_api_save_documents()
        return response, written.get("content", "")

    def test_removal_preview_and_apply_report_custom_cross_file_references(self):
        original = (
            INTERVIEW_WITH_TWO_DOCUMENTS
            + "---\nid: custom\ncode: |\n  title = petition.title\n"
        )
        related = "code: |\n  published = petition.as_pdf()\n"
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "_project_yaml_filenames",
                return_value=["main.yml", "related.yml"],
            ),
            patch.object(
                api_editor,
                "playground_read_yaml",
                side_effect=lambda uid, project, filename: (
                    original if filename == "main.yml" else related
                ),
            ),
            patch.object(api_editor, "playground_write_yaml") as write,
        ):
            for preview, status in [(True, 200), (False, 409)]:
                with api_editor.app.test_request_context(
                    "/al/editor/api/documents",
                    method="POST",
                    json={
                        "project": "Housing",
                        "filename": "main.yml",
                        "remove": ["petition"],
                        "expected_revision": "test-revision",
                        "preview": preview,
                    },
                ):
                    response = api_editor.editor_api_save_documents()
                self.assertEqual(response.status_code, status)
                body = response.get_json()
                plan = body["data"] if preview else body["error"]["details"]
                self.assertTrue(plan["blocked"])
                self.assertEqual(
                    {item["filename"] for item in plan["references"]},
                    {"main.yml", "related.yml"},
                )
                self.assertIn("-  - petition:", plan["diff"])
            write.assert_not_called()

    def test_reordering_a_bundle_is_written_back(self):
        response, content = self._save(
            {
                "project": "Eviction",
                "filename": "main.yml",
                "expected_revision": "test-revision",
                "bundles": [
                    {
                        "bundle": "al_user_bundle",
                        "elements": ["affidavit", "petition"],
                    }
                ],
            }
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("elements=[affidavit, petition]", content)

    def test_removing_last_document_saves_an_empty_bundle(self):
        response, content = self._save(
            {
                "project": "Eviction",
                "filename": "main.yml",
                "expected_revision": "test-revision",
                "bundles": [{"bundle": "al_user_bundle", "elements": []}],
            }
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("elements=[]", content)
        self.assertIn("pdf template file: petition.pdf", content)

    def test_deleting_document_cleans_its_attachment_and_bundle_entries(self):
        response, content = self._save(
            {
                "project": "Eviction",
                "filename": "main.yml",
                "expected_revision": "test-revision",
                "remove": ["petition"],
            }
        )
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("pdf template file: petition.pdf", content)
        self.assertNotIn("- petition: ALDocument", content)
        self.assertIn("elements=[affidavit]", content)

    def test_deletion_with_stale_revision_does_not_write(self):
        response, content = self._save(
            {
                "project": "Eviction",
                "filename": "main.yml",
                "expected_revision": "old",
                "remove": ["petition"],
            }
        )
        self.assertEqual(response.status_code, 409)
        self.assertEqual(content, "")

    def test_an_enabled_rule_is_written_into_the_declaration(self):
        response, content = self._save(
            {
                "project": "Eviction",
                "filename": "main.yml",
                "expected_revision": "test-revision",
                "enabled": [{"name": "affidavit", "expression": "user_is_low_income"}],
            }
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("enabled=user_is_low_income", content)
        self.assertIn(
            '- petition: ALDocument.using(filename="petition", enabled=True)', content
        )

    def test_a_rule_that_is_not_an_expression_is_refused(self):
        response, content = self._save(
            {
                "project": "Eviction",
                "filename": "main.yml",
                "expected_revision": "test-revision",
                "enabled": [{"name": "affidavit", "expression": "if x: y"}],
            }
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(content, "")

    def test_a_stale_revision_is_a_conflict(self):
        response, content = self._save(
            {
                "project": "Eviction",
                "filename": "main.yml",
                "expected_revision": "stale",
                "bundles": [{"bundle": "al_user_bundle", "elements": ["affidavit"]}],
            }
        )
        self.assertEqual(response.status_code, 409)
        self.assertEqual(content, "")


class TestEditorStyleCheckSeverity(unittest.TestCase):
    def test_style_findings_never_report_as_errors(self):
        demoted = api_editor._demote_style_findings(
            [
                {
                    "rule_id": "missing-metadata-fields",
                    "level": "error",
                    "severity": "error",
                    "message": "x",
                },
                {
                    "rule_id": "vague-button",
                    "level": "warning",
                    "severity": "warning",
                    "message": "y",
                },
                {
                    "rule_id": "prefer-person-objects",
                    "level": "info",
                    "severity": "info",
                    "message": "z",
                },
            ]
        )
        self.assertEqual(
            [item["level"] for item in demoted], ["warning", "warning", "info"]
        )
        self.assertEqual(
            [item["severity"] for item in demoted], ["warning", "warning", "info"]
        )
        self.assertEqual(demoted[0]["style_original_level"], "error")
        self.assertNotIn("style_original_level", demoted[1])


class TestEditorListTopics(unittest.TestCase):
    """The picker offers the same codes the question-driven Weaver offers."""

    def test_taxonomy_is_grouped_with_the_heading_code_first(self):
        groups = api_editor._list_topic_groups()
        self.assertTrue(groups)

        by_label = {group["label"]: group for group in groups}
        self.assertIn("Housing", by_label)
        housing = by_label["Housing"]
        # A group leads with its own broadest code.
        self.assertEqual(housing["topics"][0]["code"], "HO-00-00-00-00")
        self.assertTrue(housing["topics"][0]["heading"])
        codes = [topic["code"] for topic in housing["topics"]]
        self.assertIn("HO-05-00-00-00", codes)
        self.assertEqual(len(codes), len(set(codes)))
        for topic in housing["topics"]:
            self.assertTrue(topic["label"])

        # Relevance order, the same custom order get_LIST_codes applies.
        self.assertEqual(groups[0]["label"], "Housing")

    def test_endpoint_returns_the_groups_to_an_authenticated_developer(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            api_editor.app.test_request_context("/al/editor/api/list-topics"),
        ):
            payload = api_editor.editor_api_list_topics().get_json()

        self.assertTrue(payload["success"])
        self.assertTrue(payload["data"]["groups"])
        self.assertEqual(payload["data"]["docs_url"], "https://taxonomy.legal")

    def test_endpoint_refuses_an_unauthenticated_request(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=False),
            api_editor.app.test_request_context("/al/editor/api/list-topics"),
        ):
            response = api_editor.editor_api_list_topics()

        status = response[1] if isinstance(response, tuple) else response.status_code
        self.assertEqual(status, 401)


class TestEditorQuestionLibraryApi(unittest.TestCase):
    """The AssemblyLine question library, reachable after project creation.

    The Weaver copies these questions in only while it writes a new interview.
    An object declared later needs them too, so the editor offers them for
    whatever the file being edited declares now.
    """

    SOURCE = (
        "objects:\n"
        "  - users: ALPeopleList.using(there_are_any=True)\n"
        "  - children: ALPeopleList.using(ask_number=True)\n"
        "  - landlord: ALIndividual\n"
        "  - al_court_bundle: ALDocumentBundle.using(elements=[])\n"
        "---\n"
        "id: users names\n"
        "question: |\n"
        "  Who are you?\n"
        "fields:\n"
        "  - code: |\n"
        "      users[i].name_fields()\n"
        "---\n"
        "id: birthday\n"
        "question: |\n"
        "  When were you born?\n"
        "fields:\n"
        "  - Birthdate: users[i].birthdate\n"
    )

    def _real_editor_utils(self):
        from . import editor_utils as real_editor_utils

        return real_editor_utils

    def _patches(self, source):
        real_editor_utils = self._real_editor_utils()
        return [
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "playground_read_yaml", return_value=source),
            patch.object(
                api_editor,
                "parse_interview_yaml",
                side_effect=real_editor_utils.parse_interview_yaml,
            ),
            patch.object(
                api_editor,
                "insert_block_in_yaml",
                side_effect=real_editor_utils.insert_block_in_yaml,
            ),
            patch.object(
                api_editor,
                "add_object_declaration",
                side_effect=real_editor_utils.add_object_declaration,
            ),
            patch.object(
                api_editor,
                "source_revision",
                side_effect=real_editor_utils.source_revision,
            ),
        ]

    def _insert(self, questions, insert_after_id=None, source=None):
        written = {}

        def record_write(uid, project, filename, content):
            written["content"] = content

        with ExitStack() as stack:
            for patcher in self._patches(source if source is not None else self.SOURCE):
                stack.enter_context(patcher)
            stack.enter_context(
                patch.object(
                    api_editor, "playground_write_yaml", side_effect=record_write
                )
            )
            stack.enter_context(
                api_editor.app.test_request_context(
                    "/al/editor/api/question-library/insert",
                    method="POST",
                    json={
                        "project": "default",
                        "filename": "test.yml",
                        "insert_after_id": insert_after_id,
                        "questions": questions,
                    },
                )
            )
            response = api_editor.editor_api_question_library_insert()
        return response, written.get("content")

    def test_inserting_writes_the_blocks_the_weaver_would_have_written(self):
        response, written = self._insert(
            [
                {"var": "children", "kind": "birthdate"},
                {"var": "children", "kind": "how_many"},
            ],
            insert_after_id="birthday",
        )
        payload = response.get_json()
        self.assertTrue(payload["success"])
        self.assertEqual(
            payload["data"]["inserted_block_ids"],
            ["how many children", "child birthdate"],
        )
        self.assertIn("id: how many children", written)
        self.assertIn("children[i].birthdate", written)
        # Inserted after the anchor block, in the order the catalog lists them:
        # the gather flow before the questions about each child.
        self.assertLess(written.index("id: birthday"), written.index("id: how many"))
        self.assertLess(
            written.index("id: how many children"), written.index("id: child birthdate")
        )
        # Nothing that was already in the file moved or changed.
        self.assertIn("id: users names", written)

    def test_a_write_answers_with_the_revision_the_file_now_has(self):
        # Without it the editor keeps the revision from before its own write,
        # and the next metadata save is rejected as a conflict.
        real_editor_utils = self._real_editor_utils()
        response, written = self._insert([{"var": "children", "kind": "names"}])
        self.assertEqual(
            response.get_json()["data"]["revision"],
            real_editor_utils.source_revision(written),
        )
        response, written = self._declare(name="witnesses", class_name="ALPeopleList")
        self.assertEqual(
            response.get_json()["data"]["revision"],
            real_editor_utils.source_revision(written),
        )

    def test_a_question_the_file_already_has_is_not_added_a_second_time(self):
        response, written = self._insert(
            [
                {"var": "users", "kind": "names"},
                {"var": "users", "kind": "there_is_another"},
            ]
        )
        payload = response.get_json()
        self.assertEqual(payload["data"]["inserted_block_ids"], ["another user"])
        self.assertEqual(payload["data"]["skipped_block_ids"], ["users names"])
        self.assertEqual(written.count("id: users names"), 1)

    def test_the_same_question_asked_for_twice_is_added_once(self):
        response, written = self._insert(
            [
                {"var": "children", "kind": "names"},
                {"var": "children", "kind": "names"},
            ]
        )
        payload = response.get_json()
        self.assertEqual(payload["data"]["inserted_block_ids"], ["children names"])
        self.assertEqual(written.count("id: children names"), 1)

    def test_a_question_that_was_never_offered_is_refused(self):
        response, written = self._insert(
            [{"var": "al_court_bundle", "kind": "birthdate"}]
        )
        self.assertEqual(response.status_code, 400)
        self.assertIsNone(written)
        self.assertIn("al_court_bundle", response.get_json()["error"]["message"])

    def test_nothing_is_written_when_no_questions_are_asked_for(self):
        response, written = self._insert([])
        self.assertEqual(response.status_code, 400)
        self.assertIsNone(written)

    def _declare(self, source=None, **payload):
        written = {}

        def record_write(uid, project, filename, content):
            written["content"] = content

        body = {"project": "default", "filename": "test.yml"}
        body.update(payload)
        with ExitStack() as stack:
            for patcher in self._patches(source if source is not None else self.SOURCE):
                stack.enter_context(patcher)
            stack.enter_context(
                patch.object(
                    api_editor, "playground_write_yaml", side_effect=record_write
                )
            )
            stack.enter_context(
                api_editor.app.test_request_context(
                    "/al/editor/api/question-library/object", method="POST", json=body
                )
            )
            response = api_editor.editor_api_question_library_object()
        return response, written.get("content")

    def test_a_new_list_joins_the_block_that_already_declares_people(self):
        response, written = self._declare(
            name="witnesses",
            class_name="ALPeopleList",
            using_args="ask_number=True",
        )
        payload = response.get_json()
        self.assertTrue(payload["success"])
        self.assertEqual(
            payload["data"]["declared"],
            {"var": "witnesses", "expression": "ALPeopleList.using(ask_number=True)"},
        )
        self.assertIn(
            "  - witnesses: ALPeopleList.using(ask_number=True)",
            written,
        )
        # Added to the people block, not to a new one.
        self.assertEqual(written.count("objects:"), 1)
        # And its questions are on offer straight away.
        offered = {entry["var"] for entry in payload["data"]["objects"]}
        self.assertIn("witnesses", offered)

    def test_people_never_join_the_block_that_declares_the_documents(self):
        source = (
            "objects:\n"
            "  - al_court_bundle: ALDocumentBundle.using(elements=[])\n"
            "---\n"
            "id: q\n"
            "question: |\n"
            "  Hello\n"
        )
        response, written = self._declare(
            source=source,
            name="witnesses",
            class_name="ALPeopleList",
            save_as_draft=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("objects:\n  - witnesses: ALPeopleList", written)
        self.assertIn("al_court_bundle: ALDocumentBundle", written)
        # A block of their own, written before the questions.
        self.assertLess(written.index("witnesses"), written.index("id: q"))

    def test_a_people_block_a_line_cannot_join_gets_a_block_beside_it(self):
        # `objects: {users: ALPeopleList}` cannot take an indented line, and
        # rewriting it into another style is an edit nobody asked for.
        source = "objects: {users: ALPeopleList}\n---\nid: q\nquestion: |\n  Hello\n"
        response, written = self._declare(
            source=source,
            name="witnesses",
            class_name="ALPeopleList",
            save_as_draft=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("objects: {users: ALPeopleList}", written)
        self.assertIn("objects:\n  - witnesses: ALPeopleList", written)
        self.assertLess(written.index("witnesses"), written.index("id: q"))

    def test_the_quantity_choice_becomes_the_using_call(self):
        for using_args, expected in (
            ("", "ALPeopleList"),
            ("there_are_any=True", "ALPeopleList.using(there_are_any=True)"),
            ("ask_number=True", "ALPeopleList.using(ask_number=True)"),
            (
                "ask_number=True, target_number=2",
                "ALPeopleList.using(ask_number=True, target_number=2)",
            ),
        ):
            with self.subTest(using_args=using_args):
                response, _written = self._declare(
                    name="witnesses", class_name="ALPeopleList", using_args=using_args
                )
                self.assertEqual(
                    response.get_json()["data"]["declared"]["expression"], expected
                )

    def test_only_the_how_many_parameters_can_be_written(self):
        # An `objects:` entry is a Python expression the interview evaluates, so
        # what the browser sends is a quantity choice, never source to pass on.
        for using_args in (
            'filename="x"',
            "there_are_any=os.system('rm -rf /')",
            "True",
            "target_number=-1",
            "there_are_any=1",
        ):
            with self.subTest(using_args=using_args):
                response, written = self._declare(
                    name="witnesses", class_name="ALPeopleList", using_args=using_args
                )
                self.assertEqual(response.status_code, 400)
                self.assertIsNone(written)

    def test_only_the_classes_the_library_has_questions_for_are_declared(self):
        for class_name in ("ALDocumentBundle", "DAList", "", "ALCourt"):
            with self.subTest(class_name=class_name):
                response, written = self._declare(
                    name="witnesses", class_name=class_name
                )
                self.assertEqual(response.status_code, 400)
                self.assertIsNone(written)

    def test_a_name_the_file_already_uses_is_refused(self):
        response, written = self._declare(name="children", class_name="ALPeopleList")
        self.assertEqual(response.status_code, 400)
        self.assertIsNone(written)
        self.assertIn("already declared", response.get_json()["error"]["message"])

    def test_a_name_assembly_line_manages_itself_is_refused(self):
        # `plaintiffs` is derived from `users` and `other_parties`; declaring it
        # here would clobber that.
        response, written = self._declare(name="plaintiffs", class_name="ALPeopleList")
        self.assertEqual(response.status_code, 400)
        self.assertIsNone(written)
        self.assertIn("AssemblyLine", response.get_json()["error"]["message"])

    def test_a_name_that_is_not_a_variable_name_is_refused(self):
        for name in ("my witnesses", "2witnesses", "class", "", "witnesses.name"):
            with self.subTest(name=name):
                response, written = self._declare(name=name, class_name="ALPeopleList")
                self.assertEqual(response.status_code, 400)
                self.assertIsNone(written)

    def test_declaring_one_person_takes_no_quantity(self):
        response, _written = self._declare(name="landlord2", class_name="ALIndividual")
        self.assertEqual(
            response.get_json()["data"]["declared"]["expression"], "ALIndividual"
        )
        response, written = self._declare(
            name="landlord2", class_name="ALIndividual", using_args="ask_number=True"
        )
        self.assertEqual(response.status_code, 400)
        self.assertIsNone(written)

    def test_every_endpoint_refuses_an_unauthenticated_request(self):
        endpoints = (
            (api_editor.editor_api_question_library, "/al/editor/api/question-library"),
            (
                api_editor.editor_api_question_library_insert,
                "/al/editor/api/question-library/insert",
            ),
            (
                api_editor.editor_api_question_library_object,
                "/al/editor/api/question-library/object",
            ),
        )
        for view, path in endpoints:
            with self.subTest(path=path):
                with (
                    patch.object(api_editor, "_editor_auth_check", return_value=False),
                    api_editor.app.test_request_context(path, method="POST", json={}),
                ):
                    response = view()
                status = (
                    response[1] if isinstance(response, tuple) else response.status_code
                )
                self.assertEqual(status, 401)


class TestEditorBlockPayloadValidation(unittest.TestCase):
    """What the "Add a block" modal hands the API has to be accepted."""

    def accepts(self, block_yaml):
        api_editor._validate_block_yaml_payload(block_yaml)

    def test_a_standalone_comment_block_is_a_real_block(self):
        # Prose about the interview. docassemble reads it, the checker passes
        # it, and the modal offers it — so the API cannot refuse it.
        self.accepts("comment: |\n  Explain what the blocks below do.\n")

    def test_a_blank_new_block_of_only_yaml_comments_is_allowed(self):
        # What "Raw YAML block" inserts, before anything is typed over it.
        self.accepts("# replace with any docassemble YAML\n")
        self.accepts("# one\n\n# two\n")

    def test_an_id_with_nothing_to_name_is_refused(self):
        for payload in ("id: c1\n", "id: c1\ncomment: |\n  Just prose.\n"):
            with self.subTest(payload=payload):
                with self.assertRaisesRegex(ValueError, "incomplete"):
                    self.accepts(payload)

    def test_a_document_that_is_not_a_block_is_still_refused(self):
        for payload in ("- one\n- two\n", "just a string\n", "{}\n"):
            with self.subTest(payload=payload):
                with self.assertRaises(ValueError):
                    self.accepts(payload)

    def test_text_templates_require_a_safe_name_and_nonblank_content(self):
        self.accepts(
            "template: mailing_help\nsubject: |\n  Learn more\ncontent: |\n  Hi ${ user }\n"
        )
        for payload in (
            "template: 2_bad\ncontent: Text\n",
            "template: class\ncontent: Text\n",
            "template: blank_help\ncontent: '   '\n",
        ):
            with self.subTest(payload=payload):
                with self.assertRaises(ValueError):
                    self.accepts(payload)

    def test_new_template_variants_can_share_a_name(self):
        from .editor_utils import parse_interview_yaml

        source = "id: english\ntemplate: shared_help\nlanguage: en\ncontent: English\n"
        with patch.object(
            api_editor, "parse_interview_yaml", wraps=parse_interview_yaml
        ):
            for variant in (
                "template: shared_help\nlanguage: es\ncontent: Spanish\n",
                "template: shared_help\nif: special_case\ncontent: Conditional\n",
            ):
                with self.subTest(variant=variant):
                    api_editor._validate_template_against_file(source, variant)

    def test_variant_can_be_renamed_or_removed_while_another_defines_its_name(self):
        from .editor_utils import parse_interview_yaml

        source = (
            "id: english\ntemplate: shared_help\nlanguage: en\ncontent: English\n---\n"
            "id: spanish\ntemplate: shared_help\nlanguage: es\ncontent: Spanish\n---\n"
            "id: other\ntemplate: other_help\ncontent: Other\n---\n"
            "question: Address\nsubquestion: ${ collapse_template(shared_help) }\n"
        )
        with patch.object(
            api_editor, "parse_interview_yaml", wraps=parse_interview_yaml
        ):
            self.assertEqual(
                api_editor._template_references(source, "shared_help", "spanish"), []
            )
            api_editor._validate_template_against_file(
                source,
                "template: other_help\ncontent: Edited variant\n",
                current_block_id="spanish",
            )
            # Without the other variant, the same change loses the definition.
            remaining = source.split("---\n", 1)[1]
            self.assertEqual(
                len(
                    api_editor._template_references(remaining, "shared_help", "spanish")
                ),
                1,
            )
            with self.assertRaisesRegex(ValueError, "Cannot rename"):
                api_editor._validate_template_against_file(
                    remaining,
                    "template: other_help\ncontent: Edited variant\n",
                    current_block_id="spanish",
                )

    def test_imported_code_names_are_reserved_for_non_template_bindings(self):
        from .editor_utils import parse_interview_yaml

        source = (
            "code: |\n  import math as imported_help\n  import os.path\n"
            "  from math import sqrt as root_help\n  from math import ceil\n"
            "  def local_scope():\n    import json as local_help\n"
        )
        with patch.object(
            api_editor, "parse_interview_yaml", wraps=parse_interview_yaml
        ):
            for name in ("imported_help", "os", "root_help", "ceil"):
                with self.subTest(name=name):
                    with self.assertRaisesRegex(ValueError, "non-template"):
                        api_editor._validate_template_against_file(
                            source, f"template: {name}\ncontent: Text\n"
                        )
            api_editor._validate_template_against_file(
                source, "template: local_help\ncontent: Text\n"
            )

    def test_template_references_report_the_question_and_line(self):
        blocks = [
            {
                "id": "address",
                "title": "What is your address?",
                "line_start": 14,
                "yaml": "subquestion: ${ collapse_template(shared_help) }\n",
                "data": {},
            },
            {
                "id": "other",
                "title": "Other",
                "line_start": 30,
                "yaml": "subquestion: No help here.\n",
                "data": {},
            },
        ]
        with patch.object(
            api_editor, "parse_interview_yaml", return_value={"blocks": blocks}
        ):
            self.assertEqual(
                api_editor._template_references("source", "shared_help"),
                ["What is your address? (line 14)"],
            )

    def test_advanced_template_names_remain_source_editable(self):
        self.accepts("template: person[i].help\ncontent: Advanced text\n")

    def test_template_names_do_not_collide_with_local_function_variables(self):
        from .editor_utils import parse_interview_yaml

        source = (
            "code: |\n  def helper():\n    local_help = 'local'\n"
            "  assigned_help, other = ('global', True)\n---\n"
            "question: Choice\nyesno: answer_help\n"
        )
        with patch.object(
            api_editor, "parse_interview_yaml", wraps=parse_interview_yaml
        ):
            api_editor._validate_template_against_file(
                source, "template: local_help\ncontent: Text\n"
            )
            for name in ("assigned_help", "answer_help", "helper"):
                with self.subTest(name=name):
                    with self.assertRaisesRegex(ValueError, "already defined"):
                        api_editor._validate_template_against_file(
                            source, f"template: {name}\ncontent: Text\n"
                        )

    def test_existing_template_language_variants_can_be_edited(self):
        from .editor_utils import parse_interview_yaml

        source = (
            "template: help_text\nlanguage: en\ncontent: English\n---\n"
            "template: help_text\nlanguage: es\ncontent: Spanish\n"
        )
        block_id = parse_interview_yaml(source)["blocks"][0]["id"]
        with patch.object(
            api_editor, "parse_interview_yaml", wraps=parse_interview_yaml
        ):
            api_editor._validate_template_against_file(
                source,
                "template: help_text\ncontent: Edited\n",
                current_block_id=block_id,
            )

    def test_content_file_templates_remain_source_editable(self):
        from .editor_utils import parse_interview_yaml

        for name in ("shared_help", "person[i].help"):
            source = f"template: {name}\nsubject: Learn more\ncontent file: help.md\n"
            self.accepts(source)
            self.assertEqual(
                parse_interview_yaml(source)["blocks"][0]["type"], "template"
            )

    def test_template_references_match_the_exact_argument(self):
        from .editor_utils import parse_interview_yaml

        source = (
            "id: exact\nquestion: Exact\nsubquestion: |\n"
            "  ${ collapse_template(shared_help, collapsed=False) }\n---\n"
            "id: attribute\nquestion: Attribute\nsubquestion: |\n"
            "  ${ collapse_template(shared_help.other) }\n---\n"
            "id: prefix\nquestion: Prefix\nsubquestion: |\n"
            "  ${ collapse_template(shared_help_longer) }\n"
        )
        with patch.object(
            api_editor, "parse_interview_yaml", wraps=parse_interview_yaml
        ):
            refs = api_editor._template_references(source, "shared_help")
        self.assertEqual(len(refs), 1)
        self.assertIn("Exact", refs[0])

    def test_referenced_template_cannot_be_renamed(self):
        from .editor_utils import parse_interview_yaml

        source = (
            "template: shared_help\nsubject: Label\ncontent: Text\n---\n"
            "id: address\nquestion: Address\n"
            "subquestion: ${ collapse_template(shared_help) }\n"
        )
        block_id = parse_interview_yaml(source)["blocks"][0]["id"]
        with patch.object(
            api_editor, "parse_interview_yaml", wraps=parse_interview_yaml
        ):
            with self.assertRaisesRegex(ValueError, "Cannot rename.*Address"):
                api_editor._validate_template_against_file(
                    source,
                    "template: renamed_help\ncontent: Text\n",
                    current_block_id=block_id,
                )
            api_editor._validate_template_against_file(
                source,
                "template: shared_help\ncontent: Edited\n",
                current_block_id=block_id,
            )

    def test_question_text_is_required_when_saving(self):
        for value in ('""', '"   "', "null", "[]"):
            with self.subTest(value=value):
                with self.assertRaisesRegex(ValueError, "Question text is required"):
                    self.accepts(f"id: q1\nquestion: {value}\nfields: []\n")
        self.accepts("id: q1\nquestion: What is your name?\nfields: []\n")

    def test_new_question_may_be_inserted_as_an_empty_draft(self):
        api_editor._validate_block_yaml_payload(
            'id: q1\nquestion: ""\nfields: []\n', allow_empty_question=True
        )

    def test_new_interview_starts_with_an_empty_question(self):
        self.assertIn('question: ""\n', api_editor._default_new_interview_yaml())


class TestOrderBlockLookup(unittest.TestCase):
    """`order_blocks` holds document indices, not positions in `blocks`."""

    SOURCE = (
        "---\nmetadata:\n  title: Example\n"
        "---\nid: interview_order_form\ncode: |\n  rent_amount\n  interview_order_form = True\n"
        "---\n---\nid: main\nmandatory: True\ncode: |\n  intro\n  interview_order_form\n  download\n"
        "---\nid: intro\nquestion: Hello\ncontinue button field: intro\n"
    )

    def test_an_included_named_order_needs_no_id_or_mandatory_key(self):
        from .editor_utils import parse_interview_yaml, parse_order_code

        source = (
            "---\ncomment: Child flow\ncode: |\n"
            "  first_screen\n"
            "  if has_children:\n"
            "    children.gather()\n"
            "  child_flow_done = True\n"
        )
        with patch.object(api_editor, "parse_order_code", parse_order_code):
            named = api_editor._named_order_steps_from_model(
                parse_interview_yaml(source)
            )
        self.assertEqual(set(named), {"child_flow_done"})
        self.assertEqual(named["child_flow_done"][0]["invoke"], "first_screen")

    def test_a_simple_assignment_is_not_mistaken_for_an_order(self):
        from .editor_utils import parse_interview_yaml

        source = "---\ncode: |\n  ordinary_flag = True\n"
        self.assertEqual(
            api_editor._named_order_steps_from_model(parse_interview_yaml(source)),
            {},
        )

    def _post_order_edit(self, path, payload):
        from . import editor_utils

        with ExitStack() as stack:
            for name in (
                "parse_interview_yaml",
                "parse_order_code",
                "serialize_order_steps",
                "validate_order_steps",
                "canonical_block_yaml",
                "update_block_in_yaml",
            ):
                stack.enter_context(
                    patch.object(api_editor, name, getattr(editor_utils, name))
                )
            stack.enter_context(
                patch.object(api_editor, "_editor_auth_check", return_value=True)
            )
            stack.enter_context(
                patch.object(api_editor, "_current_user_id", return_value=7)
            )
            stack.enter_context(
                patch.object(
                    api_editor, "playground_read_yaml", return_value=self.SOURCE
                )
            )
            writer = stack.enter_context(
                patch.object(api_editor, "playground_write_yaml")
            )
            client = stack.enter_context(api_editor.app.test_client())
            response = client.post(
                path, json={"project": "default", "filename": "test.yml", **payload}
            )
        return response, writer

    def test_save_block_returns_steps_for_the_correct_order_ids(self):
        from .editor_utils import serialize_order_steps

        response, writer = self._post_order_edit(
            "/al/editor/api/block",
            {
                "block_id": "intro",
                "block_yaml": "id: intro\nquestion: Updated\ncontinue button field: intro\n",
            },
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        steps = response.get_json()["data"]["order_step_map"]
        self.assertEqual(set(steps), {"interview_order_form", "main"})
        self.assertIn(
            "rent_amount", serialize_order_steps(steps["interview_order_form"])
        )
        self.assertNotIn(
            "download", serialize_order_steps(steps["interview_order_form"])
        )
        self.assertIn("download", serialize_order_steps(steps["main"]))
        writer.assert_called_once()

    def test_save_question_without_text_does_not_write(self):
        for block_yaml in (
            'id: intro\nquestion: ""\nfields: []\n',
            "id: intro\nfields: []\n",
        ):
            with self.subTest(block_yaml=block_yaml):
                response, writer = self._post_order_edit(
                    "/al/editor/api/block",
                    {"block_id": "intro", "block_yaml": block_yaml},
                )
                self.assertEqual(response.status_code, 400)
                self.assertIn(
                    "Question text is required",
                    response.get_json()["error"]["message"],
                )
                writer.assert_not_called()

    def test_save_question_turned_into_another_block_type(self):
        for block_yaml in (
            "id: intro\ncode: |\n  intro = True\n",
            "id: intro\nobjects:\n  - user: Individual\n",
            "id: intro\ntemplate: intro\ncontent: Hello\n",
        ):
            with self.subTest(block_yaml=block_yaml):
                response, writer = self._post_order_edit(
                    "/al/editor/api/block",
                    {"block_id": "intro", "block_yaml": block_yaml},
                )
                self.assertEqual(response.status_code, 200, response.get_json())
                writer.assert_called_once()

    def test_save_blank_question_reports_empty_question(self):
        response, writer = self._post_order_edit(
            "/al/editor/api/block",
            {
                "block_id": "intro",
                "block_yaml": 'id: intro\nquestion: ""\ncontinue button field: intro\n',
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["error"]["type"], "empty_question")
        writer.assert_not_called()

    def test_save_blank_question_requires_explicit_opt_in(self):
        response, writer = self._post_order_edit(
            "/al/editor/api/block",
            {
                "block_id": "intro",
                "block_yaml": 'id: intro\nquestion: ""\ncontinue button field: intro\n',
                "allow_empty_question": True,
            },
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertIn('question: ""', writer.call_args.args[-1])

    def test_save_block_without_id_stays_selected_and_gains_no_id(self):
        from .editor_utils import parse_interview_yaml

        self.SOURCE = TestOrderBlockLookup.SOURCE + "---\nquestion: Hi\nfield: hi\n"
        handle = parse_interview_yaml(self.SOURCE)["blocks"][-1]["id"]
        response, writer = self._post_order_edit(
            "/al/editor/api/block",
            {
                "block_id": handle,
                "block_yaml": "question: Hi\nfields:\n  - Name: user_name\n",
                "edit_mode": "graphical",
            },
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        saved = parse_interview_yaml(writer.call_args.args[-1])["blocks"][-1]
        self.assertNotIn("id", saved["data"])
        self.assertNotEqual(saved["id"], handle)
        self.assertEqual(response.get_json()["data"]["saved_block_id"], saved["id"])

    def test_save_order_without_id_updates_first_order_in_place(self):
        from .editor_utils import parse_interview_yaml, parse_order_code

        response, writer = self._post_order_edit(
            "/al/editor/api/order",
            {"steps": parse_order_code("new_question\ninterview_order_form = True\n")},
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        updated = writer.call_args.args[-1]
        before = parse_interview_yaml(self.SOURCE)
        after = parse_interview_yaml(updated)
        self.assertEqual(len(before["blocks"]), len(after["blocks"]))
        for old, new in zip(before["blocks"], after["blocks"]):
            if old["id"] == "interview_order_form":
                self.assertIn("new_question", new["data"]["code"])
                self.assertNotIn("mandatory", new["data"])
            else:
                self.assertEqual(old["yaml"], new["yaml"])

    def test_order_rejects_break_outside_a_loop_without_writing(self):
        response, writer = self._post_order_edit(
            "/al/editor/api/order", {"steps": [{"kind": "break"}]}
        )
        self.assertEqual(response.status_code, 400, response.get_json())
        writer.assert_not_called()

    def test_order_saves_nested_loop_and_assignment_in_place(self):
        from .editor_utils import parse_order_code, parse_interview_yaml

        code = "for person in users:\n  # contact\n  person.email\n  person.complete = True\n"
        response, writer = self._post_order_edit(
            "/al/editor/api/order", {"steps": parse_order_code(code)}
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        blocks = parse_interview_yaml(writer.call_args.args[-1])["blocks"]
        saved = next(b for b in blocks if b["id"] == "interview_order_form")
        self.assertEqual(saved["data"]["code"], code)

    def test_stale_order_id_does_not_append_a_duplicate(self):
        response, writer = self._post_order_edit(
            "/al/editor/api/order", {"order_block_id": "missing", "steps": []}
        )
        self.assertEqual(response.status_code, 400)
        writer.assert_not_called()

    def test_an_order_block_in_the_last_document_is_found(self):
        # Every file that opens with `---` has an empty first document, so the
        # two numberings differ by one; reading an index as a position raised
        # IndexError as soon as the order block was last.
        model = {
            "blocks": [
                {"id": "meta", "index": 1, "data": {"metadata": {}}},
                {"id": "inc", "index": 2, "data": {"include": ["questions.yml"]}},
                {"id": "order", "index": 3, "data": {"code": "rent_amount\n"}},
            ],
            "order_blocks": [3],
        }
        order_step_map, _steps = api_editor._order_steps_from_model(model)
        self.assertEqual(list(order_step_map), ["order"])

    def test_the_right_block_is_read_when_documents_are_skipped(self):
        model = {
            "blocks": [
                {"id": "meta", "index": 1, "data": {"metadata": {}}},
                {"id": "order", "index": 2, "data": {"code": "rent_amount\n"}},
                {"id": "q", "index": 3, "data": {"question": "Hi"}},
            ],
            "order_blocks": [2],
        }
        order_step_map, _steps = api_editor._order_steps_from_model(model)
        self.assertEqual(list(order_step_map), ["order"])


class TestEditorStyleCheckApi(unittest.TestCase):
    def test_missing_dashboard_is_an_actionable_unavailable_response(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "playground_read_yaml",
                return_value="---\nid: q\nquestion: Hi\n",
            ),
            patch.object(
                api_editor, "parse_interview_yaml", return_value={"blocks": []}
            ),
            patch.object(
                api_editor,
                "_run_interview_linter",
                side_effect=api_editor.ALDashboardUnavailable(
                    "Running style checks needs the ALDashboard package. Install "
                    "docassemble.ALDashboard on this server and try again."
                ),
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/weaver/style-check?project=default&filename=main.yml"
            ):
                response = api_editor.editor_api_style_check()

        self.assertEqual(response.status_code, 503)
        error = response.get_json()["error"]
        self.assertEqual(error["type"], "unavailable")
        self.assertIn("Install docassemble.ALDashboard", error["message"])


class TestEditorReviewScreenAndTemplateApi(unittest.TestCase):
    """The two endpoints that lean on ALDashboard at runtime."""

    INTERVIEW = (
        "---\ninclude:\n  - questions.yml\n"
        "---\nid: my review screen\nevent: review_my_form\n"
        "question: |\n  Check your answers\nreview:\n  - Edit: old_variable\n"
        "    button: |\n      **Old**\n"
        "---\nid: download\nevent: my_form_download\nquestion: |\n  Done\n"
    )
    QUESTIONS = "---\nid: q\nquestion: |\n  Q\nfields:\n  - Rent: rent_amount\n"

    def _files(self):
        return {"main.yml": self.INTERVIEW, "questions.yml": self.QUESTIONS}

    def test_sync_reads_the_whole_include_chain_and_replaces_in_place(self):
        files = self._files()
        seen = {}

        def fake_generate(yaml_texts, **kwargs):
            seen["count"] = len(yaml_texts)
            seen.update(kwargs)
            return (
                "id: my review screen\nevent: review_my_form\n"
                "question: |\n  Check your answers\nreview:\n"
                "  - Edit: rent_amount\n    button: |\n      **Rent**\n"
            )

        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "playground_read_yaml",
                side_effect=lambda uid, project, filename: files[filename],
            ),
            patch.object(api_editor, "generate_review_screen_yaml", fake_generate),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/draft-review-screen",
                method="POST",
                json={"project": "default", "filename": "main.yml", "mode": "sync"},
            ):
                response = api_editor.editor_api_draft_review_screen()

        self.assertEqual(response.status_code, 200)
        data = response.get_json()["data"]
        self.assertEqual(data["sources"], ["main.yml", "questions.yml"])
        self.assertEqual(seen["count"], 2)
        # The drafted screen keeps the identity the interview already links to.
        self.assertEqual(seen["event_name"], "review_my_form")
        self.assertEqual(seen["screen_id"], "my review screen")
        self.assertTrue(data["replaced"])
        self.assertIn("Edit: rent_amount", data["full_yaml"])
        self.assertIn("id: download", data["full_yaml"])

        # An entry the draft has no opinion about is carried over rather than
        # dropped: AssemblyLine asks for plenty this generator cannot see, and
        # a review screen that shrinks on every sync is the worse failure.
        self.assertIn("Edit: old_variable", data["full_yaml"])
        self.assertEqual(data["kept_entries"], 1)

        # The drafted block alone does not show what the sync will do to the
        # file, so the response carries the diff the confirmation reads from.
        self.assertIn("+  - Edit: rent_amount", data["diff"]["diff"])
        self.assertFalse(data["diff"]["truncated"])
        self.assertGreater(data["diff"]["added"], 0)
        self.assertFalse(data["unchanged"])
        self.assertTrue(data["revision"])

    def test_a_missing_dashboard_is_a_503_with_something_to_do_about_it(self):
        files = self._files()

        def unavailable(*args, **kwargs):
            raise api_editor.ALDashboardUnavailable("Install docassemble.ALDashboard")

        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "playground_read_yaml",
                side_effect=lambda uid, project, filename: files[filename],
            ),
            patch.object(api_editor, "generate_review_screen_yaml", unavailable),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/draft-review-screen",
                method="POST",
                json={"project": "default", "filename": "main.yml"},
            ):
                response = api_editor.editor_api_draft_review_screen()

        self.assertEqual(response.status_code, 503)
        self.assertIn("ALDashboard", response.get_json()["error"]["message"])

    def test_a_variable_report_lands_in_the_projects_templates_folder(self):
        files = self._files()
        with tempfile.TemporaryDirectory() as tmpdir:
            written = {}

            def fake_write(yaml_texts, output_path, **kwargs):
                written["path"] = output_path
                written["title"] = kwargs.get("report_title")
                with open(output_path, "wb") as handle:
                    handle.write(b"docx")
                return {
                    "variables_count": 4,
                    "list_count": 1,
                    "scalar_count": 3,
                    "size": 4,
                }

            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(
                    api_editor,
                    "playground_read_yaml",
                    side_effect=lambda uid, project, filename: files[filename],
                ),
                patch.object(
                    api_editor,
                    "suggested_report_names",
                    return_value={
                        "title": "Main Draft",
                        "filename": "main_draft.docx",
                    },
                ),
                patch.object(api_editor, "write_variable_report_docx", fake_write),
                patch.object(
                    api_editor,
                    "_editor_storage_directory",
                    return_value=(SimpleNamespace(finalize=lambda: None), tmpdir),
                ),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/template/variable-report",
                    method="POST",
                    json={"project": "default", "filename": "main.yml"},
                ):
                    response = api_editor.editor_api_template_variable_report()

            self.assertEqual(response.status_code, 200)
            data = response.get_json()["data"]
            self.assertEqual(data["section"], "templates")
            self.assertEqual(data["filename"], "main_draft.docx")
            self.assertEqual(data["variables_count"], 4)
            self.assertEqual(data["sources"], ["main.yml", "questions.yml"])
            self.assertEqual(written["title"], "Main Draft")
            self.assertTrue(os.path.exists(written["path"]))

    def test_the_suggestion_reports_the_shapes_this_server_can_draft(self):
        files = self._files()
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(
                api_editor,
                "playground_read_yaml",
                side_effect=lambda uid, project, filename: files[filename],
            ),
            patch.object(
                api_editor,
                "suggested_report_names",
                return_value={"title": "Main Draft", "filename": "main_draft.docx"},
            ),
            patch.object(
                api_editor,
                "court_form_options",
                return_value={
                    "supported": True,
                    "shapes": [
                        {"value": "intake", "label": "Intake summary"},
                        {"value": "motion", "label": "Motion"},
                    ],
                    "profiles": [{"value": "ma_trial_court", "label": "Massachusetts"}],
                },
            ),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/template/variable-report/suggestion"
                "?project=default&filename=main.yml",
                method="GET",
            ):
                response = api_editor.editor_api_template_variable_report_suggestion()

        self.assertEqual(response.status_code, 200)
        data = response.get_json()["data"]
        self.assertTrue(data["court_forms_supported"])
        self.assertIn("motion", {shape["value"] for shape in data["shapes"]})
        self.assertEqual(data["court_profiles"][0]["value"], "ma_trial_court")

    def test_a_court_shape_reaches_the_dashboard_and_comes_back_named(self):
        files = self._files()
        with tempfile.TemporaryDirectory() as tmpdir:
            written = {}

            def fake_write(yaml_texts, output_path, **kwargs):
                written.update(kwargs)
                with open(output_path, "wb") as handle:
                    handle.write(b"docx")
                return {
                    "variables_count": 4,
                    "list_count": 1,
                    "scalar_count": 3,
                    "size": 4,
                    "shape": kwargs.get("shape"),
                    "profile_id": kwargs.get("court_profile"),
                    "profile_name": "Massachusetts Trial Court",
                    "sections": {"caption": "yaml"},
                }

            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(
                    api_editor,
                    "playground_read_yaml",
                    side_effect=lambda uid, project, filename: files[filename],
                ),
                patch.object(
                    api_editor,
                    "suggested_report_names",
                    return_value={
                        "title": "Main Draft",
                        "filename": "main_draft.docx",
                    },
                ),
                patch.object(api_editor, "write_variable_report_docx", fake_write),
                patch.object(
                    api_editor,
                    "_editor_storage_directory",
                    return_value=(SimpleNamespace(finalize=lambda: None), tmpdir),
                ),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/template/variable-report",
                    method="POST",
                    json={
                        "project": "default",
                        "filename": "main.yml",
                        "shape": "motion",
                        "court_profile": "ma_trial_court",
                        "include_certificate_of_service": True,
                        "show_variable_types": True,
                        "max_list_cols": 6,
                        "numbered_paragraphs": False,
                    },
                ):
                    response = api_editor.editor_api_template_variable_report()

            self.assertEqual(response.status_code, 200)
            self.assertEqual(written["shape"], "motion")
            self.assertEqual(written["court_profile"], "ma_trial_court")
            self.assertIs(written["include_certificate_of_service"], True)
            self.assertIs(written["show_variable_types"], True)
            self.assertEqual(written["max_list_cols"], 6)
            self.assertIs(written["numbered_paragraphs"], False)
            data = response.get_json()["data"]
            self.assertEqual(data["profile_name"], "Massachusetts Trial Court")
            self.assertEqual(data["sections"]["caption"], "yaml")

    def test_the_markdown_draft_is_saved_beside_the_docx_when_asked(self):
        files = self._files()
        with tempfile.TemporaryDirectory() as tmpdir:
            written = {}

            def fake_write(yaml_texts, output_path, **kwargs):
                written.update(kwargs)
                with open(output_path, "wb") as handle:
                    handle.write(b"docx")
                markdown_path = kwargs.get("markdown_path")
                if markdown_path:
                    with open(markdown_path, "w", encoding="utf-8") as handle:
                        handle.write("# Main Draft\n")
                    return {
                        "variables_count": 4,
                        "list_count": 1,
                        "scalar_count": 3,
                        "markdown_size": 14,
                    }
                return {"variables_count": 4, "list_count": 1, "scalar_count": 3}

            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(
                    api_editor,
                    "playground_read_yaml",
                    side_effect=lambda uid, project, filename: files[filename],
                ),
                patch.object(
                    api_editor,
                    "suggested_report_names",
                    return_value={
                        "title": "Main Draft",
                        "filename": "main_draft.docx",
                    },
                ),
                patch.object(api_editor, "write_variable_report_docx", fake_write),
                patch.object(
                    api_editor,
                    "_editor_storage_directory",
                    return_value=(SimpleNamespace(finalize=lambda: None), tmpdir),
                ),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/template/variable-report",
                    method="POST",
                    json={
                        "project": "default",
                        "filename": "main.yml",
                        "include_markdown": True,
                    },
                ):
                    response = api_editor.editor_api_template_variable_report()

            self.assertEqual(response.status_code, 200)
            data = response.get_json()["data"]
            self.assertEqual(data["markdown_filename"], "main_draft.md")
            self.assertTrue(os.path.exists(os.path.join(tmpdir, "main_draft.md")))
            self.assertEqual(
                written["markdown_path"], os.path.join(tmpdir, "main_draft.md")
            )

    def test_a_stale_markdown_draft_does_not_outlive_the_docx_it_described(self):
        """Overwriting with a Dashboard that returns no markdown.

        The old `.md` would otherwise sit beside a freshly drafted DOCX still
        claiming to be its text, and an attachment's `content file:` would go
        on assembling it.
        """
        files = self._files()
        with tempfile.TemporaryDirectory() as tmpdir:
            stale = os.path.join(tmpdir, "main_draft.md")
            with open(stale, "w") as handle:
                handle.write("# The previous draft\n")

            def fake_write(yaml_texts, output_path, **kwargs):
                with open(output_path, "wb") as handle:
                    handle.write(b"docx")
                # No markdown_size: this Dashboard produced none.
                return {"variables_count": 4, "list_count": 1, "scalar_count": 3}

            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(
                    api_editor,
                    "playground_read_yaml",
                    side_effect=lambda uid, project, filename: files[filename],
                ),
                patch.object(
                    api_editor,
                    "suggested_report_names",
                    return_value={
                        "title": "Main Draft",
                        "filename": "main_draft.docx",
                    },
                ),
                patch.object(api_editor, "write_variable_report_docx", fake_write),
                patch.object(
                    api_editor,
                    "_editor_storage_directory",
                    return_value=(SimpleNamespace(finalize=lambda: None), tmpdir),
                ),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/template/variable-report",
                    method="POST",
                    json={
                        "project": "default",
                        "filename": "main.yml",
                        "include_markdown": True,
                        "overwrite": True,
                    },
                ):
                    response = api_editor.editor_api_template_variable_report()

            self.assertEqual(response.status_code, 200)
            self.assertFalse(os.path.exists(stale))
            data = response.get_json()["data"]
            self.assertNotIn("markdown_filename", data)
            # Said out loud rather than left for the author to notice.
            self.assertIs(data["markdown_written"], False)

    def test_an_untouched_markdown_draft_is_left_alone(self):
        """Nothing is deleted when no markdown was asked for."""
        files = self._files()
        with tempfile.TemporaryDirectory() as tmpdir:
            kept = os.path.join(tmpdir, "main_draft.md")
            with open(kept, "w") as handle:
                handle.write("# Hand written\n")

            def fake_write(yaml_texts, output_path, **kwargs):
                with open(output_path, "wb") as handle:
                    handle.write(b"docx")
                return {"variables_count": 4, "list_count": 1, "scalar_count": 3}

            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(
                    api_editor,
                    "playground_read_yaml",
                    side_effect=lambda uid, project, filename: files[filename],
                ),
                patch.object(
                    api_editor,
                    "suggested_report_names",
                    return_value={
                        "title": "Main Draft",
                        "filename": "main_draft.docx",
                    },
                ),
                patch.object(api_editor, "write_variable_report_docx", fake_write),
                patch.object(
                    api_editor,
                    "_editor_storage_directory",
                    return_value=(SimpleNamespace(finalize=lambda: None), tmpdir),
                ),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/template/variable-report",
                    method="POST",
                    json={"project": "default", "filename": "main.yml"},
                ):
                    response = api_editor.editor_api_template_variable_report()

            self.assertEqual(response.status_code, 200)
            self.assertTrue(os.path.exists(kept))
            self.assertNotIn("markdown_written", response.get_json()["data"])

    def test_a_nonsense_column_count_is_rejected(self):
        files = self._files()
        with tempfile.TemporaryDirectory() as tmpdir:
            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(
                    api_editor,
                    "playground_read_yaml",
                    side_effect=lambda uid, project, filename: files[filename],
                ),
                patch.object(
                    api_editor,
                    "_editor_storage_directory",
                    return_value=(SimpleNamespace(finalize=lambda: None), tmpdir),
                ),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/template/variable-report",
                    method="POST",
                    json={
                        "project": "default",
                        "filename": "main.yml",
                        "max_list_cols": 40,
                    },
                ):
                    response = api_editor.editor_api_template_variable_report()

            self.assertEqual(response.status_code, 400)
            self.assertIn("between 1 and 12", response.get_json()["error"]["message"])

    def test_an_existing_markdown_draft_is_not_silently_overwritten(self):
        """The DOCX name is free, but the .md it would sit beside is taken."""
        files = self._files()
        with tempfile.TemporaryDirectory() as tmpdir:
            with open(os.path.join(tmpdir, "main_draft.md"), "w") as handle:
                handle.write("hand written")
            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(
                    api_editor,
                    "playground_read_yaml",
                    side_effect=lambda uid, project, filename: files[filename],
                ),
                patch.object(
                    api_editor,
                    "suggested_report_names",
                    return_value={
                        "title": "Main Draft",
                        "filename": "main_draft.docx",
                    },
                ),
                patch.object(
                    api_editor,
                    "_editor_storage_directory",
                    return_value=(SimpleNamespace(finalize=lambda: None), tmpdir),
                ),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/template/variable-report",
                    method="POST",
                    json={
                        "project": "default",
                        "filename": "main.yml",
                        "include_markdown": True,
                    },
                ):
                    response = api_editor.editor_api_template_variable_report()

            self.assertEqual(response.status_code, 400)
            message = response.get_json()["error"]["message"]
            self.assertIn("main_draft.md already exists", message)

    def test_an_existing_template_is_not_silently_overwritten(self):
        files = self._files()
        with tempfile.TemporaryDirectory() as tmpdir:
            with open(os.path.join(tmpdir, "main_draft.docx"), "wb") as handle:
                handle.write(b"already here")
            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(
                    api_editor,
                    "playground_read_yaml",
                    side_effect=lambda uid, project, filename: files[filename],
                ),
                patch.object(
                    api_editor,
                    "suggested_report_names",
                    return_value={
                        "title": "Main Draft",
                        "filename": "main_draft.docx",
                    },
                ),
                patch.object(
                    api_editor,
                    "_editor_storage_directory",
                    return_value=(SimpleNamespace(finalize=lambda: None), tmpdir),
                ),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/template/variable-report",
                    method="POST",
                    json={"project": "default", "filename": "main.yml"},
                ):
                    response = api_editor.editor_api_template_variable_report()

            self.assertEqual(response.status_code, 400)
            self.assertIn("already exists", response.get_json()["error"]["message"])


class TestEditorPackageFileApi(unittest.TestCase):
    """Reading a YAML file out of an installed package, and nothing else."""

    def test_package_exposes_its_named_order_steps(self):
        from .editor_utils import parse_interview_yaml, parse_order_code

        source = "---\ncode: |\n  framework_screen\n  framework_done = True\n"
        with tempfile.TemporaryDirectory() as tmpdir:
            path = os.path.join(tmpdir, "flow.yml")
            with open(path, "w", encoding="utf-8") as handle:
                handle.write(source)
            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(
                    api_editor, "package_question_filename", return_value=path
                ),
                patch.object(api_editor, "parse_interview_yaml", parse_interview_yaml),
                patch.object(api_editor, "parse_order_code", parse_order_code),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/package-file?reference=docassemble.Framework:flow.yml"
                ):
                    response = api_editor.editor_api_get_package_file()

        steps = response.get_json()["data"]["named_order_steps"]["framework_done"]
        self.assertEqual(steps[0]["invoke"], "framework_screen")

    def test_reads_a_question_file_from_an_installed_package(self):
        source = (
            "---\nquestion: |\n  What is your name?\nfields:\n  - Name: x.name.first\n"
        )
        parsed = []

        def fake_parse(text):
            parsed.append(text)
            return {"blocks": [{"id": "b1", "title": "What is your name?"}]}

        with tempfile.TemporaryDirectory() as tmpdir:
            path = os.path.join(tmpdir, "ql_baseline.yml")
            with open(path, "w", encoding="utf-8") as handle:
                handle.write(source)

            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(
                    api_editor, "package_question_filename", return_value=path
                ),
                patch.object(api_editor, "parse_interview_yaml", fake_parse),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/package-file"
                    "?reference=docassemble.AssemblyLine:ql_baseline.yml"
                ):
                    response = api_editor.editor_api_get_package_file()

        payload = response.get_json()
        self.assertTrue(payload["success"])
        self.assertEqual(parsed, [source], "the package's own YAML is what is parsed")
        self.assertEqual(payload["data"]["blocks"][0]["title"], "What is your name?")
        self.assertEqual(
            payload["data"]["reference"], "docassemble.AssemblyLine:ql_baseline.yml"
        )

    def test_refuses_anything_that_is_not_a_package_yaml_reference(self):
        refused = [
            "../../../etc/passwd",
            "docassemble.AssemblyLine:../../../etc/passwd.yml",
            "docassemble.AssemblyLine:a/../b.yml",
            "notdocassemble.Thing:a.yml",
            "docassemble.AssemblyLine:secrets.txt",
            "/etc/passwd",
            "",
        ]
        for reference in refused:
            with self.subTest(reference=reference):
                with (
                    patch.object(api_editor, "_editor_auth_check", return_value=True),
                    patch.object(api_editor, "package_question_filename") as resolver,
                ):
                    with api_editor.app.test_request_context(
                        "/al/editor/api/package-file",
                        query_string={"reference": reference},
                    ):
                        response = api_editor.editor_api_get_package_file()
                self.assertEqual(response.status_code, 400)
                # The resolver is never handed a reference this shape.
                resolver.assert_not_called()

    def test_reports_a_package_that_is_not_installed_as_not_found(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "package_question_filename", return_value=None),
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/package-file"
                "?reference=docassemble.NotInstalled:questions.yml"
            ):
                response = api_editor.editor_api_get_package_file()

        self.assertEqual(response.status_code, 404)
        self.assertFalse(response.get_json()["success"])

    def test_requires_authentication(self):
        with patch.object(api_editor, "_editor_auth_check", return_value=False):
            with api_editor.app.test_request_context(
                "/al/editor/api/package-file"
                "?reference=docassemble.AssemblyLine:ql_baseline.yml"
            ):
                response = api_editor.editor_api_get_package_file()

        self.assertIn(response.status_code, (401, 403))


if __name__ == "__main__":
    unittest.main()


class TestEditorProjectFileNaming(unittest.TestCase):
    """A file a project stores has to have a name Docassemble can resolve.

    https://github.com/SuffolkLITLab/docassemble-ALWeaver/issues/1059
    """

    def test_an_upload_is_renamed_on_its_way_into_the_project(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            upload = FileStorage(
                stream=BytesIO(b"template bytes"),
                filename="93A demand letter (1).docx",
            )
            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(
                    api_editor,
                    "_editor_storage_directory",
                    return_value=(SimpleNamespace(finalize=lambda: None), tmpdir),
                ),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/section-file/upload",
                    method="POST",
                    data={
                        "project": "default",
                        "section": "templates",
                        "files": upload,
                    },
                    content_type="multipart/form-data",
                ):
                    response = api_editor.editor_api_upload_section_file()

            self.assertEqual(response.status_code, 200)
            data = response.get_json()["data"]
            self.assertEqual(data["saved_files"], ["93A_demand_letter_1.docx"])
            self.assertEqual(os.listdir(tmpdir), ["93A_demand_letter_1.docx"])
            # The author has to be told: the name they will refer to the file
            # by in the interview is not the name they uploaded.
            self.assertEqual(
                data["renamed_files"],
                [
                    {
                        "from": "93A demand letter (1).docx",
                        "to": "93A_demand_letter_1.docx",
                        "reason": "unsupported_characters",
                        "message": data["renamed_files"][0]["message"],
                    }
                ],
            )
            self.assertIn(
                "93A demand letter (1).docx", data["renamed_files"][0]["message"]
            )
            self.assertIn(
                "93A_demand_letter_1.docx", data["renamed_files"][0]["message"]
            )

    def test_an_upload_that_needed_no_renaming_reports_none(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            upload = FileStorage(
                stream=BytesIO(b"template bytes"), filename="petition.docx"
            )
            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(
                    api_editor,
                    "_editor_storage_directory",
                    return_value=(SimpleNamespace(finalize=lambda: None), tmpdir),
                ),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/section-file/upload",
                    method="POST",
                    data={
                        "project": "default",
                        "section": "templates",
                        "files": upload,
                    },
                    content_type="multipart/form-data",
                ):
                    response = api_editor.editor_api_upload_section_file()

            data = response.get_json()["data"]
            self.assertEqual(data["saved_files"], ["petition.docx"])
            self.assertEqual(data["renamed_files"], [])

    def test_an_upload_whose_name_is_taken_says_so_rather_than_blaming_the_characters(
        self,
    ):
        with tempfile.TemporaryDirectory() as tmpdir:
            with open(os.path.join(tmpdir, "petition.docx"), "wb") as handle:
                handle.write(b"first")
            upload = FileStorage(stream=BytesIO(b"second"), filename="petition.docx")
            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(
                    api_editor,
                    "_editor_storage_directory",
                    return_value=(SimpleNamespace(finalize=lambda: None), tmpdir),
                ),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/section-file/upload",
                    method="POST",
                    data={
                        "project": "default",
                        "section": "templates",
                        "files": upload,
                    },
                    content_type="multipart/form-data",
                ):
                    response = api_editor.editor_api_upload_section_file()

            data = response.get_json()["data"]
            self.assertEqual(data["saved_files"], ["petition_1.docx"])
            self.assertEqual(data["renamed_files"][0]["reason"], "name_taken")
            self.assertIn(
                "already has a file with that name",
                data["renamed_files"][0]["message"],
            )

    def test_renaming_a_file_to_an_unusable_name_reports_what_it_became(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            with open(os.path.join(tmpdir, "petition.docx"), "wb") as handle:
                handle.write(b"template bytes")
            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(
                    api_editor,
                    "_editor_storage_directory",
                    return_value=(SimpleNamespace(finalize=lambda: None), tmpdir),
                ),
                patch.object(api_editor, "rename_saved_file"),
            ):
                with api_editor.app.test_request_context(
                    "/al/editor/api/section-file/rename",
                    method="POST",
                    json={
                        "project": "default",
                        "section": "templates",
                        "filename": "petition.docx",
                        "new_filename": "demand letter (final).docx",
                    },
                ):
                    response = api_editor.editor_api_rename_section_file()

            data = response.get_json()["data"]
            self.assertEqual(data["filename"], "demand_letter_final.docx")
            self.assertEqual(
                data["renamed_files"][0]["from"], "demand letter (final).docx"
            )
            self.assertEqual(data["renamed_files"][0]["to"], "demand_letter_final.docx")

    def test_binary_template_revision_from_listing_allows_rename_and_delete(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            binary_pdf = b"%PDF-1.4\n\x80\xffbinary template bytes\n%%EOF"
            with open(os.path.join(tmpdir, "petition.pdf"), "wb") as handle:
                handle.write(binary_pdf)
            area = SimpleNamespace(finalize=lambda: None)
            with (
                patch.object(api_editor, "_editor_auth_check", return_value=True),
                patch.object(api_editor, "_current_user_id", return_value=7),
                patch.object(
                    api_editor,
                    "_editor_storage_directory",
                    return_value=(area, tmpdir),
                ),
                patch.object(
                    api_editor, "_file_dependency_conflict", return_value=None
                ),
                patch.object(api_editor, "rename_saved_file") as rename_file,
                patch.object(api_editor, "delete_saved_file") as delete_file,
            ):
                listed = api_editor._list_editor_section_files(
                    7, "default", "templates"
                )
                revision = next(
                    item["revision"]
                    for item in listed
                    if item["filename"] == "petition.pdf"
                )
                self.assertEqual(revision, hashlib.sha256(binary_pdf).hexdigest())

                with api_editor.app.test_request_context(
                    "/al/editor/api/section-file/rename",
                    method="POST",
                    json={
                        "project": "default",
                        "section": "templates",
                        "filename": "petition.pdf",
                        "new_filename": "renamed.pdf",
                        "expected_revision": revision,
                    },
                ):
                    renamed = api_editor.editor_api_rename_section_file()

                with api_editor.app.test_request_context(
                    "/al/editor/api/section-file/delete",
                    method="POST",
                    json={
                        "project": "default",
                        "section": "templates",
                        "filename": "petition.pdf",
                        "expected_revision": revision,
                    },
                ):
                    deleted = api_editor.editor_api_delete_section_file()

            self.assertEqual(renamed.status_code, 200, renamed.get_data(as_text=True))
            self.assertEqual(deleted.status_code, 200, deleted.get_data(as_text=True))
            rename_file.assert_called_once_with(
                area, tmpdir, "petition.pdf", "renamed.pdf"
            )
            delete_file.assert_called_once_with(area, tmpdir, "petition.pdf")

    def test_new_project_collision_notices_match_stored_template_names(self):
        uploads = [
            FileStorage(stream=BytesIO(b"pdf"), filename="petition.pdf"),
            FileStorage(stream=BytesIO(b"docx"), filename="petition.docx"),
            FileStorage(stream=BytesIO(b"first"), filename="demand (1).docx"),
            FileStorage(stream=BytesIO(b"second"), filename="demand_1_.docx"),
        ]

        def validate_upload(**kwargs):
            filename = safe_project_filename(kwargs["filename"])
            return filename, os.path.splitext(filename)[1].lower()

        with (
            patch.object(api_editor, "_editor_async_is_configured", return_value=True),
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "get_list_of_projects", return_value=[]),
            patch.object(
                api_editor, "next_available_project_name", return_value="Collision"
            ),
            patch.object(api_editor, "get_default_github_owner", return_value=""),
            patch.object(api_editor, "create_project"),
            patch.object(
                api_editor, "validate_upload_metadata", side_effect=validate_upload
            ),
            patch.object(api_editor, "validate_document_content"),
            patch.object(
                api_editor,
                "_start_new_project_upload_job",
                return_value={
                    "job_id": "test-job",
                    "job_url": "/job/test-job",
                    "state": {"status": "queued"},
                },
            ) as start_job,
        ):
            with api_editor.app.test_request_context(
                "/al/editor/api/new-project",
                method="POST",
                data={"project_name": "Collision"},
            ):
                response = api_editor._new_project_from_uploads(7, "req-1", uploads)

        self.assertEqual(response.status_code, 202, response.get_json())
        stored_names = [
            upload["filename"]
            for upload in start_job.call_args.kwargs["uploaded_files"]
        ]
        self.assertEqual(
            stored_names,
            ["petition.pdf", "petition.docx", "demand_1.docx", "demand_1_2.docx"],
        )
        renamed = start_job.call_args.kwargs["renamed_files"]
        self.assertEqual(
            [(entry["from"], entry["to"]) for entry in renamed],
            [
                ("demand (1).docx", "demand_1.docx"),
                ("demand_1_.docx", "demand_1_2.docx"),
            ],
        )
        self.assertEqual(renamed[1]["reason"], "name_collision")
        self.assertIn("demand_1_2.docx", renamed[1]["message"])

    def test_importing_an_older_template_renames_it_first(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            stored = os.path.join(tmpdir, "demand letter (1).docx")
            with open(stored, "wb") as handle:
                handle.write(b"template bytes")
            renames = []

            def fake_rename(area, directory, old_name, new_name):
                renames.append((old_name, new_name))
                os.rename(
                    os.path.join(directory, old_name),
                    os.path.join(directory, new_name),
                )

            with (
                patch.object(
                    api_editor,
                    "_editor_storage_directory",
                    return_value=(SimpleNamespace(finalize=lambda: None), tmpdir),
                ),
                patch.object(api_editor, "rename_saved_file", fake_rename),
            ):
                path, filename = api_editor._template_import_target(
                    7, "Eviction", "demand letter (1).docx"
                )

            self.assertEqual(filename, "demand_letter_1.docx")
            self.assertEqual(path, os.path.join(tmpdir, "demand_letter_1.docx"))
            self.assertTrue(os.path.isfile(path))
            self.assertEqual(renames, [("demand letter (1).docx", filename)])

    def test_a_template_whose_safe_name_is_taken_is_reported_rather_than_replaced(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            for name in ("demand letter (1).docx", "demand_letter_1.docx"):
                with open(os.path.join(tmpdir, name), "wb") as handle:
                    handle.write(b"template bytes")

            with (
                patch.object(
                    api_editor,
                    "_editor_storage_directory",
                    return_value=(SimpleNamespace(finalize=lambda: None), tmpdir),
                ),
                patch.object(api_editor, "rename_saved_file") as mock_rename,
            ):
                with self.assertRaises(ValueError):
                    api_editor._template_import_target(
                        7, "Eviction", "demand letter (1).docx"
                    )
            mock_rename.assert_not_called()


class TestMatrixRefactorSafety(unittest.TestCase):
    def test_binary_template_reference_blocks_the_entire_variable_rename(self):
        with (
            patch.object(api_editor, "_editor_auth_check", return_value=True),
            patch.object(api_editor, "_current_user_id", return_value=7),
            patch.object(api_editor, "_project_text_files", return_value=([], [])),
            patch.object(
                api_editor, "_project_search_revision", return_value="current"
            ),
            patch.object(
                api_editor,
                "_binary_template_rename_problems",
                return_value=["binary.docx contains old_name"],
            ),
            patch.object(api_editor, "_commit_project_replacements") as commit,
            api_editor.app.test_request_context(
                "/al/editor/api/project/replace",
                method="POST",
                json={
                    "project": "default",
                    "mode": "variable",
                    "query": "old_name",
                    "replacement": "new_name",
                    "project_revision": "current",
                },
            ),
        ):
            response = api_editor.editor_api_project_replace()
        self.assertEqual(response.status_code, 422)
        self.assertIn("binary.docx", response.get_json()["error"]["message"])
        commit.assert_not_called()

    def test_partial_write_failure_restores_even_the_failing_file(self):
        self._check_rollback(False)

    def test_rollback_failure_returns_original_content_for_recovery(self):
        self._check_rollback(True)

    def _check_rollback(self, fail_restore):
        contents = {"one.yml": "one original", "two.yml": "two original"}
        changes = [
            {
                "section": "interview",
                "filename": name,
                "original": value,
                "updated": "new content",
            }
            for name, value in contents.items()
        ]

        def write(uid, project, section, filename, content):
            if filename == "two.yml" and (content == "new content" or fail_restore):
                contents[filename] = "partially written"
                raise OSError("injected disk write failure")
            contents[filename] = content

        with (
            patch.object(api_editor, "_write_project_text_file", side_effect=write),
            patch.object(
                api_editor,
                "_read_project_text_file",
                side_effect=lambda uid, project, section, filename: contents[filename],
            ),
            self.assertRaises(api_editor.ProjectReplacementWriteError) as raised,
        ):
            api_editor._commit_project_replacements_locked(7, "default", changes)
        self.assertEqual(contents["one.yml"], "one original")
        if fail_restore:
            self.assertEqual(
                raised.exception.recovery,
                [
                    {
                        "section": "interview",
                        "filename": "two.yml",
                        "original_content": "two original",
                    }
                ],
            )
        else:
            self.assertEqual(contents["two.yml"], "two original")
            self.assertEqual(raised.exception.recovery, [])
            self.assertEqual(len(raised.exception.restored), 2)


class TestTemplateFieldReadWithoutInterview(unittest.TestCase):
    def test_docx_fields_can_be_read_without_current_question(self):
        from docx import Document
        import tempfile
        from pathlib import Path

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "binary.docx"
            document = Document()
            document.add_paragraph("Name: {{ old_name }}")
            document.add_paragraph("Again: {{ old_name }}")
            document.save(str(path))
            self.assertEqual(
                api_editor._local_template_field_names(str(path)), ["old_name"]
            )
            with (
                patch.object(
                    api_editor,
                    "_list_editor_section_files",
                    return_value=[
                        {"filename": "binary.docx", "size": path.stat().st_size}
                    ],
                ),
                patch.object(
                    api_editor,
                    "_editor_storage_directory",
                    return_value=(None, directory),
                ),
            ):
                warnings = api_editor._binary_template_rename_problems(
                    7, "default", "old_name"
                )
                self.assertIn("binary.docx contains old_name", warnings[0])
                self.assertEqual(
                    api_editor._binary_template_rename_problems(
                        7, "default", "unrelated"
                    ),
                    [],
                )
